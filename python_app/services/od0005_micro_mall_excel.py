from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = (
    "门店",
    "部门",
    "柜组",
    "柜组名称",
    "销售数量",
    "售价金额",
    "销售收入+总折扣(A)",
    "销售收入",
    "毛利",
    "毛利率",
    "YZQ",
    "其他支付",
    "NZD",
)

DEPARTMENT_HEADERS = (
    "门店",
    "部门",
    "销售数量",
    "应收金额",
    "有赞卡券",
    "礼券",
    "销售收入",
    "销售收入(内转)",
    "毛利额",
    "毛利率",
)

DEPARTMENT_AMOUNT_KEYS = (
    "sales_quantity",
    "price_amount",
    "yzq_amount",
    "other_payment_amount",
    "sales_before_discount",
    "sales_revenue",
    "gross_profit",
)


def _display_department(row: dict[str, Any]) -> str:
    code = str(row.get("department_code") or "").strip()
    name = str(row.get("department_name") or "").strip()
    return " ".join(value for value in (code, name) if value)


def _department_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    departments: dict[tuple[str, str], dict[str, Any]] = {}
    for source in rows:
        department_code = str(source.get("department_code") or "").strip()
        department_name = str(source.get("department_name") or "未匹配").strip()
        key = (department_code, department_name)
        row = departments.setdefault(
            key,
            {
                "store_code": source.get("store_code", ""),
                "department_code": department_code,
                "department_name": department_name,
                **{amount_key: 0.0 for amount_key in DEPARTMENT_AMOUNT_KEYS},
            },
        )
        for amount_key in DEPARTMENT_AMOUNT_KEYS:
            row[amount_key] += float(source.get(amount_key) or 0)

    for row in departments.values():
        sales_revenue = row["sales_revenue"]
        row["gross_margin"] = (
            row["gross_profit"] / sales_revenue if sales_revenue else None
        )
    return list(departments.values())


def _prepare_sheet(
    sheet: Any,
    *,
    report: dict[str, Any],
    title: str,
    headers: tuple[str, ...],
) -> tuple[str, PatternFill, PatternFill, Border]:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = "普灵仕集团百货事业部"
    sheet["A1"].font = Font(name="宋体", size=24, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 38

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = title
    sheet["A2"].font = Font(name="宋体", size=18, bold=True)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 30

    metadata_split = min(6, len(headers) - 1)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=metadata_split)
    sheet["A3"] = f"日期：{report['start_date']}  至  {report['end_date']}"
    store_column = metadata_split + 1
    sheet.merge_cells(
        start_row=3,
        start_column=store_column,
        end_row=3,
        end_column=len(headers),
    )
    store_cell = sheet.cell(3, store_column)
    store_cell.value = f"门店：{report['store_code']} {report['store_name']}"
    for cell in (sheet["A3"], store_cell):
        cell.font = Font(name="宋体", size=11, bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[3].height = 24

    header_fill = PatternFill("solid", fgColor="E2F0D9")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(bottom=thin)
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(4, column, label)
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[4].height = 24
    return last_column, header_fill, total_fill, border


def _add_department_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("部门销售统计")
    last_column, _, total_fill, border = _prepare_sheet(
        sheet,
        report=report,
        title="微商城部门销售统计",
        headers=DEPARTMENT_HEADERS,
    )

    rows = _department_rows(list(report.get("rows", [])))
    for row_index, row in enumerate(rows, start=5):
        values = (
            row.get("store_code", ""),
            _display_department(row),
            row.get("sales_quantity", 0),
            row.get("price_amount", 0),
            row.get("yzq_amount", 0),
            row.get("other_payment_amount", 0),
            row.get("sales_before_discount", 0),
            row.get("sales_revenue", 0),
            row.get("gross_profit", 0),
            row.get("gross_margin"),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="宋体", size=10)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column == 2 else "right" if column >= 3 else "center",
                vertical="center",
            )

    total_row = 5 + len(rows)
    totals = report.get("totals", {})
    total_values = (
        "合计",
        "",
        totals.get("sales_quantity", 0),
        totals.get("price_amount", 0),
        totals.get("yzq_amount", 0),
        totals.get("other_payment_amount", 0),
        totals.get("sales_before_discount", 0),
        totals.get("sales_revenue", 0),
        totals.get("gross_profit", 0),
        totals.get("gross_margin"),
    )
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    for column, value in enumerate(total_values, start=1):
        if column == 2:
            continue
        sheet.cell(total_row, column, value)
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_index in range(5, total_row + 1):
        sheet.cell(row_index, 3).number_format = "#,##0.####"
        for column in range(4, 10):
            sheet.cell(row_index, column).number_format = "#,##0.00;[Red]-#,##0.00"
        sheet.cell(row_index, 10).number_format = "0.00%"

    widths = (9, 28, 12, 15, 15, 15, 15, 18, 15, 11)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.auto_filter.ref = f"A4:{last_column}{max(total_row - 1, 4)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{total_row}"


def _daily_rows(report: dict[str, Any], *, current_only: bool) -> list[dict[str, Any]]:
    rows = list(report.get("daily", {}).get("rows", []))
    if not current_only:
        return rows
    return [
        row
        for row in rows
        if any(float(day.get("sales_current") or 0) != 0 for day in row.get("daily", []))
    ]


def _add_daily_sales_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    daily = report.get("daily", {})
    days = list(daily.get("days", []))
    headers = ("门店", "部门", "柜组", "柜组名称", *(day.get("date", "") for day in days))
    sheet = workbook.create_sheet("逐日销售")
    last_column, _, total_fill, border = _prepare_sheet(
        sheet,
        report=report,
        title="微商城逐日销售",
        headers=headers,
    )

    rows = _daily_rows(report, current_only=True)
    for row_index, row in enumerate(rows, start=5):
        values = (
            row.get("store_code", ""),
            _display_department(row),
            row.get("group_code", ""),
            row.get("group_name", ""),
            *(day.get("sales_current", 0) for day in row.get("daily", [])),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="宋体", size=10)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column in (2, 4) else "right" if column >= 5 else "center",
                vertical="center",
            )

    total_row = 5 + len(rows)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    sheet.cell(total_row, 1, "合计")
    for day_index, day_total in enumerate(daily.get("totals", []), start=5):
        sheet.cell(total_row, day_index, day_total.get("sales_current", 0))
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_index in range(5, total_row + 1):
        for column in range(5, len(headers) + 1):
            sheet.cell(row_index, column).number_format = "#,##0.00;[Red]-#,##0.00"
    for column, width in enumerate((9, 25, 15, 30), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(5, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 14

    sheet.auto_filter.ref = f"A4:{last_column}{max(total_row - 1, 4)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{total_row}"


def _add_daily_yoy_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    daily = report.get("daily", {})
    days = list(daily.get("days", []))
    column_count = 4 + len(days) * 3
    headers = tuple("" for _ in range(column_count))
    sheet = workbook.create_sheet("逐日同期同比")
    last_column, header_fill, total_fill, border = _prepare_sheet(
        sheet,
        report=report,
        title="微商城逐日同期同比",
        headers=headers,
    )
    sheet.freeze_panes = "A6"
    sheet["A3"] = (
        f"本期：{report['start_date']} 至 {report['end_date']}；"
        f"同期：{daily.get('prior_start_date', '')} 至 {daily.get('prior_end_date', '')}"
    )

    fixed_headers = ("门店", "部门", "柜组", "柜组名称")
    for column, label in enumerate(fixed_headers, start=1):
        sheet.merge_cells(start_row=4, start_column=column, end_row=5, end_column=column)
        sheet.cell(4, column, label)
    for day_index, day in enumerate(days):
        start_column = 5 + day_index * 3
        sheet.merge_cells(
            start_row=4,
            start_column=start_column,
            end_row=4,
            end_column=start_column + 2,
        )
        sheet.cell(4, start_column, day.get("date", ""))
        for offset, label in enumerate(("本期销售", "同期销售", "同比")):
            sheet.cell(5, start_column + offset, label)

    for row_index in (4, 5):
        for column in range(1, column_count + 1):
            cell = sheet.cell(row_index, column)
            cell.fill = header_fill
            cell.font = Font(name="宋体", size=10, bold=True)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[5].height = 24

    rows = _daily_rows(report, current_only=False)
    for row_index, row in enumerate(rows, start=6):
        values: list[Any] = [
            row.get("store_code", ""),
            _display_department(row),
            row.get("group_code", ""),
            row.get("group_name", ""),
        ]
        for day in row.get("daily", []):
            values.extend(
                (
                    day.get("sales_current", 0),
                    day.get("sales_prior", 0),
                    day.get("sales_yoy"),
                )
            )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="宋体", size=10)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column in (2, 4) else "right" if column >= 5 else "center",
                vertical="center",
            )

    total_row = 6 + len(rows)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    sheet.cell(total_row, 1, "合计")
    for day_index, day_total in enumerate(daily.get("totals", [])):
        start_column = 5 + day_index * 3
        sheet.cell(total_row, start_column, day_total.get("sales_current", 0))
        sheet.cell(total_row, start_column + 1, day_total.get("sales_prior", 0))
        sheet.cell(total_row, start_column + 2, day_total.get("sales_yoy"))
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_index in range(6, total_row + 1):
        for day_index in range(len(days)):
            start_column = 5 + day_index * 3
            sheet.cell(row_index, start_column).number_format = "#,##0.00;[Red]-#,##0.00"
            sheet.cell(row_index, start_column + 1).number_format = "#,##0.00;[Red]-#,##0.00"
            sheet.cell(row_index, start_column + 2).number_format = "0.00%"
    for column, width in enumerate((9, 25, 15, 30), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(5, column_count + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12

    sheet.auto_filter.ref = f"A5:{last_column}{max(total_row - 1, 5)}"
    sheet.print_title_rows = "1:5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{total_row}"


def _add_brand_yoy_sheet(workbook: Workbook, report: dict[str, Any]) -> None:
    comparison = report.get("brand_yoy", {})
    periods = comparison.get("periods", {})
    current = periods.get("current", {})
    prior = periods.get("prior", {})
    two_year_prior = periods.get("two_year_prior", {})
    current_year = int(current.get("year") or str(report.get("start_date", "0000"))[:4] or 0)
    prior_year = int(prior.get("year") or current_year - 1)
    two_year_prior_year = int(two_year_prior.get("year") or current_year - 2)
    headers = (
        "部门",
        "品牌",
        f"{str(two_year_prior_year)[-2:]}年",
        f"{str(prior_year)[-2:]}年",
        f"{str(current_year)[-2:]}年",
        "差额",
        "同比",
        "原因分析",
        "缺失商品",
    )
    start_date = str(report.get("start_date", ""))
    end_date = str(report.get("end_date", ""))
    title = f"{str(current_year)[-2:]}年品牌销售情况及同比"
    if len(start_date) >= 7 and start_date[:7] == end_date[:7]:
        title = f"{str(current_year)[-2:]}年{int(start_date[5:7])}月品牌销售情况及同比"

    sheet = workbook.create_sheet("品牌同比")
    last_column, _, total_fill, border = _prepare_sheet(
        sheet,
        report=report,
        title=title,
        headers=headers,
    )
    sheet["A3"] = (
        f"前两年：{two_year_prior.get('start_date', '')} 至 {two_year_prior.get('end_date', '')}；"
        f"上年：{prior.get('start_date', '')} 至 {prior.get('end_date', '')}；"
        f"本期：{current.get('start_date', start_date)} 至 {current.get('end_date', end_date)}"
    )

    rows = list(comparison.get("rows", []))
    for row_index, row in enumerate(rows, start=5):
        values = (
            _display_department(row),
            row.get("group_name", ""),
            row.get("sales_two_year_prior", 0),
            row.get("sales_prior", 0),
            row.get("sales_current", 0),
            row.get("difference", 0),
            row.get("yoy"),
            "",
            "",
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="宋体", size=10)
            if column in (6, 7) and isinstance(value, (int, float)) and value < 0:
                cell.font = Font(name="宋体", size=10, color="FF0000")
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if 3 <= column <= 7 else "left",
                vertical="center",
            )

    total_row = 5 + len(rows)
    totals = comparison.get("totals", {})
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    sheet.cell(total_row, 1, "合计")
    for column, key in enumerate(
        ("sales_two_year_prior", "sales_prior", "sales_current", "difference", "yoy"),
        start=3,
    ):
        sheet.cell(total_row, column, totals.get(key))
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_index in range(5, total_row + 1):
        for column in range(3, 7):
            sheet.cell(row_index, column).number_format = "#,##0.00;[Red]-#,##0.00"
        sheet.cell(row_index, 7).number_format = "0.00%"
    for column, width in enumerate((28, 30, 14, 14, 14, 14, 12, 20, 20), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.auto_filter.ref = f"A4:{last_column}{max(total_row - 1, 4)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{total_row}"


def build_od0005_workbook(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "微商城品牌销售统计"
    last_column, _, total_fill, border = _prepare_sheet(
        sheet,
        report=report,
        title="微商城品牌销售统计",
        headers=HEADERS,
    )

    rows = report.get("rows", [])
    for row_index, row in enumerate(rows, start=5):
        values = (
            row.get("store_code", ""),
            _display_department(row),
            row.get("group_code", ""),
            row.get("group_name", ""),
            row.get("sales_quantity", 0),
            row.get("price_amount", 0),
            row.get("sales_before_discount", 0),
            row.get("sales_revenue", 0),
            row.get("gross_profit", 0),
            row.get("gross_margin"),
            row.get("yzq_amount", 0),
            row.get("other_payment_amount", 0),
            row.get("nzd_amount", 0),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="宋体", size=10)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column in (2, 4) else "right" if column >= 5 else "center",
                vertical="center",
            )

    total_row = 5 + len(rows)
    totals = report.get("totals", {})
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    total_values = (
        "合计",
        totals.get("sales_quantity", 0),
        totals.get("price_amount", 0),
        totals.get("sales_before_discount", 0),
        totals.get("sales_revenue", 0),
        totals.get("gross_profit", 0),
        totals.get("gross_margin"),
        totals.get("yzq_amount", 0),
        totals.get("other_payment_amount", 0),
        totals.get("nzd_amount", 0),
    )
    sheet.cell(total_row, 1, total_values[0])
    for offset, value in enumerate(total_values[1:], start=5):
        sheet.cell(total_row, offset, value)
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_index in range(5, total_row + 1):
        sheet.cell(row_index, 5).number_format = "#,##0.####"
        for column in range(6, 10):
            sheet.cell(row_index, column).number_format = "#,##0.00;[Red]-#,##0.00"
        sheet.cell(row_index, 10).number_format = "0.00%"
        for column in range(11, 14):
            sheet.cell(row_index, column).number_format = "#,##0.00;[Red]-#,##0.00"

    widths = (9, 25, 15, 30, 12, 15, 20, 15, 15, 11, 15, 15, 15)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.auto_filter.ref = f"A4:{last_column}{max(total_row - 1, 4)}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{total_row}"

    _add_department_sheet(workbook, report)
    _add_daily_sales_sheet(workbook, report)
    _add_daily_yoy_sheet(workbook, report)
    _add_brand_yoy_sheet(workbook, report)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_od0005_workbook_file(report: dict[str, Any]) -> BytesIO:
    output = BytesIO(build_od0005_workbook(report))
    output.seek(0)
    return output
