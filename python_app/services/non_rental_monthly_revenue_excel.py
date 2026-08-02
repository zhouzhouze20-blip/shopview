from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REGULAR_DIMENSION_HEADERS = ("大类", "部门", "楼层", "区域名", "类别名", "品牌厅")
SPECIAL_DIMENSION_HEADERS = ("大类", "部门", "楼层", "区域名", "类别名", "品牌", "品牌厅")
METRICS = (
    ("tax_included_sales", "销售含税", "amount"),
    ("tax_excluded_sales", "销售不含税", "amount"),
    ("gross_profit", "毛利额", "amount"),
    ("fee", "收费", "amount"),
    ("contribution", "贡献值", "amount"),
    ("gross_margin", "毛利率", "rate"),
    ("contract_profit", "合同毛利", "amount"),
    ("concession_loss", "让利损失额", "amount"),
    ("concession_loss_rate", "让利损失率", "rate"),
)

HEADER_GREEN = "C6E0B4"
CONTRIBUTION_BLUE = "9DC3E6"
TOTAL_BLUE = "B4C6E7"
WHITE = "FFFFFF"
RED = "FF0000"
GRID = Side(style="thin", color="7F7F7F")
MEDIUM = Side(style="medium", color="606060")
BASE_BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _amount_wan(value: Any) -> float:
    return float(value or 0) / 10_000


def _dimension_values(
    row: dict[str, Any],
    dimension_headers: tuple[str, ...],
) -> list[Any]:
    include_brand = "品牌" in dimension_headers
    if row.get("row_type") in {"department_total", "special_total", "store_total"}:
        values = [None, None, row.get("floor_name") or None, None, None]
        if include_brand:
            values.append(None)
        values.append(row.get("group_name"))
        return values

    values = [
        row.get("big_category"),
        row.get("department_name"),
        row.get("floor_name"),
        row.get("area_name"),
        row.get("category_name"),
    ]
    if include_brand:
        values.append(row.get("brand_name"))
    values.append(row.get("group_name"))
    return values


def _metric_values(metrics: dict[str, Any]) -> list[Any]:
    values: list[Any] = []
    for key, _, kind in METRICS:
        value = metrics.get(key)
        values.append(value if kind == "rate" else _amount_wan(value))
    return values


def _write_headers(
    sheet,
    financial_year: int,
    dimension_headers: tuple[str, ...],
) -> None:
    metric_count = len(METRICS)
    final_column = len(dimension_headers) + metric_count * 13

    for column, label in enumerate(dimension_headers, 1):
        sheet.merge_cells(
            start_row=1,
            start_column=column,
            end_row=3,
            end_column=column,
        )
        cell = sheet.cell(1, column, label)
        cell.fill = PatternFill("solid", fgColor=HEADER_GREEN)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    block_labels = [f"{financial_year}年小计"] + [f"{month}月份" for month in range(1, 13)]
    for block_index, block_label in enumerate(block_labels):
        start_column = len(dimension_headers) + 1 + block_index * metric_count
        end_column = start_column + metric_count - 1
        sheet.merge_cells(
            start_row=1,
            start_column=start_column,
            end_row=1,
            end_column=end_column,
        )
        sheet.cell(1, start_column, block_label)
        for metric_index, (_, label, _) in enumerate(METRICS):
            column = start_column + metric_index
            sheet.merge_cells(
                start_row=2,
                start_column=column,
                end_row=3,
                end_column=column,
            )
            sheet.cell(2, column, label)
            if metric_index == 4:
                sheet.cell(2, column).fill = PatternFill("solid", fgColor=CONTRIBUTION_BLUE)

    for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=final_column):
        for cell in row:
            if cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor=HEADER_GREEN)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BASE_BORDER


def _write_row(
    sheet,
    row_number: int,
    row: dict[str, Any],
    dimension_headers: tuple[str, ...],
) -> None:
    metric_count = len(METRICS)
    is_total = row.get("row_type") in {"department_total", "special_total", "store_total"}
    values = _dimension_values(row, dimension_headers)
    values.extend(_metric_values(row.get("annual", {})))
    for month in range(1, 13):
        values.extend(_metric_values(row.get("months", {}).get(str(month), {})))

    for column, value in enumerate(values, 1):
        cell = sheet.cell(row_number, column, _safe_text(value))
        cell.border = BASE_BORDER
        cell.alignment = Alignment(
            horizontal="right" if column > len(dimension_headers) else "center",
            vertical="center",
        )
        if column > len(dimension_headers):
            metric_index = (column - len(dimension_headers) - 1) % metric_count
            _, _, kind = METRICS[metric_index]
            cell.number_format = "0.00%" if kind == "rate" else '0.00;[Red](0.00);-'
            if metric_index == 4:
                cell.fill = PatternFill("solid", fgColor=CONTRIBUTION_BLUE)
            if metric_index in (7, 8) and isinstance(value, (int, float)) and value < 0:
                cell.font = Font(color=RED, bold=is_total)
        if is_total:
            cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
            cell.font = Font(bold=True, color=cell.font.color if cell.font.color else "000000")

    if is_total:
        for cell in sheet[row_number]:
            cell.border = Border(
                left=cell.border.left,
                right=cell.border.right,
                top=MEDIUM,
                bottom=cell.border.bottom,
            )


def _write_notes(sheet, report: dict[str, Any]) -> None:
    rows = [
        ("报表口径", "非租赁品牌月度收益"),
        ("财务年度", str(report["financial_year"])),
        ("统计期间", f"{report['dates']['start_date']} 至 {report['dates']['end_date']}"),
        ("月份口径", "1月为1月1日至1月28日；2—11月为上月29日至本月28日；12月为11月29日至12月31日。"),
        ("买单额", "第一版暂不展示；贡献值＝毛利额＋收费。"),
        ("特卖", "按柜组楼层编码16识别，在独立工作表按品牌展示；柜组收费无品牌维度，列为“收费未分配品牌”且不重复分摊。"),
        ("销售口径", "含税销售＝Σ(sglxssr＋sglpfsr)；不含税销售逐笔按销售税率还原。"),
        ("毛利口径", "经销按销售与进项税率分拆计算；其他非租赁方式按sgln2去税。"),
        ("收费口径", "按supsetcharge费用发生月顺延一个报表月，剔除38/61/94/95，并按CODECHARGE税率逐笔去税。"),
        ("合同毛利", "逐笔不含税销售×原合同扣率sglbasekl。"),
        ("让利损失", "毛利额－合同毛利；让利损失率＝让利损失额÷不含税销售。"),
        ("数据范围", report.get("scope_description", "当前用户权限范围")),
    ]
    sheet.append(["项目", "说明"])
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 110
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_GREEN)
        cell.font = Font(bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = BASE_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _write_report_sheet(
    sheet,
    report: dict[str, Any],
    rows: list[dict[str, Any]],
    *,
    is_special: bool = False,
) -> None:
    dimension_headers = (
        SPECIAL_DIMENSION_HEADERS if is_special else REGULAR_DIMENSION_HEADERS
    )
    _write_headers(sheet, int(report["financial_year"]), dimension_headers)
    row_number = 4
    for row in rows:
        _write_row(sheet, row_number, row, dimension_headers)
        row_number += 1

    first_metric_column = len(dimension_headers) + 1
    sheet.freeze_panes = f"{get_column_letter(first_metric_column)}4"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[3].height = 8
    widths = (12, 16, 9, 14, 16, 22, 28) if is_special else (12, 16, 9, 14, 16, 28)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in range(first_metric_column, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12


def build_non_rental_monthly_revenue_workbook(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    regular_sheet = workbook.active
    regular_sheet.title = f"{report['financial_year']}年"
    _write_report_sheet(regular_sheet, report, report.get("regular_rows", []))

    special_sheet = workbook.create_sheet(f"{report['financial_year']}年特卖")
    _write_report_sheet(
        special_sheet,
        report,
        report.get("special_rows", []),
        is_special=True,
    )

    notes = workbook.create_sheet("口径说明")
    _write_notes(notes, report)

    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    workbook.save(output)
    output.seek(0)
    payload = output.read()
    output.close()
    return payload


def build_non_rental_monthly_revenue_workbook_file(report: dict[str, Any]):
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    output.write(build_non_rental_monthly_revenue_workbook(report))
    output.seek(0)
    return output
