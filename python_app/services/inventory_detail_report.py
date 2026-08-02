from __future__ import annotations

from datetime import date
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session


MAX_EXPORT_ROWS = 50_000

CATEGORY_NAME_SQL = (
    "CASE "
    "WHEN gc.catcname ~ '[\u0080-\u00ff]' "
    "THEN convert_from(convert_to(gc.catcname, 'LATIN1'), 'GBK') "
    "ELSE gc.catcname END"
)

INVENTORY_COLUMNS: tuple[tuple[str, str], ...] = (
    ("store_display", "门店"),
    ("floor_display", "楼层"),
    ("area_display", "库区"),
    ("goods_code", "商品代码"),
    ("barcode", "商品条码"),
    ("brand_display", "品牌"),
    ("goods_name", "商品名称"),
    ("specification", "规格"),
    ("supplier_display", "供应商"),
    ("operation_method", "经营方式"),
    ("group_display", "柜组"),
    ("subinventory_display", "子库存"),
    ("category_display", "商品类别"),
    ("sample_label", "样品?"),
    ("selling_price", "售价"),
    ("average_purchase_price_tax_included", "平均进价(含税)"),
    ("inventory_quantity", "库存数量"),
    ("retail_amount", "零售金额"),
    ("inventory_purchase_amount_tax_excluded", "库存不含税进价金额"),
    ("inventory_purchase_amount_tax_included", "库存含税进价金额"),
)

HISTORICAL_INVENTORY_COLUMNS: tuple[tuple[str, str], ...] = (
    ("inventory_date", "日期"),
    ("area_display", "库区"),
    ("supplier_display", "供应商"),
    ("store_display", "门店"),
    ("group_display", "柜组"),
    ("subinventory_display", "子库存"),
    ("operation_method", "经营方式"),
    ("goods_code", "商品编码"),
    ("barcode", "商品条码"),
    ("goods_name", "商品名称"),
    ("specification", "规格型号"),
    ("brand_display", "品牌"),
    ("category_display", "商品类别"),
    ("sample_label", "样品?"),
    ("selling_price", "售价"),
    ("inventory_quantity", "库存数量"),
    ("inventory_purchase_amount_tax_included", "库存含税进价金额"),
    ("inventory_purchase_amount_tax_excluded", "库存不含税进价金额"),
    ("retail_amount", "库存售价金额"),
)

INVENTORY_MOVEMENT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("store_display", "门店"),
    ("group_display", "柜组"),
    ("supplier_display", "供应商"),
    ("goods_name", "商品名称"),
    ("goods_code", "商品代码"),
    ("barcode", "商品条码"),
    ("subinventory_display", "子库存"),
    ("accounting_date", "记账日期"),
    ("transaction_label", "摘要"),
    ("purchase_price_tax_included", "含税进价"),
    ("selling_price", "售价"),
    ("operation_method", "经营方式"),
    ("sample_label", "样品?"),
    ("batch_sequence", "批次序号"),
    ("increase_quantity", "数量"),
    ("increase_amount_tax_included", "含税进价金额"),
    ("increase_amount_tax_excluded", "不含税进价金额"),
    ("decrease_quantity", "数量"),
    ("decrease_amount_tax_included", "含税进价金额"),
    ("decrease_amount_tax_excluded", "不含税进价金额"),
    ("balance_quantity", "数量"),
    ("balance_cost_tax_included", "含税成本金额"),
    ("balance_cost_tax_excluded", "不含税成本金额"),
    ("memo", "备注"),
    ("specification", "规格"),
    ("manufacturer_article_number", "厂商货号"),
)


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _mapping(row: Mapping[str, Any]) -> dict[str, Any]:
    return {key: _json_value(value) for key, value in row.items()}


def _code_name_display(code_sql: str, name_sql: str) -> str:
    return (
        "CASE "
        f"WHEN NULLIF(TRIM(BOTH FROM COALESCE(({code_sql})::text, '')), '') IS NOT NULL "
        f"AND NULLIF(TRIM(BOTH FROM COALESCE(({name_sql})::text, '')), '') IS NOT NULL "
        f"THEN '[' || TRIM(BOTH FROM ({code_sql})::text) || '] ' || TRIM(BOTH FROM ({name_sql})::text) "
        f"WHEN NULLIF(TRIM(BOTH FROM COALESCE(({code_sql})::text, '')), '') IS NOT NULL "
        f"THEN '[' || TRIM(BOTH FROM ({code_sql})::text) || ']' "
        f"ELSE TRIM(BOTH FROM COALESCE(({name_sql})::text, '')) END"
    )


def build_inventory_filter_sql(filters: Mapping[str, Any], params: dict[str, Any]) -> str:
    clauses: list[str] = [] if filters.get("include_zero") else ["gs.gstkcsl > 0"]

    def fuzzy(key: str, expressions: tuple[str, ...]) -> None:
        value = str(filters.get(key) or "").strip()
        if not value:
            return
        param_name = f"{key}_like"
        params[param_name] = f"%{value}%"
        clauses.append(
            "(" + " OR ".join(
                f"UPPER(TRIM(BOTH FROM COALESCE(({expr})::text, ''))) LIKE UPPER(:{param_name})"
                for expr in expressions
            ) + ")"
        )

    store_id = str(filters.get("store_id") or "").strip()
    if store_id:
        params["store_id"] = store_id
        clauses.append("(st.store_id::text = :store_id OR gs.gstmarket = :store_id)")

    fuzzy("supplier", ("gs.gstsupid", "sb.sbcname"))
    fuzzy("group", ("gs.gstmfid", "mf.mfcname"))
    fuzzy("goods_code", ("gs.gstgdid",))
    fuzzy("goods_name", ("gb.gbcname",))
    fuzzy("barcode", ("gb.gbbarcode",))

    exact_group = str(filters.get("exact_group") or "").strip()
    if exact_group:
        params["exact_group"] = exact_group
        clauses.append(
            "UPPER(TRIM(BOTH FROM COALESCE(gs.gstmfid, ''))) = UPPER(:exact_group)"
        )

    exact_code = str(filters.get("exact_code") or "").strip()
    if exact_code:
        params["exact_code"] = exact_code
        exact_matches = [
            "UPPER(TRIM(BOTH FROM COALESCE(gb.gbbarcode, ''))) = UPPER(:exact_code)",
            "UPPER(TRIM(BOTH FROM COALESCE(gs.gstgdid, ''))) = UPPER(:exact_code)",
        ]
        if filters.get("has_goodsbarcode"):
            exact_matches.append(
                "EXISTS ("
                "SELECT 1 FROM goodsbarcode gc "
                "WHERE gc.gcgdid = gb.gbid "
                "AND UPPER(TRIM(BOTH FROM COALESCE(gc.gcbarcode, ''))) = UPPER(:exact_code)"
                ")"
            )
        clauses.append("(" + " OR ".join(exact_matches) + ")")

    return "" if not clauses else " AND " + " AND ".join(clauses)


def inventory_base_sql(filter_sql: str, scope_sql: str) -> str:
    store_display = _code_name_display("gs.gstmarket", "COALESCE(st.store_name, store_node.mfcname)")
    floor_display = _code_name_display("floor_node.mfcode", "floor_node.mfcname")
    area_display = _code_name_display("area_node.mfcode", "area_node.mfcname")
    group_display = _code_name_display("gs.gstmfid", "mf.mfcname")
    supplier_display = _code_name_display("gs.gstsupid", "sb.sbcname")
    brand_display = _code_name_display("gb.gbppcode", "cb.cbcname")
    category_display = _code_name_display("gb.gbcatcode", CATEGORY_NAME_SQL)
    subinventory_name = (
        "CASE gs.gstmd WHEN '1' THEN '柜台' WHEN '2' THEN 'D01' ELSE '' END"
    )
    subinventory_display = _code_name_display("gs.gstmd", subinventory_name)

    return f"""
      SELECT
        st.store_id::text AS scope_store_id,
        gs.gstmarket AS store_code,
        COALESCE(st.store_name, store_node.mfcname, '') AS store_name,
        {store_display} AS store_display,
        floor_node.mfcode AS floor_code,
        floor_node.mfcname AS floor_name,
        {floor_display} AS floor_display,
        area_node.mfcode AS area_code,
        area_node.mfcname AS area_name,
        {area_display} AS area_display,
        gs.gstmfid AS group_code,
        mf.mfcname AS group_name,
        {group_display} AS group_display,
        gs.gstgdid AS goods_code,
        gb.gbbarcode AS barcode,
        gb.gbcname AS goods_name,
        gb.gbspec AS specification,
        gb.gbppcode AS brand_code,
        cb.cbcname AS brand_name,
        {brand_display} AS brand_display,
        gs.gstsupid AS supplier_code,
        sb.sbcname AS supplier_name,
        {supplier_display} AS supplier_display,
        gs.gstwmid AS operation_method_code,
        CASE gs.gstwmid
          WHEN '1' THEN '经销'
          WHEN '2' THEN '成本代销'
          WHEN '3' THEN '扣率代销'
          WHEN '4' THEN '联营'
          WHEN '5' THEN '租赁'
          ELSE COALESCE(gs.gstwmid, '')
        END AS operation_method,
        gs.gstmd AS subinventory_code,
        {subinventory_display} AS subinventory_display,
        gb.gbcatcode AS category_code,
        {CATEGORY_NAME_SQL} AS category_name,
        {category_display} AS category_display,
        gs.gstsample AS sample_code,
        CASE WHEN gs.gstsample = 'Y' THEN '是' ELSE '' END AS sample_label,
        gmp.gmpsj AS selling_price,
        CASE
          WHEN COALESCE(gs.gstkcsl, 0) = 0 THEN gs.gsthsjj
          ELSE gs.gstkcsuphsje / NULLIF(gs.gstkcsl, 0)
        END AS average_purchase_price_tax_included,
        gs.gstkcsl AS inventory_quantity,
        CASE
          WHEN gb.gbmanamode = '1' THEN gmp.gmpsj * gs.gstkcsl
          ELSE gs.gstkcsuphsje / NULLIF(
            1 - CASE WHEN gmp.gmpkl = 1 THEN 0 ELSE gmp.gmpkl END,
            0
          )
        END AS retail_amount,
        gs.gstkcbhsjjje AS inventory_purchase_amount_tax_excluded,
        gs.gstkchsjjje AS inventory_purchase_amount_tax_included
      FROM goodsstock gs
      JOIN goodsbase gb ON gb.gbid = gs.gstgdid
      JOIN goodsmfprice gmp
        ON gmp.gmpgdid = gs.gstgdid
       AND gmp.gmpmfid = gs.gstmfid
       AND gmp.gmpuid = '00'
      JOIN manaframe mf ON mf.mfcode = gs.gstmfid
      LEFT JOIN manaframe area_node ON area_node.mfcode = mf.mfpcode
      LEFT JOIN manaframe floor_node ON floor_node.mfcode = area_node.mfpcode
      LEFT JOIN manaframe store_node ON store_node.mfcode = floor_node.mfpcode
      LEFT JOIN stores st ON st.store_code = gs.gstmarket
      LEFT JOIN supplierbase sb ON sb.sbid = gs.gstsupid
      LEFT JOIN codebrand cb ON cb.cbid = gb.gbppcode
      LEFT JOIN goodscat gc ON gc.catcode = gb.gbcatcode
      WHERE 1=1 {filter_sql} {scope_sql}
    """


def build_historical_inventory_filter_sql(
    filters: Mapping[str, Any],
    params: dict[str, Any],
) -> str:
    clauses: list[str] = []

    start_date = filters.get("start_date")
    end_date = filters.get("end_date")
    if start_date:
        params["start_date"] = start_date
        clauses.append("gs.gstdate >= CAST(:start_date AS date)")
    if end_date:
        params["end_date"] = end_date
        clauses.append("gs.gstdate < CAST(:end_date AS date) + INTERVAL '1 day'")

    def fuzzy(key: str, expressions: tuple[str, ...]) -> None:
        value = str(filters.get(key) or "").strip()
        if not value:
            return
        param_name = f"{key}_like"
        params[param_name] = f"%{value}%"
        clauses.append(
            "(" + " OR ".join(
                f"UPPER(TRIM(BOTH FROM COALESCE(({expr})::text, ''))) LIKE UPPER(:{param_name})"
                for expr in expressions
            ) + ")"
        )

    store_id = str(filters.get("store_id") or "").strip()
    if store_id:
        params["store_id"] = store_id
        clauses.append("(st.store_id::text = :store_id OR gs.gstmarket = :store_id)")

    fuzzy("supplier", ("gs.gstsupid", "sb.sbcname"))
    fuzzy("group", ("gs.gstmfid", "mf.mfcname"))
    fuzzy("goods_code", ("gs.gstgdid",))
    fuzzy("goods_name", ("gb.gbcname",))
    fuzzy("barcode", ("gb.gbbarcode",))

    return "" if not clauses else " AND " + " AND ".join(clauses)


def historical_inventory_base_sql(filter_sql: str, scope_sql: str) -> str:
    store_display = _code_name_display("gs.gstmarket", "COALESCE(st.store_name, store_node.mfcname)")
    area_display = _code_name_display("area_node.mfcode", "area_node.mfcname")
    group_display = _code_name_display("gs.gstmfid", "mf.mfcname")
    supplier_display = _code_name_display("gs.gstsupid", "sb.sbcname")
    brand_display = _code_name_display("gb.gbppcode", "cb.cbcname")
    category_display = _code_name_display("gb.gbcatcode", CATEGORY_NAME_SQL)
    subinventory_name = "CASE gs.gstmd WHEN '1' THEN '柜台' WHEN '2' THEN 'D01' ELSE '' END"
    subinventory_display = _code_name_display("gs.gstmd", subinventory_name)

    return f"""
      WITH historical_rows AS (
        SELECT
          gs.gstdate::date AS inventory_date,
          st.store_id::text AS scope_store_id,
          gs.gstmarket AS store_code,
          {store_display} AS store_display,
          area_node.mfcode AS area_code,
          {area_display} AS area_display,
          gs.gstmfid AS group_code,
          {group_display} AS group_display,
          gs.gstgdid AS goods_code,
          gb.gbbarcode AS barcode,
          gb.gbcname AS goods_name,
          gb.gbspec AS specification,
          {brand_display} AS brand_display,
          gs.gstsupid AS supplier_code,
          {supplier_display} AS supplier_display,
          gs.gstwmid AS operation_method_code,
          CASE gs.gstwmid
            WHEN '1' THEN '经销'
            WHEN '2' THEN '成本代销'
            WHEN '3' THEN '扣率代销'
            WHEN '4' THEN '联营'
            WHEN '5' THEN '租赁'
            ELSE COALESCE(gs.gstwmid, '')
          END AS operation_method,
          gs.gstmd AS subinventory_code,
          {subinventory_display} AS subinventory_display,
          {category_display} AS category_display,
          gs.gstsample AS sample_code,
          CASE WHEN gs.gstsample = 'Y' THEN '是' ELSE '' END AS sample_label,
          COALESCE(NULLIF(gs.gstsj, 0), gmp.gmpsj) AS selling_price,
          gs.gstkcsl AS inventory_quantity,
          gs.gstkchsjjje AS inventory_purchase_amount_tax_included,
          gs.gstkcbhsjjje AS inventory_purchase_amount_tax_excluded
        FROM goodsstock_bak gs
        LEFT JOIN goodsbase gb ON gb.gbid = gs.gstgdid
        LEFT JOIN goodsmfprice gmp
          ON gmp.gmpgdid = gs.gstgdid
         AND gmp.gmpmfid = gs.gstmfid
         AND gmp.gmpuid = '00'
        JOIN manaframe mf ON mf.mfcode = gs.gstmfid
        LEFT JOIN manaframe area_node ON area_node.mfcode = mf.mfpcode
        LEFT JOIN manaframe floor_node ON floor_node.mfcode = area_node.mfpcode
        LEFT JOIN manaframe store_node ON store_node.mfcode = floor_node.mfpcode
        LEFT JOIN stores st ON st.store_code = gs.gstmarket
        LEFT JOIN supplierbase sb ON sb.sbid = gs.gstsupid
        LEFT JOIN codebrand cb ON cb.cbid = gb.gbppcode
        LEFT JOIN goodscat gc ON gc.catcode = gb.gbcatcode
        WHERE 1=1 {filter_sql} {scope_sql}
      )
      SELECT
        inventory_date,
        scope_store_id,
        store_code,
        store_display,
        area_code,
        area_display,
        group_code,
        group_display,
        goods_code,
        barcode,
        goods_name,
        specification,
        brand_display,
        supplier_code,
        supplier_display,
        operation_method_code,
        operation_method,
        subinventory_code,
        subinventory_display,
        category_display,
        sample_code,
        sample_label,
        selling_price,
        SUM(inventory_quantity) AS inventory_quantity,
        SUM(inventory_purchase_amount_tax_included) AS inventory_purchase_amount_tax_included,
        SUM(inventory_purchase_amount_tax_excluded) AS inventory_purchase_amount_tax_excluded,
        selling_price * SUM(inventory_quantity) AS retail_amount
      FROM historical_rows
      GROUP BY
        inventory_date, scope_store_id, store_code, store_display,
        area_code, area_display, group_code, group_display,
        goods_code, barcode, goods_name, specification, brand_display,
        supplier_code, supplier_display, operation_method_code, operation_method,
        subinventory_code, subinventory_display, category_display,
        sample_code, sample_label, selling_price
    """


def build_inventory_movement_filter_sql(
    filters: Mapping[str, Any],
    params: dict[str, Any],
) -> str:
    clauses: list[str] = []

    start_date = filters.get("start_date")
    end_date = filters.get("end_date")
    if start_date:
        params["start_date"] = start_date
        clauses.append("j.jglfsdate >= CAST(:start_date AS date)")
    if end_date:
        params["end_date"] = end_date
        clauses.append("j.jglfsdate < CAST(:end_date AS date) + INTERVAL '1 day'")

    accounting_start_date = filters.get("accounting_start_date")
    accounting_end_date = filters.get("accounting_end_date")
    if accounting_start_date:
        params["accounting_start_date"] = accounting_start_date
        clauses.append("j.jgldate >= CAST(:accounting_start_date AS date)")
    if accounting_end_date:
        params["accounting_end_date"] = accounting_end_date
        clauses.append(
            "j.jgldate < CAST(:accounting_end_date AS date) + INTERVAL '1 day'"
        )

    def fuzzy(key: str, expressions: tuple[str, ...]) -> None:
        value = str(filters.get(key) or "").strip()
        if not value:
            return
        param_name = f"{key}_like"
        params[param_name] = f"%{value}%"
        clauses.append(
            "(" + " OR ".join(
                f"UPPER(TRIM(BOTH FROM COALESCE(({expr})::text, ''))) LIKE UPPER(:{param_name})"
                for expr in expressions
            ) + ")"
        )

    fuzzy("store", ("j.jglmarket", "st.store_name", "store_node.mfcname"))
    fuzzy("group", ("j.jglmfid", "mf.mfcname"))
    fuzzy("supplier", ("j.jglsupid", "sb.sbcname"))
    fuzzy("goods_code", ("j.jglgdid",))
    fuzzy("goods_name", ("gb.gbcname",))
    fuzzy("barcode", ("gb.gbbarcode",))
    fuzzy("specification", ("gb.gbspec",))

    subinventory = str(filters.get("subinventory") or "").strip()
    if subinventory:
        params["subinventory"] = subinventory
        clauses.append(
            "UPPER(TRIM(BOTH FROM COALESCE(j.jglmd, ''))) = UPPER(:subinventory)"
        )

    return "" if not clauses else " AND " + " AND ".join(clauses)


def inventory_movement_base_sql(filter_sql: str, scope_sql: str) -> str:
    store_display = _code_name_display(
        "j.jglmarket",
        "COALESCE(st.store_name, store_node.mfcname)",
    )
    group_display = _code_name_display("j.jglmfid", "mf.mfcname")
    supplier_display = _code_name_display("j.jglsupid", "sb.sbcname")
    subinventory_name = (
        "CASE TRIM(BOTH FROM j.jglmd) WHEN '1' THEN '柜台' "
        "WHEN '2' THEN 'D01' ELSE '' END"
    )
    subinventory_display = _code_name_display("j.jglmd", subinventory_name)
    adjustment_multiplier = (
        "CASE WHEN TRIM(BOTH FROM j.jgltran) IN ('W', 'X') THEN 0 ELSE 1 END"
    )
    increase_tax_included = (
        f"j.jglhsjjje + j.jgln13 * {adjustment_multiplier}"
    )
    increase_tax_excluded = (
        f"j.jglbhsjjje + j.jgln15 * {adjustment_multiplier}"
    )

    return f"""
      SELECT
        st.store_id::text AS scope_store_id,
        j.jglmarket AS store_code,
        {store_display} AS store_display,
        area_node.mfcode AS department_code,
        area_node.mfcname AS department_name,
        floor_node.mfcode AS floor_code,
        floor_node.mfcname AS floor_name,
        j.jglmfid AS group_code,
        mf.mfcname AS group_name,
        {group_display} AS group_display,
        j.jglsupid AS supplier_code,
        sb.sbcname AS supplier_name,
        {supplier_display} AS supplier_display,
        j.jglgdid AS goods_code,
        gb.gbcname AS goods_name,
        gb.gbbarcode AS barcode,
        gb.gbspec AS specification,
        gb.gbcshh AS manufacturer_article_number,
        j.jglcatid AS category_code,
        j.jglppcode AS brand_code,
        j.jglgdtype AS goods_type,
        j.jglmd AS subinventory_code,
        {subinventory_display} AS subinventory_display,
        j.jgldate::date AS accounting_date,
        j.jglfsdate::date AS occurrence_date,
        TRIM(BOTH FROM j.jgltran) AS transaction_code,
        CASE TRIM(BOTH FROM j.jgltran)
          WHEN '1' THEN '进货'
          WHEN 'E' THEN '销售'
          WHEN 'F' THEN '退货'
          WHEN 'U' THEN '进货调整'
          WHEN 'W' THEN '销售调整'
          ELSE TRIM(BOTH FROM COALESCE(j.jgltran, ''))
        END AS transaction_label,
        j.jglhsjj AS purchase_price_tax_included,
        j.jglbhsjj AS purchase_price_tax_excluded,
        j.jglsj AS selling_price,
        j.jgln2 AS tax_rate,
        j.jglkl AS discount_rate,
        j.jglwmid AS operation_method_code,
        CASE TRIM(BOTH FROM j.jglwmid)
          WHEN '1' THEN '经销'
          WHEN '2' THEN '成本代销'
          WHEN '3' THEN '扣率代销'
          WHEN '4' THEN '联营'
          WHEN '5' THEN '租赁'
          ELSE TRIM(BOTH FROM COALESCE(j.jglwmid, ''))
        END AS operation_method,
        COALESCE(NULLIF(TRIM(BOTH FROM j.jglsample), ''), 'N') AS sample_code,
        CASE WHEN COALESCE(NULLIF(TRIM(BOTH FROM j.jglsample), ''), 'N') = 'Y'
          THEN '是' ELSE '' END AS sample_label,
        j.jglbatchseq AS batch_sequence,
        CASE WHEN j.jgldac = 'D' THEN j.jglsl ELSE 0 END AS increase_quantity,
        CASE WHEN j.jgldac = 'D' THEN {increase_tax_included} ELSE 0 END
          AS increase_amount_tax_included,
        CASE WHEN j.jgldac = 'D' THEN {increase_tax_excluded} ELSE 0 END
          AS increase_amount_tax_excluded,
        CASE WHEN j.jgldac = 'C' THEN j.jglsl ELSE 0 END AS decrease_quantity,
        CASE WHEN j.jgldac = 'C' THEN {increase_tax_included} ELSE 0 END
          AS decrease_amount_tax_included,
        CASE WHEN j.jgldac = 'C' THEN {increase_tax_excluded} ELSE 0 END
          AS decrease_amount_tax_excluded,
        j.jglqmsl AS balance_quantity,
        j.jglqmhscbje AS balance_cost_tax_included,
        j.jglqmbhscbje AS balance_cost_tax_excluded,
        j.jglbillno || ' ' || TO_CHAR(j.jglfsdate, 'YYYY-MM-DD') AS memo,
        j.jglhsjjje AS movement_purchase_amount_tax_included,
        j.jglsjje AS movement_retail_amount,
        j.jglseq AS sequence
      FROM jxcgoodslist j
      JOIN goodsbase gb ON gb.gbid = j.jglgdid
      LEFT JOIN manaframe mf ON mf.mfcode = j.jglmfid
      LEFT JOIN manaframe area_node ON area_node.mfcode = mf.mfpcode
      LEFT JOIN manaframe floor_node ON floor_node.mfcode = area_node.mfpcode
      LEFT JOIN manaframe store_node ON store_node.mfcode = floor_node.mfpcode
      LEFT JOIN stores st ON st.store_code = j.jglmarket
      LEFT JOIN supplierbase sb ON sb.sbid = j.jglsupid
      LEFT JOIN codebrand cb ON cb.cbid = j.jglppcode
      LEFT JOIN goodscat gc ON gc.catcode = j.jglcatid
      WHERE 1=1 {filter_sql} {scope_sql}
    """


def load_inventory_detail_report(
    db: Session,
    *,
    filters: Mapping[str, Any],
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
    offset: int,
) -> dict[str, Any]:
    params: dict[str, Any] = dict(scope_params)
    filter_sql = build_inventory_filter_sql(filters, params)
    base_sql = inventory_base_sql(filter_sql, scope_sql)

    summary = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql})
            SELECT
              COUNT(*) AS total_count,
              COALESCE(SUM(inventory_quantity), 0) AS inventory_quantity,
              COALESCE(SUM(retail_amount), 0) AS retail_amount,
              COALESCE(SUM(inventory_purchase_amount_tax_excluded), 0)
                AS inventory_purchase_amount_tax_excluded,
              COALESCE(SUM(inventory_purchase_amount_tax_included), 0)
                AS inventory_purchase_amount_tax_included
            FROM inventory
            """
        ),
        params,
    ).mappings().one()

    page_params = {**params, "limit": limit, "offset": offset}
    rows = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql})
            SELECT *
            FROM inventory
            ORDER BY store_code, floor_code, area_code, group_code, goods_code,
                     subinventory_code, supplier_code, sample_code
            LIMIT :limit OFFSET :offset
            """
        ),
        page_params,
    ).mappings().all()

    return {
        "rows": [_mapping(row) for row in rows],
        "summary": _mapping(summary),
        "limit": limit,
        "offset": offset,
        "source_note": (
            "沿用原 ERP SQL 的商品、柜组售价内连接口径；"
            "VIEW_MFRAME_ALL 由 manaframe 三级父级关系替代。"
        ),
    }


def load_historical_inventory_detail_report(
    db: Session,
    *,
    filters: Mapping[str, Any],
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
    offset: int,
) -> dict[str, Any]:
    params: dict[str, Any] = dict(scope_params)
    filter_sql = build_historical_inventory_filter_sql(filters, params)
    base_sql = historical_inventory_base_sql(filter_sql, scope_sql)

    summary = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql})
            SELECT
              COUNT(*) AS total_count,
              COALESCE(SUM(inventory_quantity), 0) AS inventory_quantity,
              COALESCE(SUM(retail_amount), 0) AS retail_amount,
              COALESCE(SUM(inventory_purchase_amount_tax_excluded), 0)
                AS inventory_purchase_amount_tax_excluded,
              COALESCE(SUM(inventory_purchase_amount_tax_included), 0)
                AS inventory_purchase_amount_tax_included
            FROM inventory
            """
        ),
        params,
    ).mappings().one()

    page_params = {**params, "limit": limit, "offset": offset}
    rows = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql})
            SELECT *
            FROM inventory
            ORDER BY inventory_date, store_code, area_code, group_code, goods_code,
                     subinventory_code, supplier_code, sample_code
            LIMIT :limit OFFSET :offset
            """
        ),
        page_params,
    ).mappings().all()

    return {
        "rows": [_mapping(row) for row in rows],
        "summary": _mapping(summary),
        "limit": limit,
        "offset": offset,
        "source_note": (
            "数据来自 goodsstock_bak；库存日期按所选自然日闭区间查询，"
            "数量和进价金额按原 ERP 历史库存 SQL 的维度汇总，库存售价金额=售价×库存数量。"
        ),
    }


def load_inventory_movement_detail_report(
    db: Session,
    *,
    filters: Mapping[str, Any],
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
    offset: int,
) -> dict[str, Any]:
    params: dict[str, Any] = dict(scope_params)
    filter_sql = build_inventory_movement_filter_sql(filters, params)
    base_sql = inventory_movement_base_sql(filter_sql, scope_sql)

    summary = db.execute(
        text(
            f"""
            WITH movements AS ({base_sql})
            SELECT
              COUNT(*) AS total_count,
              COALESCE(SUM(increase_quantity), 0) AS increase_quantity,
              COALESCE(SUM(increase_amount_tax_included), 0)
                AS increase_amount_tax_included,
              COALESCE(SUM(increase_amount_tax_excluded), 0)
                AS increase_amount_tax_excluded,
              COALESCE(SUM(decrease_quantity), 0) AS decrease_quantity,
              COALESCE(SUM(decrease_amount_tax_included), 0)
                AS decrease_amount_tax_included,
              COALESCE(SUM(decrease_amount_tax_excluded), 0)
                AS decrease_amount_tax_excluded
            FROM movements
            """
        ),
        params,
    ).mappings().one()

    page_params = {**params, "limit": limit, "offset": offset}
    rows = db.execute(
        text(
            f"""
            WITH movements AS ({base_sql})
            SELECT *
            FROM movements
            ORDER BY sequence
            LIMIT :limit OFFSET :offset
            """
        ),
        page_params,
    ).mappings().all()

    coverage = db.execute(
        text(
            """
            SELECT
              MIN(jglfsdate)::date AS min_occurrence_date,
              MAX(jglfsdate)::date AS max_occurrence_date,
              MAX(jgldate)::date AS max_accounting_date
            FROM jxcgoodslist
            """
        )
    ).mappings().one()
    max_occurrence_date = _json_value(coverage.get("max_occurrence_date"))
    coverage_note = (
        f"；当前本地同步数据覆盖至 {max_occurrence_date}"
        if max_occurrence_date
        else "；当前本地同步表暂无数据"
    )

    return {
        "rows": [_mapping(row) for row in rows],
        "summary": _mapping(summary),
        "coverage": _mapping(coverage),
        "limit": limit,
        "offset": offset,
        "source_note": (
            "数据来自本地 jxcgoodslist，由 PAPI 从 Oracle DBUSRMKT.JXCGOODSLIST "
            "增量同步；增/减金额沿用原 SQL 的 JGLN13/JGLN15 与 W/X 调整规则，"
            "期末金额为成本金额，不能按明细行直接累计"
            + coverage_note
            + "。"
        ),
    }


def load_inventory_movement_filter_options(
    db: Session,
    *,
    field: str,
    query: str,
    filters: Mapping[str, Any],
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
) -> list[dict[str, str]]:
    option_columns = {
        "supplier": ("supplier_code", "supplier_display", "supplier_code", "supplier_name"),
        "group": ("group_code", "group_display", "group_code", "group_name"),
        "goods_code": ("goods_code", "goods_code", "goods_code", "goods_name"),
        "goods_name": ("goods_name", "goods_name", "goods_code", "goods_name"),
        "barcode": ("barcode", "barcode", "barcode", "goods_name"),
    }
    if field not in option_columns:
        raise ValueError(f"Unsupported inventory movement option field: {field}")

    params: dict[str, Any] = dict(scope_params)
    option_filters = {**filters, field: query}
    filter_sql = build_inventory_movement_filter_sql(option_filters, params)
    base_sql = inventory_movement_base_sql(filter_sql, scope_sql)
    value_column, display_column, code_column, name_column = option_columns[field]
    params["option_limit"] = limit

    rows = db.execute(
        text(
            f"""
            WITH movements AS ({base_sql}),
            options AS (
              SELECT DISTINCT
                TRIM(BOTH FROM COALESCE({value_column}::text, '')) AS value,
                TRIM(BOTH FROM COALESCE({display_column}::text, '')) AS display_value,
                TRIM(BOTH FROM COALESCE({code_column}::text, '')) AS code,
                TRIM(BOTH FROM COALESCE({name_column}::text, '')) AS name
              FROM movements
              WHERE NULLIF(TRIM(BOTH FROM COALESCE({value_column}::text, '')), '') IS NOT NULL
            )
            SELECT
              value,
              CASE
                WHEN field_value.name <> '' AND field_value.code <> ''
                  THEN '[' || field_value.code || '] ' || field_value.name
                WHEN field_value.name <> '' THEN field_value.name
                ELSE field_value.display_value
              END AS label,
              code,
              name
            FROM options field_value
            ORDER BY label, value
            LIMIT :option_limit
            """
        ),
        params,
    ).mappings().all()
    return [
        {
            "value": str(row["value"] or ""),
            "label": str(row["label"] or row["value"] or ""),
            "code": str(row["code"] or ""),
            "name": str(row["name"] or ""),
        }
        for row in rows
    ]


def load_inventory_filter_options(
    db: Session,
    *,
    field: str,
    query: str,
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
) -> list[dict[str, str]]:
    option_columns = {
        "supplier": ("supplier_code", "supplier_display", "supplier_code", "supplier_name"),
        "group": ("group_code", "group_display", "group_code", "group_name"),
        "goods_code": ("goods_code", "goods_code", "goods_code", "goods_name"),
        "goods_name": ("goods_name", "goods_name", "goods_code", "goods_name"),
        "barcode": ("barcode", "barcode", "barcode", "goods_name"),
    }
    if field not in option_columns:
        raise ValueError(f"Unsupported inventory option field: {field}")

    params: dict[str, Any] = dict(scope_params)
    filters: dict[str, Any] = {field: query}
    filter_sql = build_inventory_filter_sql(filters, params)
    base_sql = inventory_base_sql(filter_sql, scope_sql)
    value_column, display_column, code_column, name_column = option_columns[field]
    params["option_limit"] = limit

    rows = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql}),
            options AS (
              SELECT DISTINCT
                TRIM(BOTH FROM COALESCE({value_column}::text, '')) AS value,
                TRIM(BOTH FROM COALESCE({display_column}::text, '')) AS display_value,
                TRIM(BOTH FROM COALESCE({code_column}::text, '')) AS code,
                TRIM(BOTH FROM COALESCE({name_column}::text, '')) AS name
              FROM inventory
              WHERE NULLIF(TRIM(BOTH FROM COALESCE({value_column}::text, '')), '') IS NOT NULL
            )
            SELECT
              value,
              CASE
                WHEN field_value.name <> '' AND field_value.code <> ''
                  THEN '[' || field_value.code || '] ' || field_value.name
                WHEN field_value.name <> '' THEN field_value.name
                ELSE field_value.display_value
              END AS label,
              code,
              name
            FROM options field_value
            ORDER BY label, value
            LIMIT :option_limit
            """
        ),
        params,
    ).mappings().all()
    return [
        {
            "value": str(row["value"] or ""),
            "label": str(row["label"] or row["value"] or ""),
            "code": str(row["code"] or ""),
            "name": str(row["name"] or ""),
        }
        for row in rows
    ]


def load_historical_inventory_filter_options(
    db: Session,
    *,
    field: str,
    query: str,
    filters: Mapping[str, Any],
    scope_sql: str,
    scope_params: Mapping[str, Any],
    limit: int,
) -> list[dict[str, str]]:
    option_columns = {
        "supplier": ("supplier_code", "supplier_display", "supplier_code", "supplier_display"),
        "group": ("group_code", "group_display", "group_code", "group_display"),
        "goods_code": ("goods_code", "goods_code", "goods_code", "goods_name"),
        "goods_name": ("goods_name", "goods_name", "goods_code", "goods_name"),
        "barcode": ("barcode", "barcode", "barcode", "goods_name"),
    }
    if field not in option_columns:
        raise ValueError(f"Unsupported historical inventory option field: {field}")

    params: dict[str, Any] = dict(scope_params)
    option_filters = {**filters, field: query}
    filter_sql = build_historical_inventory_filter_sql(option_filters, params)
    base_sql = historical_inventory_base_sql(filter_sql, scope_sql)
    value_column, display_column, code_column, name_column = option_columns[field]
    params["option_limit"] = limit

    rows = db.execute(
        text(
            f"""
            WITH inventory AS ({base_sql}),
            options AS (
              SELECT DISTINCT
                TRIM(BOTH FROM COALESCE({value_column}::text, '')) AS value,
                TRIM(BOTH FROM COALESCE({display_column}::text, '')) AS display_value,
                TRIM(BOTH FROM COALESCE({code_column}::text, '')) AS code,
                TRIM(BOTH FROM COALESCE({name_column}::text, '')) AS name
              FROM inventory
              WHERE NULLIF(TRIM(BOTH FROM COALESCE({value_column}::text, '')), '') IS NOT NULL
            )
            SELECT value, display_value AS label, code, name
            FROM options
            ORDER BY label, value
            LIMIT :option_limit
            """
        ),
        params,
    ).mappings().all()
    return [
        {
            "value": str(row["value"] or ""),
            "label": str(row["label"] or row["value"] or ""),
            "code": str(row["code"] or ""),
            "name": str(row["name"] or ""),
        }
        for row in rows
    ]


def _safe_excel_text(value: Any) -> Any:
    text_value = "" if value is None else str(value)
    if text_value.startswith(("=", "+", "-", "@")):
        return "'" + text_value
    return text_value


def build_inventory_workbook_file(
    report: Mapping[str, Any],
    *,
    filter_description: str,
    scope_description: str,
) -> SpooledTemporaryFile:
    return _build_inventory_workbook_file(
        report,
        title="实时库存查询",
        columns=INVENTORY_COLUMNS,
        filter_description=filter_description,
        scope_description=scope_description,
        widths=(8, 20, 20, 24, 14, 18, 22, 28, 14, 34, 12, 28, 14, 18, 10, 12, 16, 14, 16, 18, 18),
    )


def build_historical_inventory_workbook_file(
    report: Mapping[str, Any],
    *,
    filter_description: str,
    scope_description: str,
) -> SpooledTemporaryFile:
    return _build_inventory_workbook_file(
        report,
        title="历史库存明细报表",
        columns=HISTORICAL_INVENTORY_COLUMNS,
        filter_description=filter_description,
        scope_description=scope_description,
        widths=(8, 13, 24, 34, 20, 28, 14, 12, 14, 18, 30, 14, 22, 18, 10, 12, 14, 18, 18, 18),
    )


def build_inventory_movement_workbook_file(
    report: Mapping[str, Any],
    *,
    filter_description: str,
    scope_description: str,
) -> SpooledTemporaryFile:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品进销存明细报表"
    sheet.sheet_view.showGridLines = False

    last_column = len(INVENTORY_MOVEMENT_COLUMNS) + 1
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet["A1"] = "商品进销存明细报表"
    sheet["A1"].font = Font(size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet["A2"] = f"查询条件：{filter_description or '全部'}"
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    sheet["A3"] = scope_description
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=last_column)
    sheet["A4"] = str(report.get("source_note") or "")
    for cell_ref in ("A2", "A3", "A4"):
        sheet[cell_ref].alignment = Alignment(wrap_text=True, vertical="center")

    header_fill = PatternFill("solid", fgColor="4472C4")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    static_headers = ["行号", *(label for _, label in INVENTORY_MOVEMENT_COLUMNS[:14])]
    for column, label in enumerate(static_headers, 1):
        sheet.merge_cells(start_row=5, start_column=column, end_row=6, end_column=column)
        cell = sheet.cell(5, column, label)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    group_specs = (
        ("增", 16, 18, ("数量", "含税进价金额", "不含税进价金额")),
        ("减", 19, 21, ("数量", "含税进价金额", "不含税进价金额")),
        ("存", 22, 24, ("数量", "含税成本金额", "不含税成本金额")),
    )
    for label, start_column, end_column, subheaders in group_specs:
        sheet.merge_cells(
            start_row=5,
            start_column=start_column,
            end_row=5,
            end_column=end_column,
        )
        sheet.cell(5, start_column, label).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        for offset, subheader in enumerate(subheaders):
            sheet.cell(6, start_column + offset, subheader).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    trailing_headers = [label for _, label in INVENTORY_MOVEMENT_COLUMNS[23:]]
    for column, label in enumerate(trailing_headers, 25):
        sheet.merge_cells(start_row=5, start_column=column, end_row=6, end_column=column)
        sheet.cell(5, column, label).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_number in (5, 6):
        for column in range(1, last_column + 1):
            cell = sheet.cell(row_number, column)
            cell.fill = header_fill
            cell.font = white_font
            cell.border = border
    sheet.row_dimensions[5].height = 24
    sheet.row_dimensions[6].height = 36

    numeric_keys = {
        "purchase_price_tax_included",
        "selling_price",
        "batch_sequence",
        "increase_quantity",
        "increase_amount_tax_included",
        "increase_amount_tax_excluded",
        "decrease_quantity",
        "decrease_amount_tax_included",
        "decrease_amount_tax_excluded",
        "balance_quantity",
        "balance_cost_tax_included",
        "balance_cost_tax_excluded",
    }
    money_keys = {
        "selling_price",
        "purchase_price_tax_included",
    }
    identifier_keys = {
        "goods_code",
        "barcode",
        "batch_sequence",
        "manufacturer_article_number",
    }

    for row_number, row in enumerate(report.get("rows", []), 7):
        sheet.cell(row_number, 1, row_number - 6)
        for column, (key, _label) in enumerate(INVENTORY_MOVEMENT_COLUMNS, 2):
            value = row.get(key)
            if key == "accounting_date" and isinstance(value, str):
                try:
                    value = date.fromisoformat(value[:10])
                except ValueError:
                    pass
            if key not in numeric_keys and key != "accounting_date":
                value = _safe_excel_text(value)
            cell = sheet.cell(row_number, column, value)
            if key == "accounting_date":
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif key in identifier_keys:
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="center")
            elif key in numeric_keys:
                cell.number_format = "0.00" if key in money_keys else "0.0000"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="center")
            cell.border = border

    total_row = 7 + len(report.get("rows", []))
    sheet.cell(total_row, 1, "合计")
    summary = report.get("summary", {})
    summary_columns = {
        "increase_quantity": 16,
        "increase_amount_tax_included": 17,
        "increase_amount_tax_excluded": 18,
        "decrease_quantity": 19,
        "decrease_amount_tax_included": 20,
        "decrease_amount_tax_excluded": 21,
    }
    for column in range(1, last_column + 1):
        cell = sheet.cell(total_row, column)
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = border
    for key, column in summary_columns.items():
        sheet.cell(total_row, column, summary.get(key, 0))
        sheet.cell(total_row, column).number_format = "0.0000"

    widths = (
        8, 20, 28, 34, 24, 14, 18, 14, 13, 14, 12, 12, 14, 10, 16,
        14, 18, 18, 14, 18, 18, 14, 18, 18, 24, 18, 18,
    )
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:{last_letter}{max(6, total_row - 1)}"
    sheet.print_title_rows = "1:6"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    output = SpooledTemporaryFile(max_size=12 * 1024 * 1024, mode="w+b")
    try:
        workbook.save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise


def _build_inventory_workbook_file(
    report: Mapping[str, Any],
    *,
    title: str,
    columns: tuple[tuple[str, str], ...],
    filter_description: str,
    scope_description: str,
    widths: tuple[int, ...],
) -> SpooledTemporaryFile:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.sheet_view.showGridLines = False

    last_column = len(columns) + 1
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet["A1"] = title
    sheet["A1"].font = Font(size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet["A2"] = f"查询条件：{filter_description or '全部'}"
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    sheet["A3"] = scope_description
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet["A3"].alignment = Alignment(wrap_text=True)

    header_fill = PatternFill("solid", fgColor="4472C4")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["行号", *(label for _, label in columns)]
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(5, column, label)
        cell.fill = header_fill
        cell.font = white_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[5].height = 36

    number_formats = {
        "selling_price": "0.00",
        "average_purchase_price_tax_included": "0.0000",
        "inventory_quantity": "0.0000",
        "retail_amount": "0.00",
        "inventory_purchase_amount_tax_excluded": "0.0000",
        "inventory_purchase_amount_tax_included": "0.0000",
    }

    for row_number, row in enumerate(report.get("rows", []), 6):
        sheet.cell(row_number, 1, row_number - 5)
        for column, (key, _label) in enumerate(columns, 2):
            value = row.get(key)
            if key == "inventory_date" and isinstance(value, str):
                try:
                    value = date.fromisoformat(value[:10])
                except ValueError:
                    pass
            if key not in number_formats and key != "inventory_date":
                value = _safe_excel_text(value)
            cell = sheet.cell(row_number, column, value)
            if key == "inventory_date":
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif key not in number_formats:
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="center")
            else:
                cell.number_format = number_formats[key]
                cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = border

    total_row = 6 + len(report.get("rows", []))
    sheet.cell(total_row, 1, "合计")
    summary = report.get("summary", {})
    summary_columns = {
        key: next(
            (index + 2 for index, (column_key, _label) in enumerate(columns) if column_key == key),
            0,
        )
        for key in (
            "inventory_quantity",
            "retail_amount",
            "inventory_purchase_amount_tax_excluded",
            "inventory_purchase_amount_tax_included",
        )
    }
    for column in range(1, last_column + 1):
        cell = sheet.cell(total_row, column)
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = border
    for key, column in summary_columns.items():
        if not column:
            continue
        sheet.cell(total_row, column, summary.get(key, 0))
        sheet.cell(total_row, column).number_format = number_formats[key]

    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{last_letter}{max(5, total_row - 1)}"
    sheet.print_title_rows = "1:5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    output = SpooledTemporaryFile(max_size=12 * 1024 * 1024, mode="w+b")
    try:
        workbook.save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise
