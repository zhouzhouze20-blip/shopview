from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
RATE_FORMAT = '0.00%'

HEADERS = [
    "门店编码", "门店名称", "部门编码", "部门名称", "柜组编码", "柜组名称",
    "供应商编码", "供应商名称", "合同号", "结算单号", "结算截止日起", "结算截止日止",
    "结算单应收未收", "行号", "行状态", "项目编码", "项目名称", "项目期间起",
    "项目期间止", "财务月", "项目金额", "已核金额", "已收款", "抵扣金额", "明细余额",
    "应收影响", "是否计入应收", "是否销售返款", "调整金额", "销售参考金额", "税率",
    "不含税金额", "备注", "计算来源", "是否预收", "建筑面积", "租赁面积",
]

MONEY_COLUMNS = {13, 21, 22, 23, 24, 25, 26, 29, 30, 32, 36, 37}


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _row_values(row: dict[str, Any]) -> list[Any]:
    return [
        row.get("store_code"), row.get("store_name"), row.get("department_code"),
        row.get("department_name"), row.get("group_code"), row.get("group_name"),
        row.get("supplier_id"), row.get("supplier_name"), row.get("contract_no"),
        row.get("bill_no"), row.get("settle_from"), row.get("settle_to"),
        row.get("bill_receivable_amount"), row.get("row_no"), row.get("status"),
        row.get("item_code"), row.get("item_name"), row.get("period_from"),
        row.get("period_to"), row.get("finance_month"), row.get("amount"),
        row.get("checked_amount"), row.get("paid_amount"), row.get("deducted_amount"),
        row.get("balance_amount"), row.get("receivable_component"),
        "是" if row.get("is_receivable_line") else "否",
        "是" if row.get("is_sales_refund") else "否",
        row.get("adjustment_amount"), row.get("sales_reference_amount"), row.get("tax_rate"),
        row.get("no_tax_amount"), row.get("memo"), row.get("calculation_source"),
        row.get("is_advance"), row.get("area"), row.get("rental_area"),
    ]


def build_rental_receivable_expense_workbook_file(report: dict[str, Any]):
    workbook = Workbook(write_only=True)
    detail_sheet = workbook.create_sheet("费用明细")
    detail_sheet.freeze_panes = "A2"
    detail_sheet.sheet_view.showGridLines = False

    header_cells = []
    for label in HEADERS:
        cell = WriteOnlyCell(detail_sheet, label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    detail_sheet.append(header_cells)

    for row in report.get("items", []):
        output_cells = []
        for column, value in enumerate(_row_values(row), 1):
            cell = WriteOnlyCell(detail_sheet, _safe_text(value))
            if column in MONEY_COLUMNS:
                cell.number_format = MONEY_FORMAT
            elif column == 31:
                cell.number_format = RATE_FORMAT
            cell.alignment = Alignment(vertical="center")
            output_cells.append(cell)
        detail_sheet.append(output_cells)

    last_row = max(1, int(report.get("detail_count") or 0) + 1)
    detail_sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"
    widths = {
        1: 12, 2: 20, 3: 14, 4: 20, 5: 16, 6: 22, 7: 14, 8: 28,
        9: 16, 10: 24, 11: 14, 12: 14, 13: 17, 14: 8, 15: 10, 16: 14,
        17: 24, 18: 14, 19: 14, 20: 12, 33: 28, 34: 18,
    }
    for column in range(1, len(HEADERS) + 1):
        detail_sheet.column_dimensions[get_column_letter(column)].width = widths.get(column, 14)

    notes_sheet = workbook.create_sheet("导出说明")
    filters = report.get("filters", {})
    notes = [
        ["租赁应收未收—费用明细导出说明"],
        ["导出时间", report.get("generated_at")],
        ["数据来源", report.get("source", {}).get("name")],
        ["数据截至", report.get("source_loaded_at")],
        ["结算截止日范围", f"{filters.get('settle_from', '')} 至 {filters.get('settle_to', '')}"],
        ["门店筛选", filters.get("mkt") or "全部（受账号权限限制）"],
        ["部门筛选", filters.get("department_code") or "全部（受账号权限限制）"],
        ["柜组编码开头", filters.get("group_prefix") or "未限制"],
        ["关键词", filters.get("keyword") or "未限制"],
        ["结算单数", report.get("bill_count", 0)],
        ["费用明细行数", report.get("detail_count", 0)],
        ["导出口径", report.get("scope_note")],
        ["应收影响", "非 00-* 行取明细余额并保留正负号；00-* 销售返款行的应收影响为 0。"],
    ]
    for row_index, values in enumerate(notes, 1):
        cells = []
        for value in values:
            cell = WriteOnlyCell(notes_sheet, _safe_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index == 1:
                cell.font = Font(bold=True, size=14, color="1F4E78")
            cells.append(cell)
        notes_sheet.append(cells)
    notes_sheet.column_dimensions["A"].width = 20
    notes_sheet.column_dimensions["B"].width = 80

    output = SpooledTemporaryFile(max_size=16 * 1024 * 1024, mode="w+b")
    workbook.save(output)
    output.seek(0)
    return output
