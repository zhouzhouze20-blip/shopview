"""Import an audited revenue month-close workbook as an immutable snapshot."""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy import bindparam, text
from sqlalchemy.engine import Engine

from services.monthly_followup_report import financial_month_period


MONEY = Decimal("0.01")


def _money(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(MONEY, rounding=ROUND_HALF_UP)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


@dataclass(frozen=True)
class MonthCloseAdjustment:
    workbook_row: int
    target_component: str
    adjustment_category: str
    source_department_code: str | None
    source_department_name: str | None
    source_subject_code: str | None
    source_subject_name: str | None
    source_business_type: str | None
    source_group_code: str | None
    source_group_name: str | None
    supplier_code: str | None
    supplier_name: str | None
    fee_type_code: str | None
    fee_type_name: str | None
    source_bill_count: int
    source_row_count: int
    raw_amount: Decimal
    accrued_tax_amount: Decimal
    adjustment_amount: Decimal
    final_amount: Decimal
    allocation_basis: str | None
    adjustment_reason: str | None
    period_month: str
    binding_status: str


@dataclass(frozen=True)
class MonthCloseWorkbook:
    path: Path
    source_snapshot_id: str
    period_month: str
    nc_amount_before_tax: Decimal
    accrued_tax_amount: Decimal
    nc_control_amount: Decimal
    raw_fee_amount: Decimal
    raw_extra_amount: Decimal
    close_adjustment_amount: Decimal
    final_fee_extra_amount: Decimal
    adjustments: tuple[MonthCloseAdjustment, ...]


def _label_amounts(sheet: Any) -> dict[str, Decimal]:
    values: dict[str, Decimal] = {}
    for row in sheet.iter_rows(values_only=True):
        label = _text(row[0] if row else None)
        amount = row[1] if len(row) > 1 else None
        if label and isinstance(amount, (int, float, Decimal)):
            values[label] = _money(amount)
    return values


def parse_month_close_workbook(path: str | Path) -> MonthCloseWorkbook:
    """Parse and fully reconcile the finance-approved workbook before any write."""
    workbook_path = Path(path).expanduser().resolve()
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    required_sheets = {"月结结论", "部门科目调平", "柜位月结调整", "数据检查"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"月结工作簿缺少工作表: {', '.join(sorted(missing))}")

    failed_checks = []
    for row in workbook["数据检查"].iter_rows(min_row=5, values_only=True):
        label = _text(row[0] if row else None)
        if not label:
            break
        status = _text(row[5] if len(row) > 5 else None)
        if not status or status.upper() != "PASS":
            failed_checks.append(f"{label}={status or '缺失'}")
    if failed_checks:
        raise ValueError("月结工作簿存在未通过检查: " + ", ".join(failed_checks))

    summary = _label_amounts(workbook["月结结论"])
    required_labels = {
        "NC6051不含计提税",
        "NC计提税净额",
        "NC6051含计提税控制数",
        "NC非富基收费",
        "月结净调整",
        "调平后收益",
    }
    missing_labels = required_labels.difference(summary)
    if missing_labels:
        raise ValueError(f"月结结论缺少金额: {', '.join(sorted(missing_labels))}")
    fee_amount_label = next(
        (
            label
            for label in ("富基6051可比联营租赁收费", "富基全部联营租赁收费")
            if label in summary
        ),
        None,
    )
    if fee_amount_label is None:
        raise ValueError("月结结论缺少金额: 富基6051可比联营租赁收费")

    department_subject_names: dict[tuple[str, str], tuple[str | None, str | None]] = {}
    for row in workbook["部门科目调平"].iter_rows(min_row=5, values_only=True):
        department_code = _text(row[0] if row else None)
        subject_code = _text(row[2] if len(row) > 2 else None)
        if department_code and subject_code:
            department_subject_names[(department_code, subject_code)] = (
                _text(row[1]),
                _text(row[3]),
            )

    adjustment_rows: list[MonthCloseAdjustment] = []
    sheet = workbook["柜位月结调整"]
    for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        category = _text(row[0] if row else None)
        if not category:
            continue
        department_code = _text(row[1])
        subject_code = _text(row[2])
        department_name, subject_name = department_subject_names.get(
            (department_code or "", subject_code or ""),
            (None, None),
        )
        is_pending = category.startswith("待人工绑定")
        target_component = (
            "EXTRA"
            if "NC非富基" in category or "后台" in category
            else "FEE"
        )
        adjustment = _money(row[13])
        period_month = _text(row[17])
        if not period_month:
            raise ValueError(f"柜位月结调整第 {row_number} 行缺少月结月份")
        adjustment_rows.append(
            MonthCloseAdjustment(
                workbook_row=row_number,
                target_component=target_component,
                adjustment_category=category,
                source_department_code=department_code,
                source_department_name=department_name,
                source_subject_code=subject_code,
                source_subject_name=subject_name,
                source_business_type=_text(row[3]),
                source_group_code=_text(row[4]),
                source_group_name=_text(row[5]),
                supplier_code=_text(row[6]),
                supplier_name=_text(row[7]),
                fee_type_code=_text(row[8]),
                fee_type_name=_text(row[9]),
                source_bill_count=int(row[10] or 0),
                source_row_count=int(row[11] or 0),
                raw_amount=_money(row[12]),
                accrued_tax_amount=(
                    _money(row[18])
                    if len(row) > 18 and row[18] is not None
                    else adjustment if target_component == "EXTRA" else Decimal("0.00")
                ),
                adjustment_amount=adjustment,
                final_amount=_money(row[14]),
                allocation_basis=_text(row[15]),
                adjustment_reason=_text(row[16]),
                period_month=period_month,
                binding_status="PENDING" if is_pending else "BOUND",
            )
        )

    if not adjustment_rows:
        raise ValueError("柜位月结调整没有可导入明细")
    months = {row.period_month for row in adjustment_rows}
    if len(months) != 1:
        raise ValueError(f"柜位月结调整包含多个期间: {', '.join(sorted(months))}")

    parsed = MonthCloseWorkbook(
        path=workbook_path,
        source_snapshot_id=(
            f"revenue-close-{next(iter(months))}-"
            f"{hashlib.sha256(workbook_path.read_bytes()).hexdigest()[:20]}"
        ),
        period_month=next(iter(months)),
        nc_amount_before_tax=summary["NC6051不含计提税"],
        accrued_tax_amount=summary["NC计提税净额"],
        nc_control_amount=summary["NC6051含计提税控制数"],
        raw_fee_amount=summary[fee_amount_label],
        raw_extra_amount=summary["NC非富基收费"],
        close_adjustment_amount=summary["月结净调整"],
        final_fee_extra_amount=summary["调平后收益"],
        adjustments=tuple(adjustment_rows),
    )
    validate_month_close(parsed)
    return parsed


def validate_month_close(month_close: MonthCloseWorkbook) -> None:
    checks = {
        "NC含税控制数": (
            month_close.nc_amount_before_tax + month_close.accrued_tax_amount,
            month_close.nc_control_amount,
        ),
        "月结调整合计": (
            sum((row.adjustment_amount for row in month_close.adjustments), Decimal("0.00")),
            month_close.close_adjustment_amount,
        ),
        "调平后收益": (
            month_close.raw_fee_amount
            + month_close.raw_extra_amount
            + month_close.close_adjustment_amount,
            month_close.final_fee_extra_amount,
        ),
        "调平后收益等于NC": (
            month_close.final_fee_extra_amount,
            month_close.nc_control_amount,
        ),
    }
    failures = []
    for label, (actual, expected) in checks.items():
        difference = (_money(actual) - _money(expected)).copy_abs()
        if difference > MONEY:
            failures.append(f"{label}: 实际 {_money(actual)}，预期 {_money(expected)}")
    if failures:
        raise ValueError("月结工作簿勾稽失败: " + "; ".join(failures))


def _unique_unit_map(rows: Iterable[Any], key_fields: tuple[str, ...]) -> dict[tuple[str, ...], int]:
    candidates: dict[tuple[str, ...], set[int]] = {}
    for row in rows:
        key = tuple(str(row[field]).strip() for field in key_fields)
        candidates.setdefault(key, set()).add(int(row["unit_id"]))
    return {key: next(iter(ids)) for key, ids in candidates.items() if len(ids) == 1}


def _repair_confirmed_fee_unit_links(
    connection: Any,
    *,
    month_close_id: int,
    store_code: str,
) -> int:
    """Attach unambiguous Fuji groups to units without changing financial values."""
    before = connection.execute(
        text(
            """
            SELECT COUNT(*) AS row_count,
                   COALESCE(SUM(adjustment_amount), 0) AS adjustment_amount
            FROM revenue_month_close_adjustments
            WHERE month_close_id = :month_close_id
            """
        ),
        {"month_close_id": month_close_id},
    ).mappings().one()
    updated = connection.execute(
        text(
            """
            WITH candidate_units AS (
              SELECT
                TRIM(fee.source_group_code) AS source_group_code,
                MIN(fee.unit_id) AS unit_id
              FROM unit_revenue_fee_detail fee
              JOIN business_units unit ON unit.id = fee.unit_id
              JOIN floors unit_floor ON unit_floor.id = unit.floor_id
              WHERE TRIM(unit_floor.store_code) = :store_code
                AND fee.unit_id IS NOT NULL
                AND NULLIF(TRIM(fee.source_group_code), '') IS NOT NULL
              GROUP BY TRIM(fee.source_group_code)
              HAVING COUNT(DISTINCT fee.unit_id) = 1
            )
            UPDATE revenue_month_close_adjustments adjustment
            SET unit_id = candidate.unit_id,
                unit_code = unit.unit_code
            FROM candidate_units candidate
            JOIN business_units unit ON unit.id = candidate.unit_id
            WHERE adjustment.month_close_id = :month_close_id
              AND adjustment.target_component = 'FEE'
              AND adjustment.binding_status = 'BOUND'
              AND TRIM(adjustment.source_group_code) = candidate.source_group_code
              AND adjustment.unit_id IS DISTINCT FROM candidate.unit_id
            """
        ),
        {"month_close_id": month_close_id, "store_code": store_code.strip()},
    )
    after = connection.execute(
        text(
            """
            SELECT COUNT(*) AS row_count,
                   COALESCE(SUM(adjustment_amount), 0) AS adjustment_amount
            FROM revenue_month_close_adjustments
            WHERE month_close_id = :month_close_id
            """
        ),
        {"month_close_id": month_close_id},
    ).mappings().one()
    if int(before["row_count"]) != int(after["row_count"]):
        raise RuntimeError("补充柜位关联时明细行数发生变化")
    if _money(before["adjustment_amount"]) != _money(after["adjustment_amount"]):
        raise RuntimeError("补充柜位关联时月结金额发生变化")
    return int(updated.rowcount or 0)


def import_month_close(
    engine: Engine,
    month_close: MonthCloseWorkbook,
    *,
    store_code: str,
    period_start_date: date,
    period_end_date: date,
    supersede_existing: bool = False,
) -> dict[str, Any]:
    """Insert one confirmed snapshot and all adjustments atomically."""
    if period_end_date < period_start_date:
        raise ValueError("月结结束日期不能早于开始日期")
    period_year, period_number = (int(part) for part in month_close.period_month.split("-"))
    expected_start_date, expected_end_date = financial_month_period(period_year, period_number)
    if (period_start_date, period_end_date) != (expected_start_date, expected_end_date):
        raise ValueError(
            f"{month_close.period_month}财务月应为 "
            f"{expected_start_date.isoformat()} 至 {expected_end_date.isoformat()}"
        )

    with engine.begin() as connection:
        store = connection.execute(
            text("SELECT store_id, store_name FROM stores WHERE TRIM(store_code) = :store_code"),
            {"store_code": store_code.strip()},
        ).mappings().one_or_none()
        if not store:
            raise ValueError(f"门店编码不存在: {store_code}")
        store_id = int(store["store_id"])
        store_name = str(store["store_name"] or store_code).strip()

        existing = connection.execute(
            text(
                """
                SELECT id, source_snapshot_id, version
                FROM revenue_month_closes
                WHERE store_id = :store_id
                  AND period_month = :period_month
                  AND status = 'CONFIRMED'
                FOR UPDATE
                """
            ),
            {
                "store_id": store_id,
                "period_month": month_close.period_month,
            },
        ).mappings().one_or_none()
        if existing:
            if existing["source_snapshot_id"] == month_close.source_snapshot_id:
                repaired_fee_rows = _repair_confirmed_fee_unit_links(
                    connection,
                    month_close_id=int(existing["id"]),
                    store_code=store_code,
                )
                count = connection.execute(
                    text(
                        "SELECT COUNT(*) FROM revenue_month_close_adjustments "
                        "WHERE month_close_id = :month_close_id"
                    ),
                    {"month_close_id": existing["id"]},
                ).scalar_one()
                return {
                    "status": "already_imported",
                    "month_close_id": int(existing["id"]),
                    "adjustment_rows": int(count),
                    "store_id": store_id,
                    "repaired_fee_rows": repaired_fee_rows,
                }
            if not supersede_existing:
                raise ValueError("该门店和财务期间已有其他已确认月结；如需保留历史并更正，请使用 supersede")
            connection.execute(
                text(
                    """
                    UPDATE revenue_month_closes
                    SET status = 'REOPENED',
                        note = CONCAT(COALESCE(note, ''), '；由新版本月结快照替代。')
                    WHERE id = :month_close_id
                    """
                ),
                {"month_close_id": existing["id"]},
            )

        next_version = int(
            connection.execute(
                text(
                    """
                    SELECT COALESCE(MAX(version), 0) + 1
                    FROM revenue_month_closes
                    WHERE store_id = :store_id
                      AND period_month = :period_month
                    """
                ),
                {"store_id": store_id, "period_month": month_close.period_month},
            ).scalar_one()
        )

        group_codes = sorted(
            {
                row.source_group_code
                for row in month_close.adjustments
                if row.binding_status == "BOUND" and row.source_group_code
            }
        )
        fee_unit_map: dict[tuple[str, ...], int] = {}
        if group_codes:
            fee_rows = connection.execute(
                text(
                    """
                    SELECT TRIM(fee.source_group_code) AS source_group_code, fee.unit_id
                    FROM unit_revenue_fee_detail fee
                    JOIN business_units unit ON unit.id = fee.unit_id
                    JOIN floors unit_floor ON unit_floor.id = unit.floor_id
                    WHERE TRIM(unit_floor.store_code) = :store_code
                      AND fee.unit_id IS NOT NULL
                      AND TRIM(fee.source_group_code) IN :group_codes
                    """
                ).bindparams(bindparam("group_codes", expanding=True)),
                {"store_code": store_code.strip(), "group_codes": group_codes},
            ).mappings().all()
            fee_unit_map = _unique_unit_map(fee_rows, ("source_group_code",))

        department_codes = sorted(
            {
                row.source_department_code
                for row in month_close.adjustments
                if row.binding_status == "BOUND"
                and row.target_component == "EXTRA"
                and row.source_department_code
            }
        )
        extra_unit_map: dict[tuple[str, ...], int] = {}
        if department_codes:
            extra_rows = connection.execute(
                text(
                    """
                    SELECT TRIM(source_department_code) AS source_department_code, unit_id
                    FROM revenue_extra_receipts
                    WHERE store_id = :store_id
                      AND revenue_month = :period_month
                      AND source_type = 'NC6051'
                      AND status = 'CONFIRMED'
                      AND unit_id IS NOT NULL
                      AND TRIM(source_department_code) IN :department_codes
                    """
                ).bindparams(bindparam("department_codes", expanding=True)),
                {
                    "store_id": store_id,
                    "period_month": month_close.period_month,
                    "department_codes": department_codes,
                },
            ).mappings().all()
            extra_unit_map = _unique_unit_map(extra_rows, ("source_department_code",))

        month_close_id = connection.execute(
            text(
                """
                INSERT INTO revenue_month_closes (
                  store_id, period_month, period_start_date, period_end_date,
                  version, status, nc_amount_before_tax, accrued_tax_amount,
                  nc_control_amount, raw_fee_amount, raw_extra_amount,
                  close_adjustment_amount, final_fee_extra_amount,
                  source_snapshot_id, source_file_name, note, confirmed_at, raw_payload
                ) VALUES (
                  :store_id, :period_month, :period_start_date, :period_end_date,
                  :version, 'CONFIRMED', :nc_amount_before_tax, :accrued_tax_amount,
                  :nc_control_amount, :raw_fee_amount, :raw_extra_amount,
                  :close_adjustment_amount, :final_fee_extra_amount,
                  :source_snapshot_id, :source_file_name, :note, NOW(), CAST(:raw_payload AS JSONB)
                )
                RETURNING id
                """
            ),
            {
                "store_id": store_id,
                "period_month": month_close.period_month,
                "period_start_date": period_start_date,
                "period_end_date": period_end_date,
                "version": next_version,
                "nc_amount_before_tax": month_close.nc_amount_before_tax,
                "accrued_tax_amount": month_close.accrued_tax_amount,
                "nc_control_amount": month_close.nc_control_amount,
                "raw_fee_amount": month_close.raw_fee_amount,
                "raw_extra_amount": month_close.raw_extra_amount,
                "close_adjustment_amount": month_close.close_adjustment_amount,
                "final_fee_extra_amount": month_close.final_fee_extra_amount,
                "source_snapshot_id": month_close.source_snapshot_id,
                "source_file_name": month_close.path.name,
                "note": (
                    f"{store_name}{period_year}财务{period_number}月"
                    f"（{period_start_date:%m-%d}至{period_end_date:%m-%d}）"
                    "NC6051含计提税月结；差异不分摊，待人工绑定。"
                ),
                "raw_payload": json.dumps(
                    {
                        "source_path": str(month_close.path),
                        "adjustment_rows": len(month_close.adjustments),
                        "import_rule": "finance_workbook_manual_binding_v2",
                    },
                    ensure_ascii=False,
                ),
            },
        ).scalar_one()

        parameters = []
        mapped_fee_rows = 0
        mapped_extra_rows = 0
        for row in month_close.adjustments:
            if row.binding_status == "PENDING":
                unit_id = None
            elif row.target_component == "EXTRA":
                unit_id = extra_unit_map.get((row.source_department_code or "",))
                mapped_extra_rows += int(unit_id is not None)
            else:
                unit_id = fee_unit_map.get((row.source_group_code or "",))
                mapped_fee_rows += int(unit_id is not None)
            source_row_key = (
                f"{month_close.period_month}:{row.workbook_row}:"
                f"{row.adjustment_category}:{row.source_department_code or '-'}:"
                f"{row.source_subject_code or '-'}:{row.source_group_code or '-'}"
            )[:160]
            parameters.append(
                {
                    "month_close_id": month_close_id,
                    "target_component": row.target_component,
                    "adjustment_category": row.adjustment_category,
                    "unit_id": unit_id,
                    "unit_code": None,
                    "source_group_code": row.source_group_code,
                    "source_group_name": row.source_group_name,
                    "source_department_code": row.source_department_code,
                    "source_department_name": row.source_department_name,
                    "source_subject_code": row.source_subject_code,
                    "source_subject_name": row.source_subject_name,
                    "source_business_type": row.source_business_type,
                    "supplier_code": row.supplier_code,
                    "supplier_name": row.supplier_name,
                    "fee_type_code": row.fee_type_code,
                    "fee_type_name": row.fee_type_name,
                    "raw_amount": row.raw_amount,
                    "accrued_tax_amount": row.accrued_tax_amount,
                    "adjustment_amount": row.adjustment_amount,
                    "final_amount": row.final_amount,
                    "allocation_basis": row.allocation_basis,
                    "adjustment_reason": row.adjustment_reason,
                    "source_row_key": source_row_key,
                    "binding_status": row.binding_status,
                    "raw_payload": json.dumps(
                        {
                            "workbook_row": row.workbook_row,
                            "period_month": row.period_month,
                            "source_bill_count": row.source_bill_count,
                            "source_row_count": row.source_row_count,
                        },
                        ensure_ascii=False,
                    ),
                }
            )

        connection.execute(
            text(
                """
                INSERT INTO revenue_month_close_adjustments (
                  month_close_id, target_component, adjustment_category,
                  unit_id, unit_code, source_group_code, source_group_name,
                  source_department_code, source_department_name,
                  source_subject_code, source_subject_name, source_business_type,
                  supplier_code, supplier_name, fee_type_code, fee_type_name,
                  raw_amount, accrued_tax_amount, adjustment_amount, final_amount,
                  allocation_basis, adjustment_reason, source_row_key,
                  binding_status, raw_payload
                ) VALUES (
                  :month_close_id, :target_component, :adjustment_category,
                  :unit_id, :unit_code, :source_group_code, :source_group_name,
                  :source_department_code, :source_department_name,
                  :source_subject_code, :source_subject_name, :source_business_type,
                  :supplier_code, :supplier_name, :fee_type_code, :fee_type_name,
                  :raw_amount, :accrued_tax_amount, :adjustment_amount, :final_amount,
                  :allocation_basis, :adjustment_reason, :source_row_key,
                  :binding_status, CAST(:raw_payload AS JSONB)
                )
                """
            ),
            parameters,
        )

        persisted = connection.execute(
            text(
                """
                SELECT COUNT(*) AS row_count,
                       COALESCE(SUM(adjustment_amount), 0) AS adjustment_amount
                FROM revenue_month_close_adjustments
                WHERE month_close_id = :month_close_id
                """
            ),
            {"month_close_id": month_close_id},
        ).mappings().one()
        if int(persisted["row_count"]) != len(month_close.adjustments):
            raise RuntimeError("月结调整写入行数校验失败")
        if _money(persisted["adjustment_amount"]) != month_close.close_adjustment_amount:
            raise RuntimeError("月结调整写入金额校验失败")

        return {
            "status": "imported",
            "month_close_id": int(month_close_id),
            "store_id": store_id,
            "version": next_version,
            "superseded_month_close_id": int(existing["id"]) if existing else None,
            "adjustment_rows": len(month_close.adjustments),
            "mapped_fee_rows": mapped_fee_rows,
            "mapped_extra_rows": mapped_extra_rows,
            "nc_control_amount": str(month_close.nc_control_amount),
            "final_fee_extra_amount": str(month_close.final_fee_extra_amount),
        }
