"""New Century payment report: intentionally preserves the supplied Oracle SQL grain."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import text

from .od0002_report import TrustedScopeSql, _trusted_scope_value

PAY_CODES = ('0328', '0329', '0330', '0331', '0332', '0333')
NOTE = ('按原SQL口径：每条支付记录关联整票毛利和销售收入，不去重、不分摊；'
        '柜组取小票商品明细MAX(gz)。商品明细自2026-06-01起，销售流水自2026-01-01起；'
        '支付金额仅djlb=1取正，其余取负；毛利、销售收入保留源值符号。')
COLUMNS = [('rq', '日期'), ('mkt', '门店'), ('bmname', '部门'), ('gz', '柜组编码'),
           ('gzname', '柜组名称'), ('paycode', '支付编码'), ('pname', '支付名称'),
           ('money', '支付金额'), ('bl', '折算收入比例'), ('ml', '毛利'), ('xssr', '销售收入')]


def build_query(*, start_date: date, end_date: date, department_id=None, group_code=None,
                pay_codes=PAY_CODES, scope_filter_sql=TrustedScopeSql(' AND 1=0'), scope_params=None):
    if end_date < start_date or (end_date - start_date).days > 365:
        raise ValueError('日期范围须为1至366天')
    if not pay_codes or any(code not in PAY_CODES for code in pay_codes):
        raise ValueError('支付方式仅支持0328至0333')
    params = dict(scope_params or {})
    params.update(start_date=start_date, end_exclusive=end_date + timedelta(days=1), pay_codes=list(pay_codes))
    filters = ''
    for key, expression, value in [('department_id', 'mf.mfpcode', department_id), ('group_code', 'r.gz', group_code)]:
        if value and value.strip():
            params[key] = value.strip()
            filters += f' AND {expression} = :{key}'
    sql = f"""
WITH selected_pay AS MATERIALIZED (
  SELECT aa.* FROM salepay aa
  WHERE aa.mkt = '603' AND aa.paycode = ANY(:pay_codes)
    AND aa.rqsj >= :start_date AND aa.rqsj < :end_exclusive
), bill_groups AS (
  SELECT g.billno, MAX(g.gz) AS gz FROM salegoods g
  WHERE g.rqsj >= DATE '2026-06-01'
    AND EXISTS (SELECT 1 FROM selected_pay p WHERE p.billno = g.billno)
  GROUP BY g.billno
), bill_sales AS (
  SELECT s.sglbillno, SUM(s.sgln2) AS ml, SUM(s.sglxssr) AS xssr
  FROM salegoodslist s WHERE s.sgldate >= DATE '2026-01-01'
    AND EXISTS (SELECT 1 FROM selected_pay p WHERE p.billno = s.sglbillno)
  GROUP BY s.sglbillno
), report AS (
  SELECT aa.rqsj::date AS rq, aa.mkt, dd.gz, aa.paycode, aa.payname,
    SUM(aa.je * CASE WHEN cc.djlb = '1' THEN 1 ELSE -1 END) AS money,
    SUM(ee.ml) AS ml, SUM(ee.xssr) AS xssr
  FROM selected_pay aa
  JOIN salehead cc ON aa.billno = cc.billno
  JOIN bill_groups dd ON aa.billno = dd.billno
  JOIN bill_sales ee ON aa.billno = ee.sglbillno
  GROUP BY aa.rqsj::date, aa.mkt, aa.paycode, aa.payname, dd.gz
)
SELECT r.*, mf.mfpcode AS bm,
  (SELECT f.mfcname FROM manaframe f WHERE f.mfcode = mf.mfpcode) AS bmname,
  mf.mfcname AS gzname,
  (SELECT pm.pmname FROM paymode pm WHERE pm.pmcode = r.paycode) AS pname,
  (SELECT ROUND(pm.pmrevrate, 4) FROM paymode pm WHERE pm.pmcode = r.paycode) AS bl
FROM report r
LEFT JOIN manaframe mf ON mf.mfcode = r.gz
WHERE 1=1 {filters}
AND EXISTS (
  SELECT 1 FROM (SELECT 1) scope_anchor
  LEFT JOIN stores st ON st.store_code = r.mkt
  LEFT JOIN manaframe dept ON dept.mfcode = mf.mfpcode
  LEFT JOIN area_category ac ON UPPER(TRIM(ac.category_code)) = UPPER(TRIM(mf.mfchr1))
  WHERE 1=1 {_trusted_scope_value(scope_filter_sql)}
)
ORDER BY r.rq, r.mkt, r.gz, r.paycode, r.payname
"""
    return sql, params


def load_report(db, **kwargs):
    sql, params = build_query(**kwargs)
    db.execute(text("SET LOCAL statement_timeout = '120s'"))
    rows = [dict(row) for row in db.execute(text(sql), params).mappings()]
    totals = {key: sum((row[key] or Decimal(0) for row in rows), Decimal(0))
              for key in ('money', 'ml', 'xssr')}
    return dict(rows=rows, totals=totals, note=NOTE,
                start_date=kwargs['start_date'], end_date=kwargs['end_date'], store_code='603')


def build_workbook(report):
    wb = Workbook()
    ws = wb.active
    ws.title = '新世纪支付方式销售毛利'
    ws.append(['新世纪支付方式销售毛利报表'])
    ws.append([f"查询期间：{report['start_date']} 至 {report['end_date']}（含首尾）"])
    ws.append([NOTE])
    for row_number in (1, 2, 3):
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=11)
        ws.cell(row_number, 1).alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 42
    ws['A1'].font = Font(size=16, bold=True)
    ws.append([label for _, label in COLUMNS])
    for row in report['rows']:
        ws.append([row.get(key) for key, _ in COLUMNS])
    ws.append(['合计', None, None, None, None, None, None,
               report['totals']['money'], None, report['totals']['ml'], report['totals']['xssr']])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if cell.row >= 5:
                if cell.column == 8: cell.number_format = '#,##0.00'
                if cell.column in (10, 11): cell.number_format = '#,##0.00##'
                if cell.column == 9: cell.number_format = '0.0000'
                if cell.column == 1 and isinstance(cell.value, date): cell.number_format = 'yyyy-mm-dd'
    for cell in ws[4]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1D4ED8')
    for index, width in enumerate([15, 12, 24, 18, 26, 14, 24, 18, 18, 18, 18], 1):
        ws.column_dimensions[ws.cell(4, index).column_letter].width = width
    ws.freeze_panes = 'E5'
    ws.auto_filter.ref = f'A4:K{4 + len(report["rows"])}'
    ws.print_title_rows = '1:4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
