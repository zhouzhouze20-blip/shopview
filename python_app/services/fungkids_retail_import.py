"""Validated Nebula exports, kept in a private staging database before publication.

An export is a complete snapshot of an explicitly supplied date interval. Replacing
that interval preserves legitimate identical lines without guessing a line ID.
Only business fields are persisted; customer names and phone numbers are omitted.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

CENT = Decimal("0.01")
VENUE_RATE = Decimal("0.15")
REQUIRED = ("单据编号", "渠道编号", "单据日期", "货号", "计收额", "数量",
            "计收折扣", "销售方式", "原单号")
PRICE_FIELDS = ("吊牌价", "吊牌额", "零售价", "生意额", "计收价", "计收额")


def number(value, label: str) -> Decimal:
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label}缺少有效数值") from None
    if not result.is_finite():
        raise ValueError(f"{label}必须为有限数值")
    return result


def money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_free_gift(record: dict) -> bool:
    """Require source gift evidence and no paid amounts before filling blanks."""
    marked = (str(record.get("销售类型") or "").strip() in ("赠品", "赠送")
              or any(marker in str(record.get("促销信息") or "")
                     for marker in ("买满赠送", "会员积分兑换货品")))
    return marked and all(
        is_blank(record.get(field)) or number(record[field], field) == 0
        for field in ("计收额", "计收价", "生意额")
    )


def is_unpriced_line(record: dict) -> bool:
    """Recognize source rows with no prices, even without a promotion marker.

    Require the complete price schema: omitted columns are not zero evidence.
    The export control total remains mandatory after normalizing these rows.
    """
    return all(field in record for field in PRICE_FIELDS) and all(
        is_blank(record[field]) or number(record[field], field) == 0
        for field in PRICE_FIELDS
    )


def settlement_rate(discount: Decimal) -> Decimal:
    if not discount.is_finite() or discount < 0 or discount > 1:
        raise ValueError("计收折扣必须在0至1之间")
    return Decimal("0.50") if discount >= Decimal("0.85") else (
        Decimal("0.52") if discount >= Decimal("0.70") else Decimal("0.55"))


@dataclass(frozen=True)
class RetailLine:
    source_row: int
    bill: str
    channel: str
    business_date: str
    product: str
    quantity: Decimal
    sales: Decimal
    discount: Decimal
    sale_type: str
    original_bill: str
    supplier_rate: Decimal
    supplier_cost: Decimal
    venue_fee: Decimal
    expected_receipt: Decimal
    gross_profit: Decimal


@dataclass(frozen=True)
class RetailExport:
    sha256: str
    channel: str
    start: str
    end: str
    lines: tuple[RetailLine, ...]

    def summary(self) -> dict:
        return {
            "start_date": self.start, "end_date": self.end,
            "channel": self.channel, "row_count": len(self.lines),
            "return_rows": sum(line.sale_type == "退货" for line in self.lines),
            **{key: str(sum((getattr(line, key) for line in self.lines), Decimal(0)))
               for key in ("sales", "venue_fee", "expected_receipt", "supplier_cost", "gross_profit")},
        }


def parse_export(path: str | Path, *, start: date, end: date,
                 channel: str = "PBCZ6001") -> RetailExport:
    if start > end:
        raise ValueError("开始日期不能晚于结束日期")
    path = Path(path)
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if "店铺零售报表" not in book.sheetnames:
            raise ValueError("缺少店铺零售报表工作表")
        rows = book["店铺零售报表"].iter_rows(values_only=True)
        header = next(rows, ())
        if any(header.count(name) != 1 for name in REQUIRED):
            raise ValueError("必需字段缺失或重复")
        indexes = {name: header.index(name) for name in REQUIRED}
        indexes.update({name: header.index(name) for name in
                        ("销售类型", "促销信息", *PRICE_FIELDS) if name in header})
        lines = []
        control = None
        for row_no, row in enumerate(rows, 2):
            if not any(value is not None for value in row):
                continue
            if row[0] == "合计":
                if control is not None:
                    raise ValueError("出现多个合计行")
                control = number(row[indexes["计收额"]], "合计计收额")
                continue
            if control is not None:
                raise ValueError("合计行后出现明细")
            record = {name: row[index] for name, index in indexes.items()}
            if str(record["渠道编号"]) != channel:
                raise ValueError(f"第{row_no}行渠道不属于指定店铺")
            raw_date = record["单据日期"]
            day = raw_date.date() if isinstance(raw_date, datetime) else (
                raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)))
            if not start <= day <= end:
                raise ValueError(f"第{row_no}行超出声明的查询日期范围")
            if not record["单据编号"] or not record["货号"]:
                raise ValueError(f"第{row_no}行缺少单据编号或货号")
            zero_amount_line = is_free_gift(record) or is_unpriced_line(record)
            sales = (Decimal(0) if zero_amount_line and is_blank(record["计收额"])
                     else number(record["计收额"], f"第{row_no}行计收额"))
            if sales != money(sales):
                raise ValueError(f"第{row_no}行计收额超过两位小数")
            quantity = number(record["数量"], f"第{row_no}行数量")
            sale_type = str(record["销售方式"])
            if sale_type not in ("销售", "换货", "退货"):
                raise ValueError(f"第{row_no}行销售方式未支持")
            if (sale_type in ("销售", "换货") and (sales < 0 or quantity < 0)) or (
                    sale_type == "退货" and (sales > 0 or quantity > 0)):
                raise ValueError(f"第{row_no}行金额或数量符号与销售方式不符")
            discount = (Decimal(0) if zero_amount_line and is_blank(record["计收折扣"])
                        else number(record["计收折扣"], f"第{row_no}行计收折扣"))
            rate = settlement_rate(discount)
            cost, fee = money(sales * rate), money(sales * VENUE_RATE)
            lines.append(RetailLine(row_no, str(record["单据编号"]), channel,
                day.isoformat(), str(record["货号"]), quantity, sales, discount,
                sale_type, str(record["原单号"] or ""), rate, cost, fee,
                sales - fee, sales - fee - cost))
        if control is None or sum((line.sales for line in lines), Decimal(0)) != control:
            raise ValueError("明细计收额与合计不一致，或缺少合计行")
        return RetailExport(digest, channel, start.isoformat(), end.isoformat(), tuple(lines))
    finally:
        book.close()


def import_snapshot(database: str | Path, export: RetailExport, *, allow_empty=False) -> dict:
    """Local staging only. No changes to ShopView production sales or permissions."""
    if not export.lines and not allow_empty:
        raise ValueError("空报表不能自动覆盖既有数据；需显式确认空区间")
    database = Path(database)
    database.parent.mkdir(parents=True, exist_ok=True)
    # Create with private permissions before SQLite opens the file.
    import os
    fd = os.open(database, os.O_CREAT | os.O_RDWR, 0o600)
    os.close(fd)
    with sqlite3.connect(database, timeout=30) as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS fungkids_batches (
            id INTEGER PRIMARY KEY, sha256 TEXT NOT NULL, channel TEXT NOT NULL,
            start_date TEXT NOT NULL, end_date TEXT NOT NULL, imported_at TEXT NOT NULL,
            summary TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS fungkids_lines (
            batch_id INTEGER NOT NULL, source_row INTEGER NOT NULL,
            channel TEXT NOT NULL, business_date TEXT NOT NULL, payload TEXT NOT NULL,
            PRIMARY KEY(batch_id, source_row));
        CREATE INDEX IF NOT EXISTS fungkids_date_idx ON fungkids_lines(channel,business_date);
        """)
        db.execute("BEGIN IMMEDIATE")
        batch = db.execute("""INSERT INTO fungkids_batches
            (sha256,channel,start_date,end_date,imported_at,summary) VALUES (?,?,?,?,?,?)""",
            (export.sha256, export.channel, export.start, export.end,
             datetime.now(timezone.utc).isoformat(), json.dumps(export.summary())))
        db.execute("DELETE FROM fungkids_lines WHERE channel=? AND business_date BETWEEN ? AND ?",
                   (export.channel, export.start, export.end))
        db.executemany("INSERT INTO fungkids_lines VALUES (?,?,?,?,?)", [
            (batch.lastrowid, line.source_row, line.channel, line.business_date,
             json.dumps(asdict(line), default=str, ensure_ascii=False)) for line in export.lines])
    return export.summary()
