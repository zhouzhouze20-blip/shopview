"""Excel export for the star-diamond analysis, using the existing server workbook stack."""
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MONEY = '#,##0.00;[Red]-#,##0.00'
TEXT = '@'
COUNT = '#,##0.##'


def build_star_diamond_workbook(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    period = f"常州购物中心 | {report['start_date']} 至 {report['end_date']} | 金额单位：元"

    def sheet(name, columns, rows, scope):
        ws = wb.create_sheet(name)
        width = len(columns)
        for line, content in enumerate((name, period, scope), 1):
            ws.merge_cells(start_row=line, start_column=1, end_row=line, end_column=width)
            cell = ws.cell(line, 1, content)
            cell.data_type = 's'
            cell.font = Font(name='微软雅黑', size=16 if line == 1 else 10, bold=line == 1, color='17365D')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.row_dimensions[line].height = 28 if line == 1 else 32
        for col, (key, label, fmt, col_width) in enumerate(columns, 1):
            cell = ws.cell(5, col, label)
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[get_column_letter(col)].width = col_width
        for row_index, record in enumerate(rows, 6):
            for col, (key, label, fmt, col_width) in enumerate(columns, 1):
                value = record.get(key)
                if fmt == TEXT and value is not None:
                    value = value.isoformat(sep=' ') if hasattr(value, 'hour') else (value.isoformat() if hasattr(value, 'isoformat') else str(value))
                cell = ws.cell(row_index, col, value)
                if isinstance(value, str):
                    cell.data_type = 's'  # Identifiers and user data must never become formulas.
                cell.number_format = fmt
                cell.font = Font(name='微软雅黑', size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            ws.row_dimensions[row_index].height = 32 if name != 'AI整体方案' else 60
        ws.freeze_panes = 'A6'
        ws.auto_filter.ref = f'A5:{get_column_letter(width)}{max(5, ws.max_row)}'
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = '1:5'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws

    s = report['overview']['summary']
    metric_specs = [('star_member_count','星钻会员',COUNT),('active_member_count','消费会员',COUNT),('sales_amount','期间销售额',MONEY),('ticket_count','小票数',COUNT),('avg_ticket_amount','客单价',MONEY),('avg_member_amount','人均消费',MONEY),('repeat_member_count','复购会员',COUNT),('high_value_member_count','高价值维护会员',COUNT),('silent_member_count','待唤醒会员',COUNT),('net_profit','净毛利',MONEY)]
    metrics = [{'metric':label,'value':s.get(key,0)} for key,label,fmt in metric_specs]
    for label, numerator, denominator in [('消费覆盖率','active_member_count','star_member_count'),('复购率','repeat_member_count','active_member_count')]:
        metrics.append({'metric':label,'value':s.get(numerator,0)/s[denominator] if s.get(denominator) else 0})
    ws = sheet('经营汇总',[('metric','指标',TEXT,32),('value','数值',MONEY,32)],metrics,'全部星钻会员；仅按日期筛选，不受明细筛选影响。')
    for row, (_, _, fmt) in enumerate(metric_specs,6):
        ws.cell(row,2).number_format = fmt
    for row in (16,17):
        ws.cell(row,2).number_format = '0.0%'
    sheet('服务分层',[('segment','分层',TEXT,20),('member_count','会员数',COUNT,14),('sales_amount','销售额',MONEY,22),('service_action','建议服务动作',TEXT,75)],report['overview']['service_segments'],'全部星钻会员；按日期统计。')
    sheet('偏好品类Top10',[('category_display','品类',TEXT,36),('member_count','消费会员',COUNT,16),('ticket_count','小票数',COUNT,16),('sales_amount','销售额',MONEY,22)],report['overview']['top_categories'],'全部星钻会员；销售额排名前10的品类。')
    sheet('会员清单',[(k,l,f,w) for k,l,f,w in [
        ('member_no','会员号',TEXT,22),('customer_name','姓名',TEXT,16),('telephone','手机号',TEXT,20),('customer_level','等级',TEXT,14),('admission_date','入会日期',TEXT,22),('ticket_count','小票数',COUNT,14),('sales_amount','销售额',MONEY,22),('avg_ticket_amount','客单价',MONEY,20),('last_sale_time','最近消费',TEXT,24),('categories','消费品类',TEXT,50),('brands','消费品牌',TEXT,50),('service_segment','服务分层',TEXT,18)
    ]],report['members'],f"搜索：{report.get('keyword') or '全部'}；分层：{report.get('segment') or '全部'}；共 {len(report['members'])} 行，完整导出。")
    sheet('购物轨迹',[(k,l,f,w) for k,l,f,w in [
        ('sale_time','销售时间',TEXT,24),('billno','小票号',TEXT,26),('member_no','会员号',TEXT,22),('customer_name','姓名',TEXT,16),('sku_count','SKU数',COUNT,12),('quantity','件数',COUNT,12),('sales_amount','销售额',MONEY,20),('net_profit','净毛利',MONEY,20),('departments','部门',TEXT,36),('groups','柜组',TEXT,44),('categories','品类',TEXT,44),('brands','品牌',TEXT,44)
    ]],report['trails'],f"指定会员：{report.get('member_no') or '全部'}；不受会员清单搜索和分层影响；共 {len(report['trails'])} 行，完整导出。")
    if report.get('ai_report'):
        sheet('AI整体方案',[('line','段落',COUNT,10),('content','方案内容',TEXT,110)], [{'line':i,'content':line} for i,line in enumerate(report['ai_report'].splitlines(),1)],'保存页面已生成、日期一致的整体 AI 方案，导出时不重新生成。')
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
