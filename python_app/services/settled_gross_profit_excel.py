from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = (
    "门店",
    "部门",
    "柜组",
    "供应商",
    "楼层",
    "营业面积",
    "经营方式",
    "区域",
    "合同号",
    "销售数量",
    "销售收入",
    "不含税销售",
    "前台毛利",
    "不含税毛利调整",
    "保底调整",
    "含税销售保底成本",
    "原扣率毛利",
    "家电返利",
    "合同到期日期",
    "合同毛利",
)

BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
WHITE = "FFFFFF"
AMBER = "FFF2CC"
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _code_name(code: Any, name: Any, *, brackets: bool = False) -> str:
    code_text = str(code or "").strip()
    name_text = str(name or "").strip()
    if brackets and code_text:
        return f"[{code_text}]{name_text}"
    return " ".join(value for value in (code_text, name_text) if value)


def _date_text(value: Any) -> str:
    if value is None:
        return ""
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _row_values(row: dict[str, Any]) -> list[Any]:
    return [
        _code_name(row.get("store_code"), row.get("store_name")),
        _code_name(row.get("department_code"), row.get("department_name")),
        _code_name(row.get("group_code"), row.get("group_name")),
        _code_name(row.get("supplier_code"), row.get("supplier_name"), brackets=True),
        row.get("floor_name"),
        row.get("business_area"),
        row.get("operation_mode_name"),
        _code_name(row.get("area_code"), row.get("area_name")),
        row.get("contract_code"),
        row.get("sales_qty"),
        row.get("sales_revenue"),
        row.get("tax_excluded_sales"),
        row.get("front_profit"),
        row.get("tax_excluded_profit_adjustment"),
        row.get("floor_adjustment"),
        row.get("sales_floor_cost"),
        row.get("original_rate_profit"),
        row.get("appliance_rebate"),
        row.get("contract_end_date"),
        row.get("contract_profit"),
    ]


def build_settled_gross_profit_workbook(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结算后销售毛利排行"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:T1")
    sheet["A1"] = "结算后销售毛利排行表"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    dates = report["dates"]
    sheet.merge_cells("A2:T2")
    sheet["A2"] = (
        f"统计期间：{_date_text(dates['start_date'])} 至 "
        f"{_date_text(dates['end_date'])}（含首尾日）"
    )
    sheet.merge_cells("A3:T3")
    sheet["A3"] = report.get("scope_description", "当前用户权限范围")

    quality = report.get("quality", {})
    sheet.merge_cells("A4:T4")
    sheet["A4"] = (
        f"数据质量：{quality.get('row_count', 0)} 行；"
        f"未匹配合同 {quality.get('unresolved_contract_group_count', 0)} 行；"
        f"未匹配供应商名称 {quality.get('missing_supplier_name_count', 0)} 行；"
        f"未匹配区域名称 {quality.get('missing_area_name_count', 0)} 行"
    )
    if quality.get("unresolved_contract_group_count", 0):
        sheet["A4"].fill = PatternFill("solid", fgColor=AMBER)

    for column, label in enumerate(HEADERS, 1):
        cell = sheet.cell(5, column, label)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    row_number = 6
    for row in report.get("rows", []):
        for column, value in enumerate(_row_values(row), 1):
            cell = sheet.cell(row_number, column, _safe_text(value))
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="right" if 6 <= column <= 18 or column == 20 else "left",
                vertical="center",
            )
            if column in (6, 9, 19):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if column == 6:
                cell.number_format = "#,##0.00"
            elif 10 <= column <= 18 or column == 20:
                cell.number_format = "#,##0.00"
        row_number += 1

    totals = report.get("totals", {})
    total_values = ["合计"] + [None] * 8 + [
        totals.get("sales_qty"),
        totals.get("sales_revenue"),
        totals.get("tax_excluded_sales"),
        totals.get("front_profit"),
        totals.get("tax_excluded_profit_adjustment"),
        totals.get("floor_adjustment"),
        totals.get("sales_floor_cost"),
        totals.get("original_rate_profit"),
        totals.get("appliance_rebate"),
        None,
        None,
    ]
    for column, value in enumerate(total_values, 1):
        cell = sheet.cell(row_number, column, value)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True)
        cell.border = BORDER
        if 10 <= column <= 18:
            cell.number_format = "#,##0.00"

    widths = (24, 28, 30, 32, 10, 12, 12, 18, 15) + (14,) * 9 + (14, 14)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[5].height = 36
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:T{max(5, row_number - 1)}"

    notes = workbook.create_sheet("口径说明")
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 110
    note_rows = [
        ("统计期间", "按 salegoodslist.sglhsrq，开始日和结束日均包含。"),
        ("销售收入", "sum(sglxssr + sglpfsr)。"),
        ("不含税销售", "sum((sglxssr + sglpfsr) / (1 + sglxstax))，汇总后保留 2 位。"),
        ("前台毛利", "sum(sgln2 / (1 + sglxstax))，汇总后保留 2 位。"),
        ("合同号", "由 contmain + contmanaframe 按门店、柜组、供应商、经营方式和销售日解析。"),
        ("调整项目", "来自 supsetcharge，按富基收费项目代码和对应税率还原。"),
        ("销售保底", "来自 salecostday + contbd，沿用原报表有效期判断。"),
        ("数据边界", "本文件使用柜位系统 PostgreSQL 现有表；源 Oracle 与本地同步差异需单独验数。"),
    ]
    notes.append(["项目", "说明"])
    for item in note_rows:
        notes.append(item)
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
    for row in notes.iter_rows():
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes.freeze_panes = "A2"

    output = SpooledTemporaryFile(max_size=5 * 1024 * 1024, mode="w+b")
    workbook.save(output)
    output.seek(0)
    payload = output.read()
    output.close()
    return payload


def build_settled_gross_profit_workbook_file(report: dict[str, Any]):
    output = SpooledTemporaryFile(max_size=5 * 1024 * 1024, mode="w+b")
    output.write(build_settled_gross_profit_workbook(report))
    output.seek(0)
    return output
