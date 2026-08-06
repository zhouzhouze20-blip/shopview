from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BLUE = "3877A6"
LIGHT_BLUE = "DDEBF7"
TOTAL_BLUE = "D9EAF7"
WHITE = "FFFFFF"
RED = "FF0000"
THIN_BLUE = Side(style="thin", color=BLUE)
GRID_BORDER = Border(left=THIN_BLUE, right=THIN_BLUE, top=THIN_BLUE, bottom=THIN_BLUE)


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _wan(value: Any) -> float:
    return float(value or 0) / 10_000


def _dimension_label(row: dict[str, Any], dimension: str) -> str:
    return str(row.get("store_name") if dimension == "store" else row.get("department_name") or "")


def _write_sheet(sheet, report: dict[str, Any], *, dimension: str) -> None:
    rows = report["store_rows"] if dimension == "store" else report["department_rows"]
    year = int(report["financial_year"])
    prior_year = int(report["prior_year"])
    periods = list(report["periods"])
    dimension_name = "门店" if dimension == "store" else "部门"
    last_column = 4 + len(periods) * 2 + 4

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "E5"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = f"本期年份：{year}年     (单位：万元)"
    sheet.merge_cells("D1:F1")
    sheet["D1"] = f"同期年份：{prior_year}年     (单位：万元)"
    for cell in (sheet["A1"], sheet["D1"]):
        cell.font = Font(name="宋体", size=12, bold=True, color="333333")

    for column, label in ((1, dimension_name), (2, "类别"), (3, "费用")):
        sheet.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
        sheet.cell(3, column, label)

    column = 5
    for period in periods:
        sheet.merge_cells(start_row=3, start_column=column, end_row=3, end_column=column + 1)
        sheet.cell(3, column, f"{period:02d}")
        sheet.cell(4, column, "本期")
        sheet.cell(4, column + 1, "同期")
        column += 2
    sheet.merge_cells(start_row=3, start_column=column, end_row=3, end_column=column + 1)
    sheet.cell(3, column, "合计")
    sheet.cell(4, column, "本期")
    sheet.cell(4, column + 1, "同期")
    sheet.cell(4, column + 2, "同比差异（值）")
    sheet.cell(4, column + 3, "同比比率（%）")

    for header_row in sheet.iter_rows(min_row=3, max_row=4, min_col=1, max_col=last_column):
        for cell in header_row:
            cell.font = Font(name="Times New Roman", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = GRID_BORDER
            cell.fill = PatternFill("solid", fgColor=WHITE)

    previous_group = None
    previous_category = None
    output_row = 5
    for row in rows:
        group = _dimension_label(row, dimension)
        row_type = row["row_type"]
        if row_type == "detail":
            group_value = group if group != previous_group else None
            category_value = row["category"] if (group != previous_group or row["category"] != previous_category) else None
        elif row_type == "category_total":
            group_value = None
            category_value = row["category"]
        else:
            group_value = group
            category_value = "合计"

        values: list[Any] = [group_value, category_value, row.get("fee_name") or None, None]
        for period in periods:
            comparison = row["months"][str(period)]
            values.extend((_wan(comparison["current"]), _wan(comparison["prior"])))
        values.extend(
            (
                _wan(row["total"]["current"]),
                _wan(row["total"]["prior"]),
                _wan(row["total"]["difference"]),
                row["total"]["rate"],
            )
        )
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(output_row, column_index, _safe_text(value))
            cell.border = GRID_BORDER
            cell.font = Font(
                name="Times New Roman",
                size=10,
                bold=row_type != "detail",
            )
            cell.alignment = Alignment(
                horizontal="right" if column_index >= 5 else "left",
                vertical="center",
            )
            if column_index >= 5:
                if column_index == last_column:
                    cell.number_format = "0.00%"
                else:
                    cell.number_format = "0.00"
            if column_index in (last_column - 1, last_column) and isinstance(value, (int, float)) and value < 0:
                cell.font = Font(name="Times New Roman", size=10, bold=row_type != "detail", color=RED)
            if row_type == "category_total":
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            elif row_type == "grand_total":
                cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
        previous_group = group
        if row_type == "detail":
            previous_category = row["category"]
        elif row_type == "grand_total":
            previous_group = group
            previous_category = None
        output_row += 1

    widths = {"A": 24 if dimension == "department" else 14, "B": 18, "C": 18, "D": 3}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    for index in range(5, last_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 13
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 8
    sheet.row_dimensions[3].height = 22
    sheet.row_dimensions[4].height = 22


def build_store_other_business_income_workbook(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    store_sheet = workbook.active
    store_sheet.title = "门店"
    _write_sheet(store_sheet, report, dimension="store")
    department_sheet = workbook.create_sheet("部门")
    _write_sheet(department_sheet, report, dimension="department")

    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    workbook.save(output)
    output.seek(0)
    payload = output.read()
    output.close()
    return payload


def build_store_other_business_income_workbook_file(report: dict[str, Any]):
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    output.write(build_store_other_business_income_workbook(report))
    output.seek(0)
    return output
