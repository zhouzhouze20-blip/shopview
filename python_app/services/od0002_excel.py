from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .od0002_report import EXCLUDED_DEPARTMENT_CODES


SHEETS = (
    ("分店", "stores"),
    ("部门", "departments"),
    ("区域", "areas"),
    ("品类", "categories"),
    ("柜组", "groups"),
    ("特卖", "special_sales"),
    ("楼层", "floors"),
)
BLUE = "4472C4"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
METRIC_GROUPS = (
    ("来客数", "笔"),
    ("客单", "元"),
    ("销售收入", "万元"),
    ("毛利额", "万元"),
    ("毛利率", "%"),
)


def _date_text(value: Any) -> str:
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _write_headers(sheet, dimension_name: str) -> None:
    sheet.merge_cells("A5:D5")
    sheet["A5"] = "维度"
    for index, (label, unit) in enumerate(METRIC_GROUPS):
        start = 5 + index * 3
        end = start + 2
        sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        sheet.cell(5, start).value = label
        for offset, period_label in enumerate(("本期", "同期", "同比")):
            sheet.cell(6, start + offset).value = period_label
            sheet.cell(7, start + offset).value = unit if offset < 2 else "%"

    identifiers = ("门店编码", "门店名称", f"{dimension_name}编码", f"{dimension_name}名称")
    for column, label in enumerate(identifiers, 1):
        sheet.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column)
        sheet.cell(6, column).value = label

    for row in range(5, 8):
        for column in range(1, 20):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER


def _write_group_headers(sheet) -> None:
    sheet.merge_cells("A5:J5")
    sheet["A5"] = "维度"
    for index, (label, unit) in enumerate(METRIC_GROUPS):
        start = 11 + index * 3
        end = start + 2
        sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        sheet.cell(5, start).value = label
        for offset, period_label in enumerate(("本期", "同期", "同比")):
            sheet.cell(6, start + offset).value = period_label
            sheet.cell(7, start + offset).value = unit if offset < 2 else "%"

    identifiers = (
        "门店编码",
        "门店名称",
        "部门编码",
        "部门名称",
        "区域编码",
        "区域名称",
        "品类编码",
        "品类名称",
        "柜组编码",
        "柜组名称",
    )
    for column, label in enumerate(identifiers, 1):
        sheet.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column)
        sheet.cell(6, column).value = label
    for row in range(5, 8):
        for column in range(1, 26):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER


def _write_special_sale_headers(sheet) -> None:
    sheet.merge_cells("A5:H5")
    sheet["A5"] = "维度"
    for index, (label, unit) in enumerate(METRIC_GROUPS):
        start = 9 + index * 3
        end = start + 2
        sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        sheet.cell(5, start).value = label
        for offset, period_label in enumerate(("本期", "同期", "同比")):
            sheet.cell(6, start + offset).value = period_label
            sheet.cell(7, start + offset).value = unit if offset < 2 else "%"

    identifiers = (
        "门店编码",
        "门店名称",
        "部门编码",
        "部门名称",
        "柜组编码",
        "柜组名称",
        "品牌编码",
        "品牌名称",
    )
    for column, label in enumerate(identifiers, 1):
        sheet.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column)
        sheet.cell(6, column).value = label
    for row in range(5, 8):
        for column in range(1, 24):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER


def _write_hierarchy_headers(sheet) -> None:
    sheet.merge_cells("A5:D5")
    sheet["A5"] = "层级"
    for column, label in enumerate(("门店", "部门", "区域", "品类"), 1):
        sheet.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column)
        sheet.cell(6, column).value = label
    for index, (label, unit) in enumerate(METRIC_GROUPS):
        start = 5 + index * 3
        end = start + 2
        sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        sheet.cell(5, start).value = label
        for offset, period_label in enumerate(("本期", "同期", "同比")):
            sheet.cell(6, start + offset).value = period_label
            sheet.cell(7, start + offset).value = unit if offset < 2 else "%"
    for row in range(5, 8):
        for column in range(1, 20):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER


def _metric_values(metrics: dict[str, Any]) -> list[Any]:
    return [
        int(metrics.get("ticket_count_current") or 0),
        int(metrics.get("ticket_count_prior") or 0),
        metrics.get("ticket_count_yoy"),
        metrics.get("average_ticket_current"),
        metrics.get("average_ticket_prior"),
        metrics.get("average_ticket_yoy"),
        float(metrics.get("sales_current") or 0) / 10000,
        float(metrics.get("sales_prior") or 0) / 10000,
        metrics.get("sales_yoy"),
        float(metrics.get("profit_current") or 0) / 10000,
        float(metrics.get("profit_prior") or 0) / 10000,
        metrics.get("profit_yoy"),
        metrics.get("margin_current"),
        metrics.get("margin_prior"),
        metrics.get("margin_change"),
    ]


def _safe_excel_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_data_row(
    sheet,
    row_number: int,
    values: list[Any],
    *,
    identifier_columns: int = 4,
    total: bool = False,
) -> None:
    for column, value in enumerate(values, 1):
        if column <= identifier_columns:
            value = _safe_excel_text(value)
        cell = sheet.cell(row_number, column, value)
        cell.border = BORDER
        metric_column = column - identifier_columns
        if metric_column in (1, 2):
            cell.number_format = "#,##0"
        elif metric_column in (4, 5, 7, 8, 10, 11):
            cell.number_format = "0.00"
        elif metric_column >= 3:
            cell.number_format = "0.00%"
        if total:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")


def _write_report_sheet(sheet, report: dict[str, Any], label: str, dimension_key: str) -> None:
    dates = report["dates"]
    sheet.merge_cells("A1:S1")
    sheet["A1"] = f"OD0002 门店销售毛利汇总表（{label}）"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:S2")
    sheet["A2"] = f"本期：{_date_text(dates['start_date'])} 至 {_date_text(dates['end_date'])}"
    sheet.merge_cells("A3:S3")
    sheet["A3"] = f"同期：{_date_text(dates['prior_start_date'])} 至 {_date_text(dates['prior_end_date'])}"
    _write_headers(sheet, label)

    row_number = 8
    for row in report.get("dimensions", {}).get(dimension_key, []):
        identifiers = [
            row.get("store_code"), row.get("store_name"),
            row.get("dimension_code"), row.get("dimension_name"),
        ]
        _write_data_row(sheet, row_number, identifiers + _metric_values(row.get("metrics", {})))
        row_number += 1
    total = report.get("totals", {}).get(dimension_key, {})
    _write_data_row(sheet, row_number, ["合计", None, None, None] + _metric_values(total), total=True)

    sheet.freeze_panes = "A8"
    widths = (14, 18, 16, 24) + (14,) * 15
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_group_sheet(sheet, report: dict[str, Any]) -> None:
    dates = report["dates"]
    sheet.merge_cells("A1:Y1")
    sheet["A1"] = "OD0002 门店销售毛利汇总表（柜组）"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:Y2")
    sheet["A2"] = f"本期：{_date_text(dates['start_date'])} 至 {_date_text(dates['end_date'])}"
    sheet.merge_cells("A3:Y3")
    sheet["A3"] = f"同期：{_date_text(dates['prior_start_date'])} 至 {_date_text(dates['prior_end_date'])}"
    _write_group_headers(sheet)

    row_number = 8
    for row in report.get("dimensions", {}).get("groups", []):
        identifiers = [
            row.get("store_code"),
            row.get("store_name"),
            row.get("department_code"),
            row.get("department_name"),
            row.get("area_code"),
            row.get("area_name"),
            row.get("category_code"),
            row.get("category_name"),
            row.get("dimension_code"),
            row.get("dimension_name"),
        ]
        _write_data_row(
            sheet,
            row_number,
            identifiers + _metric_values(row.get("metrics", {})),
            identifier_columns=10,
        )
        row_number += 1

    total = report.get("totals", {}).get("groups", {})
    _write_data_row(
        sheet,
        row_number,
        ["合计", None, None, None, None, None, None, None, None, None]
        + _metric_values(total),
        identifier_columns=10,
        total=True,
    )
    sheet.freeze_panes = "A8"
    widths = (14, 18, 16, 24, 14, 18, 14, 20, 18, 28) + (14,) * 15
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_special_sale_sheet(sheet, report: dict[str, Any]) -> None:
    dates = report["dates"]
    sheet.merge_cells("A1:W1")
    sheet["A1"] = "OD0002 门店销售毛利汇总表（特卖）"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:W2")
    sheet["A2"] = f"本期：{_date_text(dates['start_date'])} 至 {_date_text(dates['end_date'])}"
    sheet.merge_cells("A3:W3")
    sheet["A3"] = f"同期：{_date_text(dates['prior_start_date'])} 至 {_date_text(dates['prior_end_date'])}"
    _write_special_sale_headers(sheet)

    row_number = 8
    for row in report.get("dimensions", {}).get("special_sales", []):
        identifiers = [
            row.get("store_code"),
            row.get("store_name"),
            row.get("department_code"),
            row.get("department_name"),
            row.get("dimension_code"),
            row.get("dimension_name"),
            row.get("brand_code"),
            row.get("brand_name"),
        ]
        _write_data_row(
            sheet,
            row_number,
            identifiers + _metric_values(row.get("metrics", {})),
            identifier_columns=8,
        )
        row_number += 1

    total = report.get("totals", {}).get("special_sales", {})
    _write_data_row(
        sheet,
        row_number,
        ["合计", None, None, None, None, None, None, None] + _metric_values(total),
        identifier_columns=8,
        total=True,
    )
    sheet.freeze_panes = "A8"
    widths = (14, 18, 16, 24, 18, 28, 14, 28) + (14,) * 15
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _hierarchy_text(name: Any, code: Any, *, suffix: str = "") -> Any:
    safe_name = str(name or "未匹配") + suffix
    safe_code = str(code or "—")
    return _safe_excel_text(f"{safe_name}\n{safe_code}")


def _write_department_category_sheet(sheet, report: dict[str, Any]) -> None:
    dates = report["dates"]
    label = "部门（含品类）"
    sheet.merge_cells("A1:S1")
    sheet["A1"] = f"OD0002 门店销售毛利汇总表（{label}）"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:S2")
    sheet["A2"] = f"本期：{_date_text(dates['start_date'])} 至 {_date_text(dates['end_date'])}"
    sheet.merge_cells("A3:S3")
    sheet["A3"] = f"同期：{_date_text(dates['prior_start_date'])} 至 {_date_text(dates['prior_end_date'])}"
    _write_hierarchy_headers(sheet)

    row_number = 8
    hierarchy_rows = report.get("dimensions", {}).get("department_categories", [])
    for row in hierarchy_rows:
        row_type = row.get("row_type")
        if row_type == "area_subtotal":
            identifiers = [
                None,
                None,
                _hierarchy_text(row.get("area_name"), row.get("area_code"), suffix="小计"),
                None,
            ]
        elif row_type == "department_subtotal":
            identifiers = [
                None,
                _hierarchy_text(
                    row.get("department_name"), row.get("department_code"), suffix="小计"
                ),
                None,
                None,
            ]
        else:
            identifiers = [
                _hierarchy_text(row.get("store_name"), row.get("store_code")),
                _hierarchy_text(row.get("department_name"), row.get("department_code")),
                _hierarchy_text(row.get("area_name"), row.get("area_code")),
                _hierarchy_text(row.get("category_name"), row.get("category_code")),
            ]
        _write_data_row(
            sheet,
            row_number,
            identifiers + _metric_values(row.get("metrics", {})),
            total=row_type in ("area_subtotal", "department_subtotal"),
        )
        for column in range(1, 5):
            sheet.cell(row_number, column).alignment = Alignment(wrap_text=True, vertical="center")
        row_number += 1

    total = report.get("totals", {}).get("department_categories", {})
    _write_data_row(sheet, row_number, ["合计", None, None, None] + _metric_values(total), total=True)
    sheet.freeze_panes = "A8"
    widths = (22, 24, 22, 24) + (14,) * 15
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_notes(sheet, report: dict[str, Any]) -> None:
    sheet["A1"] = "项目"
    sheet["B1"] = "说明"
    sheet["A2"] = "OD0002 口径"
    scope = report.get("scope_description") or "当前用户权限范围：以系统数据权限为准"
    excluded_departments = "、".join(sorted(EXCLUDED_DEPARTMENT_CODES))
    sheet["B2"] = (
        "销售日期取 sglhsrq；销售收入取 sglxssr；毛利额取 sgln2；"
        "来客数按 salegoodslist.sglbillno 去重小票数计算；"
        "客单=销售收入（元）/来客数，来客数为0时客单留空；"
        "排除租赁 sglwmid=5；排除楼层00；排除16部门；"
        f"排除部门编码：{excluded_departments}；{scope}；"
        "合计仅包含当前用户权限范围；"
        "特卖Sheet仅列 manaframe.mflc=16 的柜组，"
        "品牌编码取 salegoodslist.sglppcode，并关联 codebrand.cbid 显示 codebrand.cbcname；"
        "大类暂不提供。"
    )
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.border = BORDER
    sheet["A2"].border = BORDER
    sheet["B2"].border = BORDER
    sheet["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 110
    sheet.row_dimensions[2].height = 120


def _build_od0002_workbook(report: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for label, dimension_key in SHEETS:
        sheet = workbook.create_sheet(label)
        if dimension_key == "groups":
            _write_group_sheet(sheet, report)
        elif dimension_key == "special_sales":
            _write_special_sale_sheet(sheet, report)
        else:
            _write_report_sheet(sheet, report, label, dimension_key)
        if dimension_key == "departments":
            _write_department_category_sheet(workbook.create_sheet("部门（含品类）"), report)
    _write_notes(workbook.create_sheet("报表说明"), report)
    return workbook


def build_od0002_workbook_file(report: dict[str, Any]) -> SpooledTemporaryFile:
    """Save the workbook directly to a seeked temporary file for streaming."""
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        _build_od0002_workbook(report).save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise


def build_od0002_workbook(report: dict[str, Any]) -> bytes:
    """Compatibility helper for callers that explicitly need workbook bytes."""
    output = build_od0002_workbook_file(report)
    try:
        return output.read()
    finally:
        output.close()
