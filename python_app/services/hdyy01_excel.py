from __future__ import annotations

from tempfile import SpooledTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .od0002_report import EXCLUDED_DEPARTMENT_CODES


BLUE = "4472C4"
WHITE = "FFFFFF"
TOTAL_FILL = "D9EAF7"
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DETAIL_COLUMNS = (
    ("机构", "store_name", "text"),
    ("部门", "department_name", "text"),
    ("柜组编码", "group_code", "text"),
    ("柜组名称", "group_name", "text"),
    ("面积", "area", "decimal"),
    ("楼层", "floor_name", "text"),
    ("数量", "quantity", "quantity"),
    ("销售收入", "sales_amount", "money"),
    ("含税销售成本", "tax_cost", "money"),
    ("毛利", "profit", "money"),
    ("消费次数", "ticket_count", "integer"),
    ("客单", "average_ticket", "money"),
    ("会员销售", "member_sales", "money"),
    ("储值卡销售", "stored_card_sales", "money"),
)


def _date_text(value: Any) -> str:
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _safe_excel_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _cell_value(row: dict[str, Any], key: str, kind: str) -> Any:
    value = row.get(key)
    if kind == "classification" and (value is None or not str(value).strip()):
        value = "未匹配"
    if kind in ("text", "classification"):
        return _safe_excel_text(value)
    if value is None:
        return None
    if kind == "integer":
        return int(value)
    return float(value)


def _number_format(kind: str) -> str | None:
    if kind in ("decimal", "money"):
        return "0.00"
    if kind == "quantity":
        return "0.####"
    if kind == "integer":
        return "0"
    return None


def _style_header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(color=WHITE, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER


def _write_only_cell(
    sheet,
    value: Any,
    *,
    kind: str | None = None,
    total: bool = False,
) -> WriteOnlyCell:
    cell = WriteOnlyCell(sheet, value=value)
    cell.border = BORDER
    number_format = _number_format(kind) if kind is not None else None
    if number_format is not None:
        cell.number_format = number_format
    if total:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
    return cell


def _write_detail_sheet(sheet, report: dict[str, Any]) -> None:
    final_column = get_column_letter(len(DETAIL_COLUMNS))
    dates = report.get("dates", {})
    rows = report.get("rows", [])
    last_data_row = max(4, 4 + len(rows))

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{final_column}{last_data_row}"
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[4].height = 28
    widths = (18, 20, 15, 22, 12, 10, 13, 15, 15, 15, 13, 15, 15, 16)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    title_cell = WriteOnlyCell(sheet, value="HDYY01柜组经营分析表")
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.append([title_cell])

    period_cell = WriteOnlyCell(sheet, value=(
        f"日期：{_date_text(dates.get('start_date'))} 至 "
        f"{_date_text(dates.get('end_date'))}；金额单位：元；面积单位：平方米"
    ))
    period_cell.alignment = Alignment(horizontal="left")
    sheet.append([period_cell])

    selected_store = report.get("selected_store") or "全部"
    selected_department = report.get("selected_department") or "全部"
    sheet.append([f"筛选：机构 {selected_store}；部门 {selected_department}"])

    header_cells = []
    for label, _key, _kind in DETAIL_COLUMNS:
        cell = WriteOnlyCell(sheet, value=label)
        _style_header(cell)
        header_cells.append(cell)
    sheet.append(header_cells)

    for row in rows:
        sheet.append([
            _write_only_cell(sheet, _cell_value(row, key, kind), kind=kind)
            for _label, key, kind in DETAIL_COLUMNS
        ])

    total = report.get("total", {})
    total_values = {
        "quantity": total.get("quantity", 0),
        "sales_amount": total.get("sales_amount", 0),
        "tax_cost": total.get("tax_cost", 0),
        "profit": total.get("profit", 0),
        "ticket_count": total.get("ticket_count", 0),
        "average_ticket": total.get("average_ticket"),
        "member_sales": total.get("member_sales", 0),
        "stored_card_sales": total.get("stored_card_sales", 0),
    }
    total_cells = []
    for column, (_label, key, kind) in enumerate(DETAIL_COLUMNS, 1):
        value = "合计" if column == 1 else total_values.get(key)
        if key in total_values and value is not None:
            value = _cell_value(total_values, key, kind)
        total_cells.append(
            _write_only_cell(sheet, value, kind=kind, total=True)
        )
    sheet.append(total_cells)


def _write_notes_sheet(sheet, report: dict[str, Any]) -> None:
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 120
    sheet.row_dimensions[2].height = 100

    header_cells = [WriteOnlyCell(sheet, value=value) for value in ("项目", "说明")]
    for cell in header_cells:
        _style_header(cell)
    sheet.append(header_cells)

    excluded = "、".join(sorted(EXCLUDED_DEPARTMENT_CODES))
    scope = report.get("scope_description") or "当前用户权限范围：以系统数据权限为准"
    dates = report.get("dates", {})
    notes = (
        f"报表期间：{_date_text(dates.get('start_date'))} 至 "
        f"{_date_text(dates.get('end_date'))}；"
        "销售日期取 sglhsrq；销售收入取 sglxssr，金额单位为元；"
        "含税销售成本取 sgln13+sgln14-sglsupzk；毛利取 sgln2；"
        "储值卡销售取 sglfcard；会员销售以 salehead.hykh 非空识别；"
        "退货按带符号金额计入；仅小票净销售额 > 0 时计入消费次数；"
        "客单 = 带符号销售额 / 正向消费次数；排除租赁业务 sglwmid=5；"
        f"排除部门编码：{excluded}；{scope}；"
        "分类层级使用 manaframe.mfchr2 关联 mana_brand_hierarchy，按编码层级关联；"
        "分类缺失显示未匹配；"
        "未匹配会员小票数因细粒度权限无法安全归属，当前不可计算。"
    )
    label_cell = _write_only_cell(sheet, "HDYY01 口径")
    label_cell.alignment = Alignment(vertical="top")
    notes_cell = _write_only_cell(sheet, notes)
    notes_cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.append([label_cell, notes_cell])


def _build_hdyy01_workbook(report: dict[str, Any]) -> Workbook:
    workbook = Workbook(write_only=True)
    detail = workbook.create_sheet("明细")
    _write_detail_sheet(detail, report)
    _write_notes_sheet(workbook.create_sheet("报表说明"), report)
    return workbook


def build_hdyy01_workbook_file(report: dict[str, Any]) -> SpooledTemporaryFile:
    """Build a seeked temporary XLSX file suitable for chunked streaming."""
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        _build_hdyy01_workbook(report).save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise


def build_hdyy01_workbook(report: dict[str, Any]) -> bytes:
    """Build the HDYY01 workbook as bytes for compatibility and tests."""
    output = build_hdyy01_workbook_file(report)
    try:
        return output.read()
    finally:
        output.close()
