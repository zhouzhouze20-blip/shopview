"""Read the persisted financial-month inventory report with business scope."""
import re
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.orm import Session

from models.database import get_db
from models.models import User
from routers.auth import get_current_user
from routers.authz import load_business_scope, require_permission
from routers.sales import _business_scope_filter_sql

router = APIRouter(prefix='/api/inventory-turnover', tags=['inventory'])
PERMISSION = 'sales.inventory_turnover.view'
NOTE = ('经销、成本代销；按核算日期含退货净额。销售成本按每行进项税率转不含税。'
        '周转率=销售成本/平均库存成本；周转天数=平均库存成本/销售成本×已统计天数；'
        '存销比=期末库存成本/销售成本；覆盖天数=存销比×已统计天数。'
        '本月统计至前一天；缺日期或非正分母时相关指标不适用。')
COLUMNS = [
    ('store_code', '门店编码'), ('store_name', '门店名称'),
    ('department_code', '部门编码'), ('department_name', '部门名称'),
    ('financial_month', '财务年月'), ('supplier_code', '供应商编码'),
    ('supplier_name', '供应商名称'), ('group_code', '柜组编码'), ('group_name', '柜组名称'),
    ('operation_method', '经营方式'), ('period_start', '期间开始'), ('period_end', '期间结束'),
    ('as_of_date', '统计截止'), ('elapsed_days', '已统计天数'), ('inventory_days', '快照覆盖天数'),
    ('sales_days', '门店销售覆盖天数'), ('sales_quantity', '销售数量'), ('sales_revenue', '销售收入'),
    ('sales_cost_ex_tax', '不含税销售成本'), ('ending_quantity', '期末库存数量'),
    ('ending_cost_ex_tax', '期末不含税成本'), ('average_cost_ex_tax', '不含税平均库存'),
    ('turnover_days', '周转天数'), ('turnover_rate', '周转率'), ('stock_sales_ratio', '期末存销比'),
    ('stock_cover_days', '期末库存覆盖天数'), ('data_status', '数据状态'), ('updated_at', '更新时间'),
]


BRAND_COLUMNS = [(key, title) for key, title in COLUMNS
                 if key not in {'supplier_code', 'supplier_name', 'operation_method'}]
BRAND_COLUMNS = [(key, '品牌柜组' if key == 'group_name' else title) for key, title in BRAND_COLUMNS]
BRAND_COLUMNS.insert(9, ('supplier_count', '供应商数'))
BRAND_COLUMNS.append(('exception_count', '异常明细数'))
BRAND_NOTE = '沿用原Excel品牌柜组口径；同柜组供应商合并，跨柜组不按同名合并。先按权限筛选，再汇总金额并重算比率；仅包含当前账号授权数据。'


def summarize_brands(rows):
    groups = defaultdict(list)
    for row in rows:
        groups[(row['store_code'], row['department_code'], row['financial_month'], row['group_code'])].append(row)
    result = []
    amounts = ['sales_quantity', 'sales_revenue', 'sales_cost_ex_tax', 'ending_quantity',
               'ending_cost_ex_tax', 'average_cost_ex_tax']
    for members in groups.values():
        row = {key: value for key, value in members[0].items()
               if key not in {'supplier_code', 'supplier_name', 'operation_method'}}
        for key in amounts:
            values = [m[key] for m in members]
            row[key] = None if any(v is None for v in values) else sum(values, Decimal(0))
        row['supplier_count'] = len({m['supplier_code'] for m in members})
        row['exception_count'] = sum(m['data_status'] != '正常' for m in members)
        row['updated_at'] = min(m['updated_at'] for m in members)
        row['inventory_days'] = min(m['inventory_days'] for m in members)
        row['sales_days'] = min(m['sales_days'] for m in members)
        aligned = all(len({m[k] for m in members}) == 1 for k in
                      ['period_start', 'period_end', 'as_of_date', 'elapsed_days'])
        statuses = sorted({m['data_status'] for m in members if m['data_status'] != '正常'})
        if not aligned:
            statuses.append('明细统计期间不一致')
            row['average_cost_ex_tax'] = None
            row['as_of_date'] = min(m['as_of_date'] for m in members)
        row['data_status'] = ('明细异常：' + '；'.join(statuses)) if statuses else '正常'
        cost, average, ending = (row[k] for k in ['sales_cost_ex_tax', 'average_cost_ex_tax', 'ending_cost_ex_tax'])
        valid = aligned and row['sales_days'] == row['elapsed_days'] and cost is not None and cost > 0
        turnover_valid = valid and average is not None and average > 0
        stock_valid = valid and ending is not None and ending >= 0
        row['turnover_rate'] = cost / average if turnover_valid else None
        row['turnover_days'] = average / cost * row['elapsed_days'] if turnover_valid else None
        row['stock_sales_ratio'] = ending / cost if stock_valid else None
        row['stock_cover_days'] = ending / cost * row['elapsed_days'] if stock_valid else None
        result.append(row)
    return result


def scoped_where(db, user):
    require_permission(db, user, PERMISSION)
    scope = load_business_scope(db, user, fallback_resource_code='sales')
    # A cabinet/supplier aggregate cannot safely expose a partial brand scope.
    if scope.allow.get('brand') or scope.deny.get('brand'):
        raise HTTPException(403, '此汇总报表不支持品牌级数据范围，请使用柜组或部门范围授权')
    params = {}
    where = _business_scope_filter_sql(
        scope, params, prefix='turnover', store_expr='r.store_id::text',
        department_code_expr='r.department_code', department_name_expr='r.department_name',
        group_expr='r.group_code', supplier_expr='r.supplier_code',
        category_code_expr='r.category_code', category_name_expr='r.category_name', floor_expr='r.floor_code')
    return ' WHERE 1=1 ' + where, params


def filter_where(where, params, store_code=None, department_code=None, financial_month=None):
    if financial_month and not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', financial_month):
        raise HTTPException(422, '财务年月须为YYYY-MM')
    for key, value in [('store_code', store_code), ('department_code', department_code), ('financial_month', financial_month)]:
        if value:
            where += f' AND r.{key} = :{key}'
            params[key] = value
    return where


@router.get('/options')
def options(store_code: str | None = None, department_code: str | None = None,
            db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    where, params = scoped_where(db, user)
    stores = db.execute(text('SELECT DISTINCT r.store_code,r.store_name FROM inventory_turnover_monthly r' + where + ' ORDER BY r.store_code'), params).mappings().all()
    where = filter_where(where, params, store_code=store_code)
    departments = db.execute(text('SELECT DISTINCT r.department_code,r.department_name FROM inventory_turnover_monthly r' + where + ' ORDER BY r.department_code'), params).mappings().all() if store_code else []
    where = filter_where(where, params, department_code=department_code)
    months = db.execute(text('SELECT DISTINCT r.financial_month FROM inventory_turnover_monthly r' + where + ' ORDER BY r.financial_month DESC'), params).scalars().all() if store_code else []
    return dict(stores=[dict(r) for r in stores], departments=[dict(r) for r in departments], months=months)


def report(db, user, store_code, department_code, financial_month):
    where, params = scoped_where(db, user)
    where = filter_where(where, params, store_code, department_code, financial_month)
    rows = [dict(r) for r in db.execute(text('SELECT r.* FROM inventory_turnover_monthly r' + where + ' ORDER BY r.department_code,r.group_code,r.supplier_code,r.operation_method'), params).mappings()]
    return dict(rows=rows, count=len(rows), brands=summarize_brands(rows), brand_note=BRAND_NOTE, note=NOTE,
                filters=dict(store_code=store_code, department_code=department_code, financial_month=financial_month))


@router.get('')
def get_report(store_code: str = Query(..., min_length=1), financial_month: str = Query(...),
               department_code: str | None = None, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)):
    return report(db, user, store_code, department_code, financial_month)


@router.get('/export')
def export_report(store_code: str = Query(..., min_length=1), financial_month: str = Query(...),
                  department_code: str | None = None, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    result = report(db, user, store_code, department_code, financial_month)
    wb = Workbook()
    wb.remove(wb.active)
    for title, columns, rows, note in [
        ('按品牌周转率汇总', BRAND_COLUMNS, result['brands'], BRAND_NOTE + NOTE),
        ('供应商明细', COLUMNS, result['rows'], NOTE),
    ]:
        ws = wb.create_sheet(title)
        ws.append([note]); ws.append([label for _, label in columns])
        for row in rows:
            values = []
            for key, _ in columns:
                value = row[key]
                if key == 'updated_at': value = value.isoformat()
                if key == 'operation_method': value = {'1': '经销', '2': '成本代销'}.get(value, value)
                values.append(value)
            ws.append(values)
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value, str): cell.data_type = 's'
                if cell.row > 2 and isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = '#,##0.00##'
        for cell in ws[2]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78')
            ws.column_dimensions[cell.column_letter].width = 22
        ws.freeze_panes = 'I3'
        ws.auto_filter.ref = f'A2:{ws.cell(2, len(columns)).column_letter}{ws.max_row}'
    buffer = BytesIO(); wb.save(buffer); buffer.seek(0)
    filename = quote(f'商品周转财务月报_{store_code}_{financial_month}.xlsx')
    return StreamingResponse(buffer, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=UTF-8''{filename}"})
