"""品类主管品牌维护与实时绩效 API。"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from models.database import get_db
from models.models import OperationLog, User
from routers.auth import get_current_user
from routers.authz import (
    get_authz_subject,
    is_admin,
    load_business_scope,
    require_permission,
    scope_allows_business,
)
from services.category_performance import build_score_row, parse_period_month


router = APIRouter(prefix="/api/category-performance", tags=["category-performance"])

VIEW_PERMISSION = "sales.category_performance.view"
MANAGE_PERMISSION = "sales.category_performance.manage"


class BrandAssignmentUpdate(BaseModel):
    manager_user_id: int
    is_active: bool = True
    is_key_brand: bool | None = None


class KeyBrandTargetUpdate(BaseModel):
    manager_user_id: int
    sales_target: Decimal = Field(ge=0)
    is_active: bool = True


class ManagerTargetUpdate(BaseModel):
    manager_user_id: int
    area_revenue_target: Decimal = Field(ge=0)
    area_weight: Decimal = Field(default=Decimal("40"), ge=0)
    key_brand_weight: Decimal = Field(default=Decimal("40"), ge=0)
    self_weight: Decimal = Field(default=Decimal("20"), ge=0)
    self_score: Decimal | None = Field(default=None, ge=0)
    assessment_content: str | None = None


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _rows(rows) -> list[dict[str, Any]]:
    return [
        {key: _json_value(value) for key, value in row.items()}
        for row in rows
    ]


def _period_window(period: str):
    try:
        return parse_period_month(period)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(exc),
        ) from exc


def _table_exists(db: Session, table_name: str) -> bool:
    return bool(
        db.execute(
            text(
                """
                SELECT EXISTS (
                  SELECT 1
                  FROM information_schema.tables
                  WHERE table_schema = 'public' AND table_name = :table_name
                )
                """
            ),
            {"table_name": table_name},
        ).scalar()
    )


def _require_tables(db: Session) -> None:
    required = (
        "category_manager_brand_assignments",
        "category_key_brand_targets",
        "category_manager_performance_targets",
    )
    missing = [table_name for table_name in required if not _table_exists(db, table_name)]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="品类绩效维护表尚未创建，请先执行数据库升级",
        )


def _has_permission(db: Session, user: User, permission_code: str) -> bool:
    subject = get_authz_subject(db, user)
    if is_admin(db, subject):
        return True
    return bool(
        db.execute(
            text(
                """
                SELECT 1
                FROM user_roles ur
                JOIN roles r ON r.id = ur.role_id AND r.is_active
                JOIN role_permissions rp ON rp.role_id = r.id
                JOIN permissions p ON p.id = rp.permission_id
                WHERE ur.user_id = :user_id
                  AND p.permission_code = :permission_code
                  AND (ur.expires_at IS NULL OR ur.expires_at > NOW())
                LIMIT 1
                """
            ),
            {"user_id": subject.user_id, "permission_code": permission_code},
        ).first()
    )


def _store_row(db: Session, store_id: int):
    row = db.execute(
        text(
            """
            SELECT store_id, store_code, store_name
            FROM stores
            WHERE store_id = :store_id AND COALESCE(is_active, TRUE)
            """
        ),
        {"store_id": store_id},
    ).mappings().first()
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="门店不存在或已停用")
    return row


def _all_store_groups(db: Session, store_row) -> list[dict[str, Any]]:
    return _rows(
        db.execute(
            text(
                """
                SELECT
                  mf.mfcode AS group_code,
                  mf.mfcname AS group_name,
                  dept.mfcode AS department_code,
                  dept.mfcname AS department_name,
                  mf.mfjyqy AS area_name,
                  mf.mfjyfs AS operation_method,
                  COALESCE(kb.is_key_brand, FALSE) AS is_key_brand
                FROM manaframe mf
                LEFT JOIN manaframe dept
                  ON UPPER(TRIM(dept.mfcode)) = UPPER(TRIM(mf.mfpcode))
                LEFT JOIN manaframe_key_brand kb
                  ON UPPER(TRIM(kb.mfcode)) = UPPER(TRIM(mf.mfcode))
                WHERE LEFT(TRIM(mf.mfcode), 3) = :store_code
                  AND NULLIF(TRIM(mf.mfpcode), '') IS NOT NULL
                  AND NOT EXISTS (
                    SELECT 1
                    FROM manaframe child
                    WHERE UPPER(TRIM(child.mfpcode)) = UPPER(TRIM(mf.mfcode))
                  )
                ORDER BY dept.mfcode, mf.mfcode
                """
            ),
            {"store_code": str(store_row["store_code"]).strip()},
        ).mappings().all()
    )


def _allowed_groups(db: Session, current_user: User, store_row) -> list[dict[str, Any]]:
    scope = load_business_scope(db, current_user, fallback_resource_code="sales")
    result = []
    for group in _all_store_groups(db, store_row):
        if scope_allows_business(
            scope,
            store_id=store_row["store_id"],
            department_code=group.get("department_code"),
            department_name=group.get("department_name"),
            group_code=group.get("group_code"),
        ):
            result.append(group)
    return result


def _require_allowed_group(
    db: Session,
    current_user: User,
    store_row,
    group_code: str,
) -> dict[str, Any]:
    normalized = str(group_code or "").strip().upper()
    group = next(
        (
            item
            for item in _allowed_groups(db, current_user, store_row)
            if str(item.get("group_code") or "").strip().upper() == normalized
        ),
        None,
    )
    if not group:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="该品牌不在当前账号的数据范围内")
    return group


def _manager(db: Session, manager_user_id: int):
    manager = db.execute(
        text(
            """
            SELECT user_id, real_name, username, employee_no
            FROM users
            WHERE user_id = :user_id
              AND is_active
              AND COALESCE(status, 'ACTIVE') = 'ACTIVE'
            """
        ),
        {"user_id": manager_user_id},
    ).mappings().first()
    if not manager:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="品类主管账号不存在或已停用")
    if not str(manager["real_name"] or "").strip():
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="品类主管账号未维护姓名")
    return manager


def _ensure_coupon_followup_role(db: Session, *, manager_user_id: int, store_row) -> None:
    """New 601 brand managers immediately receive only the mobile follow-up module role."""
    if str(store_row["store_code"] or "").strip() != "601":
        return
    db.execute(
        text(
            """
            INSERT INTO user_roles (user_id, role_id, store_id, created_at)
            SELECT :manager_user_id, role_row.id, :store_id, NOW()
            FROM roles role_row
            WHERE role_row.role_code = 'category_coupon_followup_manager'
              AND role_row.is_active
            ON CONFLICT DO NOTHING
            """
        ),
        {
            "manager_user_id": manager_user_id,
            "store_id": store_row["store_id"],
        },
    )


def _ensure_supplier_payment_role(db: Session, *, manager_user_id: int, store_row) -> None:
    """Active category managers receive the mobile payment entry; data stays assignment-scoped."""
    db.execute(
        text(
            """
            INSERT INTO user_roles (user_id, role_id, store_id, created_at)
            SELECT :manager_user_id, role_row.id, :store_id, NOW()
            FROM roles role_row
            WHERE role_row.role_code = 'category_supplier_payment_viewer'
              AND role_row.is_active
            ON CONFLICT DO NOTHING
            """
        ),
        {
            "manager_user_id": manager_user_id,
            "store_id": store_row["store_id"],
        },
    )


def _operation_log(
    db: Session,
    current_user: User,
    *,
    action_code: str,
    target_id: str,
    detail: dict[str, Any],
) -> None:
    db.add(
        OperationLog(
            user_id=current_user.user_id,
            action_code=action_code,
            resource_code="category_performance",
            target_id=target_id,
            detail=detail,
        )
    )


@router.get("/options")
async def category_performance_options(
    store_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, VIEW_PERMISSION)
    _require_tables(db)
    store = _store_row(db, store_id)
    groups = _allowed_groups(db, current_user, store)
    latest_period = db.execute(
        text(
            """
            SELECT TO_CHAR(MAX(period_month), 'YYYY-MM')
            FROM (
              SELECT period_month
              FROM category_key_brand_targets
              WHERE store_id = :store_id AND is_active
              UNION ALL
              SELECT period_month
              FROM category_manager_performance_targets
              WHERE store_id = :store_id
              UNION ALL
              SELECT DATE_TRUNC('month', MAX(sglhsrq))::date AS period_month
              FROM salegoodslist
              WHERE sglmarket = :store_code
            ) maintained_periods
            """
        ),
        {
            "store_id": store_id,
            "store_code": str(store["store_code"]).strip(),
        },
    ).scalar()
    manager_rows = db.execute(
        text(
            """
            SELECT user_id, real_name, username, employee_no
            FROM users
            WHERE is_active
              AND COALESCE(status, 'ACTIVE') = 'ACTIVE'
              AND NULLIF(TRIM(real_name), '') IS NOT NULL
            ORDER BY real_name, user_id
            """
        )
    ).mappings().all()
    return {
        "store": dict(store),
        "can_manage": _has_permission(db, current_user, MANAGE_PERMISSION),
        "latest_period": latest_period,
        "groups": groups,
        "managers": _rows(manager_rows),
    }


@router.get("/brand-assignments")
async def list_brand_assignments(
    store_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, VIEW_PERMISSION)
    _require_tables(db)
    store = _store_row(db, store_id)
    allowed = {str(item["group_code"]).strip().upper() for item in _allowed_groups(db, current_user, store)}
    rows = _rows(
        db.execute(
            text(
                """
                SELECT
                  a.id,
                  a.group_code,
                  mf.mfcname AS group_name,
                  dept.mfcode AS department_code,
                  dept.mfcname AS department_name,
                  a.manager_user_id,
                  a.manager_name,
                  COALESCE(kb.is_key_brand, FALSE) AS is_key_brand,
                  a.is_active,
                  a.updated_at
                FROM category_manager_brand_assignments a
                LEFT JOIN manaframe mf
                  ON UPPER(TRIM(mf.mfcode)) = UPPER(TRIM(a.group_code))
                LEFT JOIN manaframe dept
                  ON UPPER(TRIM(dept.mfcode)) = UPPER(TRIM(mf.mfpcode))
                LEFT JOIN manaframe_key_brand kb
                  ON UPPER(TRIM(kb.mfcode)) = UPPER(TRIM(a.group_code))
                WHERE a.store_id = :store_id
                ORDER BY dept.mfcode, a.group_code
                """
            ),
            {"store_id": store_id},
        ).mappings().all()
    )
    return [row for row in rows if str(row["group_code"]).strip().upper() in allowed]


@router.put("/brand-assignments/{group_code}")
async def upsert_brand_assignment(
    group_code: str,
    payload: BrandAssignmentUpdate,
    store_id: int = Query(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, MANAGE_PERMISSION)
    _require_tables(db)
    store = _store_row(db, store_id)
    group = _require_allowed_group(db, current_user, store, group_code)
    manager = _manager(db, payload.manager_user_id)
    normalized_group_code = str(group["group_code"]).strip()
    manager_name = str(manager["real_name"]).strip()
    db.execute(
        text(
            """
            INSERT INTO category_manager_brand_assignments (
              store_id, group_code, manager_user_id, manager_name, is_active,
              created_by, updated_by
            )
            VALUES (
              :store_id, :group_code, :manager_user_id, :manager_name, :is_active,
              :user_id, :user_id
            )
            ON CONFLICT (store_id, group_code) DO UPDATE SET
              manager_user_id = EXCLUDED.manager_user_id,
              manager_name = EXCLUDED.manager_name,
              is_active = EXCLUDED.is_active,
              updated_by = EXCLUDED.updated_by,
              updated_at = NOW()
            """
        ),
        {
            "store_id": store_id,
            "group_code": normalized_group_code,
            "manager_user_id": manager["user_id"],
            "manager_name": manager_name,
            "is_active": payload.is_active,
            "user_id": current_user.user_id,
        },
    )
    if payload.is_active:
        _ensure_coupon_followup_role(
            db,
            manager_user_id=int(manager["user_id"]),
            store_row=store,
        )
        _ensure_supplier_payment_role(
            db,
            manager_user_id=int(manager["user_id"]),
            store_row=store,
        )
    if payload.is_key_brand is not None:
        db.execute(
            text(
                """
                INSERT INTO manaframe_key_brand (mfcode, is_key_brand, updated_at, updated_by)
                VALUES (:group_code, :is_key_brand, NOW(), :updated_by)
                ON CONFLICT (mfcode) DO UPDATE SET
                  is_key_brand = EXCLUDED.is_key_brand,
                  updated_at = NOW(),
                  updated_by = EXCLUDED.updated_by
                """
            ),
            {
                "group_code": normalized_group_code,
                "is_key_brand": payload.is_key_brand,
                "updated_by": str(current_user.user_id),
            },
        )
    _operation_log(
        db,
        current_user,
        action_code="category_performance.brand_assignment.update",
        target_id=f"{store_id}:{normalized_group_code}",
        detail={
            "manager_user_id": manager["user_id"],
            "manager_name": manager_name,
            "is_active": payload.is_active,
            "is_key_brand": payload.is_key_brand,
        },
    )
    db.commit()
    return {
        "group_code": normalized_group_code,
        "manager_user_id": manager["user_id"],
        "manager_name": manager_name,
        "is_key_brand": payload.is_key_brand,
    }


@router.get("/key-brand-targets")
async def list_key_brand_targets(
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, VIEW_PERMISSION)
    _require_tables(db)
    window = _period_window(period)
    store = _store_row(db, store_id)
    allowed = {str(item["group_code"]).strip().upper() for item in _allowed_groups(db, current_user, store)}
    rows = _rows(
        db.execute(
            text(
                """
                SELECT
                  t.id,
                  t.group_code,
                  mf.mfcname AS group_name,
                  dept.mfcode AS department_code,
                  dept.mfcname AS department_name,
                  t.manager_user_id,
                  t.manager_name,
                  t.sales_target,
                  t.is_active,
                  t.updated_at
                FROM category_key_brand_targets t
                LEFT JOIN manaframe mf
                  ON UPPER(TRIM(mf.mfcode)) = UPPER(TRIM(t.group_code))
                LEFT JOIN manaframe dept
                  ON UPPER(TRIM(dept.mfcode)) = UPPER(TRIM(mf.mfpcode))
                WHERE t.store_id = :store_id
                  AND t.period_month = :period_month
                ORDER BY dept.mfcode, t.group_code
                """
            ),
            {"store_id": store_id, "period_month": window.period_month},
        ).mappings().all()
    )
    return [row for row in rows if str(row["group_code"]).strip().upper() in allowed]


@router.put("/key-brand-targets/{group_code}")
async def upsert_key_brand_target(
    group_code: str,
    payload: KeyBrandTargetUpdate,
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, MANAGE_PERMISSION)
    _require_tables(db)
    window = _period_window(period)
    store = _store_row(db, store_id)
    group = _require_allowed_group(db, current_user, store, group_code)
    manager = _manager(db, payload.manager_user_id)
    normalized_group_code = str(group["group_code"]).strip()
    manager_name = str(manager["real_name"]).strip()
    db.execute(
        text(
            """
            INSERT INTO category_key_brand_targets (
              store_id, group_code, period_month, manager_user_id, manager_name,
              sales_target, is_active, created_by, updated_by
            )
            VALUES (
              :store_id, :group_code, :period_month, :manager_user_id, :manager_name,
              :sales_target, :is_active, :user_id, :user_id
            )
            ON CONFLICT (store_id, group_code, period_month) DO UPDATE SET
              manager_user_id = EXCLUDED.manager_user_id,
              manager_name = EXCLUDED.manager_name,
              sales_target = EXCLUDED.sales_target,
              is_active = EXCLUDED.is_active,
              updated_by = EXCLUDED.updated_by,
              updated_at = NOW()
            """
        ),
        {
            "store_id": store_id,
            "group_code": normalized_group_code,
            "period_month": window.period_month,
            "manager_user_id": manager["user_id"],
            "manager_name": manager_name,
            "sales_target": payload.sales_target,
            "is_active": payload.is_active,
            "user_id": current_user.user_id,
        },
    )
    db.execute(
        text(
            """
            INSERT INTO manaframe_key_brand (mfcode, is_key_brand, updated_at, updated_by)
            VALUES (:group_code, TRUE, NOW(), :updated_by)
            ON CONFLICT (mfcode) DO UPDATE SET
              is_key_brand = TRUE,
              updated_at = NOW(),
              updated_by = EXCLUDED.updated_by
            """
        ),
        {"group_code": normalized_group_code, "updated_by": str(current_user.user_id)},
    )
    _operation_log(
        db,
        current_user,
        action_code="category_performance.key_brand_target.update",
        target_id=f"{store_id}:{period}:{normalized_group_code}",
        detail={"manager_name": manager_name, "sales_target": float(payload.sales_target), "is_active": payload.is_active},
    )
    db.commit()
    return {"group_code": normalized_group_code, "manager_user_id": manager["user_id"], "manager_name": manager_name}


@router.get("/manager-targets")
async def list_manager_targets(
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, VIEW_PERMISSION)
    _require_tables(db)
    window = _period_window(period)
    store = _store_row(db, store_id)
    allowed_groups = _allowed_groups(db, current_user, store)
    allowed_codes = [str(item["group_code"]).strip().upper() for item in allowed_groups]
    if not allowed_codes:
        return []
    rows = db.execute(
        text(
            """
            SELECT
              t.id,
              t.manager_user_id,
              t.manager_name,
              t.area_revenue_target,
              t.area_weight,
              t.key_brand_weight,
              t.self_weight,
              t.self_score,
              t.assessment_content,
              t.updated_at
            FROM category_manager_performance_targets t
            WHERE t.store_id = :store_id
              AND t.period_month = :period_month
              AND (
                EXISTS (
                  SELECT 1 FROM category_manager_brand_assignments a
                  WHERE a.store_id = t.store_id
                    AND a.manager_user_id = t.manager_user_id
                    AND a.is_active
                    AND UPPER(TRIM(a.group_code)) = ANY(:allowed_codes)
                )
                OR EXISTS (
                  SELECT 1 FROM category_key_brand_targets k
                  WHERE k.store_id = t.store_id
                    AND k.period_month = t.period_month
                    AND k.manager_user_id = t.manager_user_id
                    AND k.is_active
                    AND UPPER(TRIM(k.group_code)) = ANY(:allowed_codes)
                )
              )
            ORDER BY t.manager_name, t.manager_user_id
            """
        ),
        {"store_id": store_id, "period_month": window.period_month, "allowed_codes": allowed_codes},
    ).mappings().all()
    return _rows(rows)


@router.put("/manager-targets/{manager_user_id}")
async def upsert_manager_target(
    manager_user_id: int,
    payload: ManagerTargetUpdate,
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, MANAGE_PERMISSION)
    _require_tables(db)
    if manager_user_id != payload.manager_user_id:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="主管参数不一致")
    weights_total = payload.area_weight + payload.key_brand_weight + payload.self_weight
    if weights_total != Decimal("100"):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="三项权重合计必须为100")
    if payload.self_score is not None and payload.self_score > payload.self_weight:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="部门自定得分不能超过其权重")
    window = _period_window(period)
    store = _store_row(db, store_id)
    manager = _manager(db, manager_user_id)
    allowed_codes = {
        str(item["group_code"]).strip().upper()
        for item in _allowed_groups(db, current_user, store)
    }
    manager_has_visible_brand = False
    if allowed_codes:
        manager_has_visible_brand = bool(
            db.execute(
                text(
                    """
                    SELECT 1
                    FROM (
                      SELECT group_code
                      FROM category_manager_brand_assignments
                      WHERE store_id = :store_id
                        AND manager_user_id = :manager_user_id
                        AND is_active
                      UNION ALL
                      SELECT group_code
                      FROM category_key_brand_targets
                      WHERE store_id = :store_id
                        AND period_month = :period_month
                        AND manager_user_id = :manager_user_id
                        AND is_active
                    ) assigned
                    WHERE UPPER(TRIM(group_code)) = ANY(:allowed_codes)
                    LIMIT 1
                    """
                ),
                {
                    "store_id": store_id,
                    "period_month": window.period_month,
                    "manager_user_id": manager_user_id,
                    "allowed_codes": list(allowed_codes),
                },
            ).first()
        )
    if not manager_has_visible_brand:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="该主管尚未关联当前账号可维护的品牌",
        )
    manager_name = str(manager["real_name"]).strip()
    db.execute(
        text(
            """
            INSERT INTO category_manager_performance_targets (
              store_id, period_month, manager_user_id, manager_name,
              area_revenue_target, area_weight, key_brand_weight, self_weight,
              self_score, assessment_content, created_by, updated_by
            )
            VALUES (
              :store_id, :period_month, :manager_user_id, :manager_name,
              :area_revenue_target, :area_weight, :key_brand_weight, :self_weight,
              :self_score, :assessment_content, :user_id, :user_id
            )
            ON CONFLICT (store_id, manager_user_id, period_month) DO UPDATE SET
              manager_name = EXCLUDED.manager_name,
              area_revenue_target = EXCLUDED.area_revenue_target,
              area_weight = EXCLUDED.area_weight,
              key_brand_weight = EXCLUDED.key_brand_weight,
              self_weight = EXCLUDED.self_weight,
              self_score = EXCLUDED.self_score,
              assessment_content = EXCLUDED.assessment_content,
              updated_by = EXCLUDED.updated_by,
              updated_at = NOW()
            """
        ),
        {
            "store_id": store_id,
            "period_month": window.period_month,
            "manager_user_id": manager_user_id,
            "manager_name": manager_name,
            "area_revenue_target": payload.area_revenue_target,
            "area_weight": payload.area_weight,
            "key_brand_weight": payload.key_brand_weight,
            "self_weight": payload.self_weight,
            "self_score": payload.self_score,
            "assessment_content": (payload.assessment_content or "").strip() or None,
            "user_id": current_user.user_id,
        },
    )
    _operation_log(
        db,
        current_user,
        action_code="category_performance.manager_target.update",
        target_id=f"{store_id}:{period}:{manager_user_id}",
        detail={
            "manager_name": manager_name,
            "area_revenue_target": float(payload.area_revenue_target),
            "weights": [float(payload.area_weight), float(payload.key_brand_weight), float(payload.self_weight)],
            "self_score": float(payload.self_score) if payload.self_score is not None else None,
        },
    )
    db.commit()
    return {"manager_user_id": manager_user_id, "manager_name": manager_name}


def _manager_actuals(
    db: Session,
    *,
    store_id: int,
    store_code: str,
    start_date: date,
    end_date: date,
    allowed_codes: list[str],
) -> dict[int, Decimal]:
    if not allowed_codes:
        return {}
    rows = db.execute(
        text(
            """
            WITH sales_by_group AS (
              SELECT
                UPPER(TRIM(s.sglmfid)) AS group_code_norm,
                COALESCE(
                  SUM(COALESCE(s.sgln2, 0) / NULLIF(1 + COALESCE(s.sglxstax, 0), 0)),
                  0
                ) AS amount
              FROM salegoodslist s
              WHERE s.sglhsrq BETWEEN :start_date AND :end_date
                AND s.sglmarket = :store_code
              GROUP BY UPPER(TRIM(s.sglmfid))
            ),
            fee_source AS (
              SELECT
                UPPER(TRIM(fee.source_group_code)) AS group_code_norm,
                CASE
                  WHEN TRIM(COALESCE(fee.fee_type_name, '')) LIKE '损失承担%'
                  THEN COALESCE(fee.tax_excluded_amount, 0) / 1.13
                  ELSE COALESCE(fee.tax_excluded_amount, 0)
                END AS amount,
                ROW_NUMBER() OVER (
                  PARTITION BY
                    fee.revenue_date,
                    UPPER(TRIM(COALESCE(fee.source_group_code, ''))),
                    UPPER(TRIM(COALESCE(fee.contract_code, ''))),
                    UPPER(TRIM(COALESCE(fee.fee_type_code, ''))),
                    TRIM(COALESCE(fee.source_doc_no, '')),
                    TRIM(COALESCE(fee.source_row_key, '')),
                    COALESCE(fee.tax_included_amount, 0),
                    COALESCE(fee.tax_excluded_amount, 0)
                  ORDER BY fee.id
                ) AS duplicate_rank,
                TRIM(COALESCE(fee.fee_type_name, '')) LIKE '损失承担%' AS is_loss_bearing
              FROM unit_revenue_fee_detail fee
              WHERE fee.revenue_date BETWEEN :start_date AND :end_date
                AND LEFT(TRIM(fee.source_group_code), 3) = :store_code
            ),
            fees_by_group AS (
              SELECT group_code_norm, SUM(amount) AS amount
              FROM fee_source
              WHERE NOT is_loss_bearing OR duplicate_rank = 1
              GROUP BY group_code_norm
            ),
            extras_by_group AS (
              SELECT UPPER(TRIM(source_group_code)) AS group_code_norm, SUM(amount) AS amount
              FROM revenue_extra_receipts
              WHERE revenue_date BETWEEN :start_date AND :end_date
                AND (
                  store_id = :store_id
                  OR LEFT(TRIM(source_group_code), 3) = :store_code
                )
                AND status = 'CONFIRMED'
                AND source_type = 'NC6051'
              GROUP BY UPPER(TRIM(source_group_code))
            ),
            group_revenue AS (
              SELECT group_code_norm, SUM(amount) AS amount
              FROM (
                SELECT group_code_norm, amount FROM sales_by_group
                UNION ALL
                SELECT group_code_norm, amount FROM fees_by_group
                UNION ALL
                SELECT group_code_norm, amount FROM extras_by_group
              ) source
              GROUP BY group_code_norm
            )
            SELECT
              a.manager_user_id,
              COALESCE(SUM(COALESCE(r.amount, 0)), 0) / 10000.0 AS area_actual
            FROM category_manager_brand_assignments a
            LEFT JOIN group_revenue r
              ON r.group_code_norm = UPPER(TRIM(a.group_code))
            WHERE a.store_id = :store_id
              AND a.is_active
              AND UPPER(TRIM(a.group_code)) = ANY(:allowed_codes)
            GROUP BY a.manager_user_id
            """
        ),
        {
            "store_id": store_id,
            "store_code": store_code,
            "start_date": start_date,
            "end_date": end_date,
            "allowed_codes": allowed_codes,
        },
    ).mappings().all()
    return {int(row["manager_user_id"]): Decimal(row["area_actual"] or 0) for row in rows}


def _key_brand_actuals(
    db: Session,
    *,
    store_id: int,
    store_code: str,
    period_month: date,
    start_date: date,
    end_date: date,
    allowed_codes: list[str],
) -> dict[int, dict[str, Decimal | int]]:
    if not allowed_codes:
        return {}
    rows = db.execute(
        text(
            """
            WITH key_brands AS (
              SELECT
                COALESCE(t.manager_user_id, a.manager_user_id) AS manager_user_id,
                UPPER(TRIM(a.group_code)) AS group_code_norm,
                COALESCE(t.sales_target, 0) AS sales_target
              FROM category_manager_brand_assignments a
              JOIN manaframe_key_brand kb
                ON UPPER(TRIM(kb.mfcode)) = UPPER(TRIM(a.group_code))
               AND kb.is_key_brand
              LEFT JOIN category_key_brand_targets t
                ON t.store_id = a.store_id
               AND UPPER(TRIM(t.group_code)) = UPPER(TRIM(a.group_code))
               AND t.period_month = :period_month
               AND t.is_active
              WHERE a.store_id = :store_id
                AND a.is_active
                AND UPPER(TRIM(a.group_code)) = ANY(:allowed_codes)
            ),
            sales_by_group AS (
              SELECT
                UPPER(TRIM(sglmfid)) AS group_code_norm,
                COALESCE(SUM(COALESCE(sglsjje, 0)), 0) / 10000.0 AS sales_actual
              FROM salegoodslist
              WHERE sglhsrq BETWEEN :start_date AND :end_date
                AND sglmarket = :store_code
              GROUP BY UPPER(TRIM(sglmfid))
            )
            SELECT
              t.manager_user_id,
              SUM(t.sales_target) AS sales_target,
              SUM(COALESCE(s.sales_actual, 0)) AS sales_actual,
              COUNT(*)::integer AS brand_count
            FROM key_brands t
            LEFT JOIN sales_by_group s ON s.group_code_norm = t.group_code_norm
            GROUP BY t.manager_user_id
            """
        ),
        {
            "store_id": store_id,
            "store_code": store_code,
            "period_month": period_month,
            "start_date": start_date,
            "end_date": end_date,
            "allowed_codes": allowed_codes,
        },
    ).mappings().all()
    return {
        int(row["manager_user_id"]): {
            "target": Decimal(row["sales_target"] or 0),
            "actual": Decimal(row["sales_actual"] or 0),
            "brand_count": int(row["brand_count"] or 0),
        }
        for row in rows
    }


@router.get("/scorecard")
async def category_performance_scorecard(
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, VIEW_PERMISSION)
    _require_tables(db)
    window = _period_window(period)
    store = _store_row(db, store_id)
    allowed_groups = _allowed_groups(db, current_user, store)
    allowed_codes = [str(item["group_code"]).strip().upper() for item in allowed_groups]
    latest_sales_date = db.execute(
        text(
            """
            SELECT MAX(sglhsrq)
            FROM salegoodslist
            WHERE sglmarket = :store_code
              AND sglhsrq BETWEEN :start_date AND :end_date
            """
        ),
        {
            "store_code": str(store["store_code"]).strip(),
            "start_date": window.start_date,
            "end_date": window.end_date,
        },
    ).scalar()
    effective_end_date = min(window.end_date, latest_sales_date) if latest_sales_date else window.end_date
    area_actuals = _manager_actuals(
        db,
        store_id=store_id,
        store_code=str(store["store_code"]).strip(),
        start_date=window.start_date,
        end_date=effective_end_date,
        allowed_codes=allowed_codes,
    )
    key_actuals = _key_brand_actuals(
        db,
        store_id=store_id,
        store_code=str(store["store_code"]).strip(),
        period_month=window.period_month,
        start_date=window.start_date,
        end_date=effective_end_date,
        allowed_codes=allowed_codes,
    )

    visible_manager_ids = set(area_actuals) | set(key_actuals)
    targets = []
    if visible_manager_ids:
        targets = db.execute(
            text(
                """
                SELECT
                  u.user_id AS manager_user_id,
                  COALESCE(t.manager_name, u.real_name) AS manager_name,
                  COALESCE(t.area_revenue_target, 0) AS area_revenue_target,
                  COALESCE(t.area_weight, 40) AS area_weight,
                  COALESCE(t.key_brand_weight, 40) AS key_brand_weight,
                  COALESCE(t.self_weight, 20) AS self_weight,
                  t.self_score,
                  t.assessment_content
                FROM users u
                LEFT JOIN category_manager_performance_targets t
                  ON t.manager_user_id = u.user_id
                 AND t.store_id = :store_id
                 AND t.period_month = :period_month
                WHERE u.user_id = ANY(:visible_manager_ids)
                ORDER BY COALESCE(t.manager_name, u.real_name), u.user_id
                """
            ),
            {
                "store_id": store_id,
                "period_month": window.period_month,
                "visible_manager_ids": sorted(visible_manager_ids),
            },
        ).mappings().all()
    items = []
    for target in targets:
        manager_user_id = int(target["manager_user_id"])
        key = key_actuals.get(
            manager_user_id,
            {"target": Decimal("0"), "actual": Decimal("0"), "brand_count": 0},
        )
        score = build_score_row(
            area_target=target["area_revenue_target"],
            area_actual=area_actuals.get(manager_user_id, Decimal("0")),
            area_weight=target["area_weight"],
            key_target=key["target"],
            key_actual=key["actual"],
            key_weight=target["key_brand_weight"],
            self_score=target["self_score"],
        )
        items.append(
            {
                "manager_user_id": manager_user_id,
                "manager_name": target["manager_name"],
                "assessment_content": target["assessment_content"],
                "area_weight": float(target["area_weight"]),
                "key_brand_weight": float(target["key_brand_weight"]),
                "self_weight": float(target["self_weight"]),
                "key_brand_count": int(key["brand_count"]),
                **{name: _json_value(value) for name, value in score.items()},
            }
        )

    latest_fee_date = db.execute(
        text(
            """
            SELECT MAX(fee.revenue_date)
            FROM unit_revenue_fee_detail fee
            WHERE fee.revenue_date BETWEEN :start_date AND :end_date
              AND LEFT(TRIM(fee.source_group_code), 3) = :store_code
            """
        ),
        {
            "store_code": str(store["store_code"]).strip(),
            "start_date": window.start_date,
            "end_date": window.end_date,
        },
    ).scalar()
    return {
        "store": dict(store),
        "period": period,
        "period_start": window.start_date.isoformat(),
        "period_end": window.end_date.isoformat(),
        "effective_end_date": effective_end_date.isoformat(),
        "latest_sales_date": latest_sales_date.isoformat() if latest_sales_date else None,
        "latest_fee_date": latest_fee_date.isoformat() if latest_fee_date else None,
        "definitions": {
            "area_actual": "分管品牌柜组的不含税毛利、富基收费及NC6051非富基收费合计，单位万元",
            "key_brand_actual": "重点品牌柜组售价金额合计，单位万元",
            "score": "达成率×对应权重；不封顶。总分达到100/90/80分时系数分别为1.2/1.0/0.9，否则0.8",
        },
        "items": items,
    }


def _normalized_name(value: object) -> str:
    return "".join(str(value or "").strip().split()).upper()


@router.post("/import-workbook")
async def import_category_performance_workbook(
    store_id: int = Query(...),
    period: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    require_permission(db, current_user, MANAGE_PERMISSION)
    _require_tables(db)
    window = _period_window(period)
    store = _store_row(db, store_id)
    allowed_groups = _allowed_groups(db, current_user, store)
    group_map: dict[str, list[dict[str, Any]]] = {}
    for group in allowed_groups:
        group_map.setdefault(_normalized_name(group["group_name"]), []).append(group)
    manager_rows = db.execute(
        text(
            """
            SELECT user_id, real_name
            FROM users
            WHERE is_active
              AND COALESCE(status, 'ACTIVE') = 'ACTIVE'
              AND NULLIF(TRIM(real_name), '') IS NOT NULL
            """
        )
    ).mappings().all()
    manager_map = {_normalized_name(row["real_name"]): row for row in manager_rows}

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="服务端缺少Excel读取组件") from exc

    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Excel文件不能超过20MB")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取Excel文件") from exc

    required_sheets = {"Sheet6", "重点品牌销售目标达成", "实际得分"}
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        raise HTTPException(status_code=422, detail=f"Excel缺少工作表：{'、'.join(missing_sheets)}")

    imported = {"brand_assignments": 0, "key_brand_targets": 0, "manager_targets": 0}
    unmatched: list[dict[str, str]] = []
    department_hints: dict[tuple[str, str], set[str]] = {}

    if "分管区域收益目标达成" in workbook.sheetnames:
        for row in workbook["分管区域收益目标达成"].iter_rows(min_row=2, values_only=True):
            department_name = row[0] if len(row) > 0 else None
            brand_name = row[1] if len(row) > 1 else None
            manager_name = row[6] if len(row) > 6 else None
            if (
                department_name
                and brand_name
                and manager_name
                and str(manager_name).strip() != "#N/A"
            ):
                department_hints.setdefault(
                    (_normalized_name(brand_name), _normalized_name(manager_name)),
                    set(),
                ).add(str(department_name).strip())

    def resolve_manager(name: object):
        return manager_map.get(_normalized_name(name))

    def resolve_group(
        name: object,
        department_name: object = None,
        manager_name: object = None,
    ):
        candidates = group_map.get(_normalized_name(name), [])
        department_norm = _normalized_name(department_name)
        if not department_norm and manager_name and len(candidates) > 1:
            hinted_departments = department_hints.get(
                (_normalized_name(name), _normalized_name(manager_name)),
                set(),
            )
            if len(hinted_departments) == 1:
                department_norm = _normalized_name(next(iter(hinted_departments)))
        if department_norm and len(candidates) > 1:
            candidates = [
                candidate
                for candidate in candidates
                if _normalized_name(candidate.get("department_name")) == department_norm
            ]
        return candidates[0] if len(candidates) == 1 else None

    for brand_name, manager_name, *_ in workbook["Sheet6"].iter_rows(values_only=True):
        if not brand_name or not manager_name:
            continue
        group = resolve_group(brand_name, manager_name=manager_name)
        manager = resolve_manager(manager_name)
        if not group or not manager:
            unmatched.append({"sheet": "Sheet6", "brand": str(brand_name), "manager": str(manager_name)})
            continue
        db.execute(
            text(
                """
                INSERT INTO category_manager_brand_assignments (
                  store_id, group_code, manager_user_id, manager_name, is_active,
                  created_by, updated_by
                )
                VALUES (:store_id, :group_code, :manager_user_id, :manager_name, TRUE, :user_id, :user_id)
                ON CONFLICT (store_id, group_code) DO UPDATE SET
                  manager_user_id = EXCLUDED.manager_user_id,
                  manager_name = EXCLUDED.manager_name,
                  is_active = TRUE,
                  updated_by = EXCLUDED.updated_by,
                  updated_at = NOW()
                """
            ),
            {
                "store_id": store_id,
                "group_code": group["group_code"],
                "manager_user_id": manager["user_id"],
                "manager_name": manager["real_name"],
                "user_id": current_user.user_id,
            },
        )
        _ensure_supplier_payment_role(
            db,
            manager_user_id=int(manager["user_id"]),
            store_row=store,
        )
        imported["brand_assignments"] += 1

    key_sheet = workbook["重点品牌销售目标达成"]
    for row in key_sheet.iter_rows(min_row=2, values_only=True):
        brand_name = row[4] if len(row) > 4 else None
        department_name = row[1] if len(row) > 1 else None
        marker = row[5] if len(row) > 5 else None
        manager_name = row[6] if len(row) > 6 else None
        target = row[7] if len(row) > 7 else None
        if "重点品牌" not in str(marker or "") or not brand_name or not manager_name:
            continue
        group = resolve_group(brand_name, department_name)
        manager = resolve_manager(manager_name)
        if not group or not manager:
            unmatched.append({"sheet": "重点品牌销售目标达成", "brand": str(brand_name), "manager": str(manager_name)})
            continue
        sales_target = Decimal(str(target or 0))
        db.execute(
            text(
                """
                INSERT INTO category_key_brand_targets (
                  store_id, group_code, period_month, manager_user_id, manager_name,
                  sales_target, is_active, created_by, updated_by
                )
                VALUES (
                  :store_id, :group_code, :period_month, :manager_user_id, :manager_name,
                  :sales_target, TRUE, :user_id, :user_id
                )
                ON CONFLICT (store_id, group_code, period_month) DO UPDATE SET
                  manager_user_id = EXCLUDED.manager_user_id,
                  manager_name = EXCLUDED.manager_name,
                  sales_target = EXCLUDED.sales_target,
                  is_active = TRUE,
                  updated_by = EXCLUDED.updated_by,
                  updated_at = NOW()
                """
            ),
            {
                "store_id": store_id,
                "group_code": group["group_code"],
                "period_month": window.period_month,
                "manager_user_id": manager["user_id"],
                "manager_name": manager["real_name"],
                "sales_target": sales_target,
                "user_id": current_user.user_id,
            },
        )
        db.execute(
            text(
                """
                INSERT INTO manaframe_key_brand (mfcode, is_key_brand, updated_at, updated_by)
                VALUES (:group_code, TRUE, NOW(), :updated_by)
                ON CONFLICT (mfcode) DO UPDATE SET
                  is_key_brand = TRUE,
                  updated_at = NOW(),
                  updated_by = EXCLUDED.updated_by
                """
            ),
            {"group_code": group["group_code"], "updated_by": str(current_user.user_id)},
        )
        imported["key_brand_targets"] += 1

    score_sheet = workbook["实际得分"]
    for row in score_sheet.iter_rows(values_only=True):
        department_name = row[0] if len(row) > 0 else None
        manager_name = row[1] if len(row) > 1 else None
        area_target = row[2] if len(row) > 2 else None
        assessment_content = row[10] if len(row) > 10 else None
        self_score = row[11] if len(row) > 11 else None
        if not department_name or not manager_name or not isinstance(area_target, (int, float, Decimal)):
            continue
        manager = resolve_manager(manager_name)
        if not manager:
            unmatched.append({"sheet": "实际得分", "brand": str(department_name), "manager": str(manager_name)})
            continue
        department_text = str(department_name)
        if "特业" in department_text:
            weights = (Decimal("60"), Decimal("20"), Decimal("20"))
        elif "儿童游乐园" in department_text:
            weights = (Decimal("80"), Decimal("0"), Decimal("20"))
        else:
            weights = (Decimal("40"), Decimal("40"), Decimal("20"))
        db.execute(
            text(
                """
                INSERT INTO category_manager_performance_targets (
                  store_id, period_month, manager_user_id, manager_name,
                  area_revenue_target, area_weight, key_brand_weight, self_weight,
                  self_score, assessment_content, created_by, updated_by
                )
                VALUES (
                  :store_id, :period_month, :manager_user_id, :manager_name,
                  :area_target, :area_weight, :key_weight, :self_weight,
                  :self_score, :assessment_content, :user_id, :user_id
                )
                ON CONFLICT (store_id, manager_user_id, period_month) DO UPDATE SET
                  manager_name = EXCLUDED.manager_name,
                  area_revenue_target = EXCLUDED.area_revenue_target,
                  area_weight = EXCLUDED.area_weight,
                  key_brand_weight = EXCLUDED.key_brand_weight,
                  self_weight = EXCLUDED.self_weight,
                  self_score = EXCLUDED.self_score,
                  assessment_content = EXCLUDED.assessment_content,
                  updated_by = EXCLUDED.updated_by,
                  updated_at = NOW()
                """
            ),
            {
                "store_id": store_id,
                "period_month": window.period_month,
                "manager_user_id": manager["user_id"],
                "manager_name": manager["real_name"],
                "area_target": Decimal(str(area_target)),
                "area_weight": weights[0],
                "key_weight": weights[1],
                "self_weight": weights[2],
                "self_score": Decimal(str(self_score)) if isinstance(self_score, (int, float, Decimal)) else None,
                "assessment_content": str(assessment_content).strip() if assessment_content else None,
                "user_id": current_user.user_id,
            },
        )
        imported["manager_targets"] += 1

    _operation_log(
        db,
        current_user,
        action_code="category_performance.workbook.import",
        target_id=f"{store_id}:{period}",
        detail={"filename": file.filename, "imported": imported, "unmatched_count": len(unmatched)},
    )
    db.commit()
    return {
        "filename": file.filename,
        "store_id": store_id,
        "period": period,
        "imported": imported,
        "unmatched_count": len(unmatched),
        "unmatched": unmatched[:100],
    }
