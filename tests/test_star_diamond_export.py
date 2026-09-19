from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import sys

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'python_app'))
from routers import activity_analysis as a
from services.activity_analysis.star_diamond_excel import build_star_diamond_workbook


def payload(**kwargs):
    return a.StarDiamondExportRequest(start_date=date(2026,1,1),end_date=date(2026,9,12),**kwargs)


def mock_data(monkeypatch):
    members=[{'member_no':str(i).zfill(6),'customer_name':'=1+1','telephone':'00123456789'} for i in range(305)]
    sales=[{'member_no':m['member_no'],'ticket_count':2,'overview_ticket_count':2,'sales_amount':Decimal('12000.25'),'overview_net_profit':Decimal('-1.25')} for m in members]
    tickets=[{'member_no':m['member_no'],'billno':Decimal('12345678901234567890')+i,'sale_time':'2026-09-12 12:00:00'} for i,m in enumerate(members)]
    monkeypatch.setattr(a,'_load_star_diamond_members',lambda db,keyword=None: members[:1] if keyword else members)
    monkeypatch.setattr(a,'_load_star_diamond_tickets',lambda *args: tickets)
    monkeypatch.setattr(a,'_load_star_diamond_member_sales',lambda *args,**kwargs:sales)
    monkeypatch.setattr(a,'_load_star_diamond_category_sales',lambda *args:[])
    monkeypatch.setattr(a,'_load_star_diamond_trail_rows',lambda db,rows:[{**r,'sales_amount':Decimal('-15.50')} for r in rows])


def test_export_is_not_limited_to_preview_and_preserves_excel_types(monkeypatch):
    mock_data(monkeypatch)
    report=a._load_star_diamond_export_data(object(),payload(ai_report='测试完整方案\n=1+1'))
    assert len(report['members'])==len(report['trails'])==305
    wb=load_workbook(BytesIO(build_star_diamond_workbook(report)))
    assert wb.sheetnames==['经营汇总','服务分层','偏好品类Top10','会员清单','购物轨迹','AI整体方案']
    assert wb['会员清单'].max_row==310
    assert wb['会员清单']['A6'].value=='000000'
    assert wb['会员清单']['C6'].value=='00123456789'
    assert wb['会员清单']['B6'].value=='=1+1'
    assert wb['会员清单']['B6'].data_type=='s'
    assert wb['会员清单']['G6'].data_type=='n'
    assert wb['购物轨迹']['B6'].value=='12345678901234567890'
    assert wb['购物轨迹']['G6'].value==-15.5
    assert wb['AI整体方案']['B7'].data_type=='s'
    assert wb['经营汇总']['B16'].value==1
    assert wb['经营汇总']['B16'].number_format=='0.0%'
    assert wb['会员清单'].freeze_panes=='A6'


def test_export_filters_are_independent_and_summary_stays_full_period(monkeypatch):
    mock_data(monkeypatch)
    report=a._load_star_diamond_export_data(object(),payload(keyword='test',segment='高价值维护',member_no='000002'))
    assert [r['member_no'] for r in report['members']]==['000000']
    assert [r['member_no'] for r in report['trails']]==['000002']
    assert report['overview']['summary']['star_member_count']==305
    assert 'AI整体方案' not in load_workbook(BytesIO(build_star_diamond_workbook(report))).sheetnames


def test_export_rejects_missing_permission_before_loading_data(monkeypatch):
    def deny(*args):
        raise HTTPException(status_code=403,detail='无功能权限')
    monkeypatch.setattr(a,'require_permission',deny)
    monkeypatch.setattr(a,'_load_star_diamond_export_data',lambda *args:pytest.fail('must not load data'))
    with pytest.raises(HTTPException) as e:
        a.star_diamond_export(payload(),object(),object())
    assert e.value.status_code==403


def test_export_rejects_out_of_scope_store_before_loading_data(monkeypatch):
    monkeypatch.setattr(a,'require_permission',lambda *args:None)
    def deny(*args):
        raise HTTPException(status_code=403,detail='无门店权限')
    monkeypatch.setattr(a,'_require_center_store_scope',deny)
    with pytest.raises(HTTPException) as e:
        a.star_diamond_export(payload(),object(),object())
    assert e.value.status_code==403


def test_empty_export_has_headers_and_valid_numeric_summary(monkeypatch):
    monkeypatch.setattr(a,'_load_star_diamond_members',lambda *args:[])
    monkeypatch.setattr(a,'_load_star_diamond_tickets',lambda *args:[])
    report=a._load_star_diamond_export_data(object(),payload())
    wb=load_workbook(BytesIO(build_star_diamond_workbook(report)))
    assert wb['会员清单'].max_row==5
    assert wb['购物轨迹'].max_row==5
    assert wb['经营汇总']['B16'].value==0


def test_export_download_response_and_invalid_date(monkeypatch):
    mock_data(monkeypatch)
    checks=[]
    monkeypatch.setattr(a,'require_permission',lambda db,user,permission: checks.append(permission))
    monkeypatch.setattr(a,'_require_center_store_scope',lambda *args: checks.append('center_scope'))
    monkeypatch.setattr(a,'_ensure_required_tables',lambda *args:None)
    monkeypatch.setattr(a,'_execute_star_diamond_query',lambda db,loader:loader())
    response=a.star_diamond_export(payload(),object(),object())
    assert checks==[a.STAR_DIAMOND_ANALYSIS_PERMISSION,'center_scope']
    assert response.status_code==200
    assert response.body.startswith(b'PK')
    assert 'filename*=UTF-8' in response.headers['content-disposition']
    assert response.headers['cache-control']=='no-store'
    assert len(load_workbook(BytesIO(response.body)).sheetnames)==5
    bad=a.StarDiamondExportRequest(start_date='2026-09-12',end_date='2026-01-01')
    with pytest.raises(HTTPException) as error:
        a.star_diamond_export(bad,object(),object())
    assert error.value.status_code==400
