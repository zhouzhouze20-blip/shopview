from datetime import date
import inspect
from pathlib import Path

from openpyxl import load_workbook
import pytest

from python_app.services.hy0001_excel import build_hy0001_workbook_file
from python_app.services.hy0001_report import (
    assemble_report,
    build_report_query,
    date_window,
)
from python_app.routers import sales


def _raw_rows():
    base = {
        "store_code": "001",
        "store_name": "中心店",
        "department_code": "00101",
        "department_name": "中心一部",
        "group_code": "0010101",
        "group_name": "重点品牌厅",
        "manager_name": "张三",
    }
    values = {
        "03": (100, 80, 4, 5),
        "04": (200, 100, 6, 4),
        "02": (50, 0, 2, 0),
        "01": (25, 50, 1, 2),
    }
    labels = {
        "03": "黑金卡会员",
        "04": "黑钻卡会员",
        "02": "金星卡会员",
        "01": "银星卡会员",
    }
    return [
        {
            **base,
            "level_code": code,
            "level_label": labels[code],
            "current_sales": current_sales,
            "prior_sales": prior_sales,
            "current_buyers": current_buyers,
            "prior_buyers": prior_buyers,
        }
        for code, (current_sales, prior_sales, current_buyers, prior_buyers) in values.items()
    ]


def test_date_window_uses_selected_dates_and_prior_year():
    assert date_window(date(2024, 2, 29), date(2024, 3, 5)) == (
        date(2024, 2, 29),
        date(2024, 3, 5),
        date(2023, 2, 28),
        date(2023, 3, 5),
    )


def test_query_keeps_key_brand_manager_and_member_level_lineage():
    sql, params = build_report_query(
        start_date=date(2026, 3, 3),
        end_date=date(2026, 3, 18),
        scope_filter_sql="AND st.store_id::text = :scope_store",
        scope_params={"scope_store": "8"},
        selected_store="001",
        selected_department="00101",
    )

    assert "manaframe_key_brand" in sql
    assert "category_manager_brand_assignments" in sql
    assert "h.custtype" in sql
    assert "h.hykh" in sql
    assert "BOOL_OR(sales_revenue > 0)" in sql
    assert "AND st.store_id::text = :scope_store" in sql
    assert params["selected_store"] == "001"
    assert params["selected_department"] == "00101"
    assert params["current_start"] == date(2026, 3, 3)
    assert params["current_end"] == date(2026, 3, 18)
    assert params["prior_start"] == date(2025, 3, 3)
    assert params["prior_end"] == date(2025, 3, 18)


def test_report_adds_formal_manager_column_and_safe_yoy():
    report = assemble_report(
        _raw_rows(),
        start_date=date(2026, 3, 3),
        end_date=date(2026, 3, 18),
        selected_store="001",
        selected_department=None,
    )

    assert report["rows"][0]["manager_name"] == "张三"
    assert report["rows"][0]["levels"][0]["sales_yoy"] == 0.25
    assert report["rows"][0]["levels"][2]["sales_yoy"] is None
    assert report["manager_summary"][0]["current_premium_buyers"] == 10
    assert report["manager_summary"][0]["prior_premium_buyers"] == 9
    assert report["manager_summary"][0]["premium_buyer_yoy"] == pytest.approx(1 / 9)
    assert report["manager_summary"][0]["current_premium_sales"] == 300
    assert report["manager_summary"][0]["prior_premium_sales"] == 180


def test_export_matches_reference_structure_and_uses_formula_summary():
    report = assemble_report(
        _raw_rows(),
        start_date=date(2026, 3, 3),
        end_date=date(2026, 3, 18),
        selected_store="001",
        selected_department=None,
    )
    report["scope_description"] = "中心店"
    export_file = build_hy0001_workbook_file(report)
    workbook = load_workbook(export_file, data_only=False)

    assert workbook.sheetnames == ["明细", "主管汇总", "数据口径"]
    detail = workbook["明细"]
    assert detail["A1"].value == "品类主管"
    assert detail["C1"].value == "重点品牌"
    assert detail["D1"].value == "销售（元）"
    assert detail["P1"].value == "消费人数"
    assert detail["A4"].value == "张三"
    assert detail["D4"].value == 100
    assert detail["F4"].font.color.rgb == "FFC00000"
    assert detail["O4"].font.color.rgb == "FF008000"
    assert detail["R4"].font.color.rgb == "FF008000"
    assert len(detail.conditional_formatting) == 8
    summary = workbook["主管汇总"]
    assert summary["C1"].value == "本期黑金+黑钻人数"
    assert "$P$4:$P$4" in summary["C2"].value
    assert "$S$4:$S$4" in summary["C2"].value
    assert summary["F2"].value.startswith("=SUMIF('明细'!")
    assert summary["E2"].font.color.rgb == "FFC00000"
    assert summary["H2"].font.color.rgb == "FFC00000"
    assert len(summary.conditional_formatting) == 2
    export_file.close()


def test_permission_migration_is_chained_after_latest_report_migration():
    migration = Path(
        "python_app/alembic/versions/t8b9c0d1e2f3_add_hy0001_permission.py"
    ).read_text(encoding="utf-8")
    assert 'down_revision: Union[str, Sequence[str], None] = "s7a8b9c0d1e2"' in migration
    assert "sales.hy0001.view" in migration
    assert "sales.category_performance.view" in migration
    assert "sales.brand_member_analysis.view" in migration


def test_report_and_export_share_the_independent_permission_scoped_loader():
    loader_source = inspect.getsource(sales._load_hy0001_for_request)
    report_source = inspect.getsource(sales.hy0001_report)
    export_source = inspect.getsource(sales.hy0001_export)

    assert sales.HY0001_PERMISSION == "sales.hy0001.view"
    assert "require_permission(db, current_user, HY0001_PERMISSION)" in loader_source
    assert "_business_scope_filter_sql" in loader_source
    assert "_load_hy0001_for_request" in report_source
    assert "_load_hy0001_for_request" in export_source
