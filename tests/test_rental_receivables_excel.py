from openpyxl import load_workbook

from python_app.services.rental_receivables_excel import (
    build_rental_receivable_expense_workbook_file,
)


def test_expense_workbook_contains_signed_fee_details_and_scope_notes():
    report = {
        "items": [
            {
                "store_code": "602",
                "store_name": "常州百货大楼",
                "department_code": "6020101",
                "department_name": "营运一部",
                "group_code": "6020101173",
                "group_name": "老凤祥厅",
                "supplier_id": "30643",
                "supplier_name": "常州泰鼎珠宝有限公司",
                "contract_no": "30643001",
                "bill_no": "260812602Z0000000006",
                "settle_from": "2026-06-30",
                "settle_to": "2026-07-31",
                "bill_receivable_amount": 1376.2,
                "row_no": 2,
                "status": "Y",
                "item_code": "00-0303",
                "item_name": "银联聚合支付",
                "period_from": "2026-07-01",
                "period_to": "2026-07-31",
                "finance_month": "202607",
                "amount": -42196.24,
                "checked_amount": 0,
                "paid_amount": 0,
                "deducted_amount": 0,
                "balance_amount": -42196.24,
                "receivable_component": 0,
                "is_receivable_line": False,
                "is_sales_refund": True,
                "adjustment_amount": 0,
                "sales_reference_amount": 0,
                "tax_rate": 0,
                "no_tax_amount": 0,
                "memo": "=unsafe",
                "calculation_source": "CHARGELIST",
                "is_advance": "N",
                "area": 88,
                "rental_area": 88,
            }
        ],
        "detail_count": 1,
        "bill_count": 1,
        "filters": {
            "settle_from": "2026-01-01",
            "settle_to": "2026-08-17",
            "mkt": "602",
            "department_code": "6020101",
            "group_prefix": None,
            "keyword": None,
        },
        "source": {"name": "本地 ODS（PAPI 总部库同步）"},
        "source_loaded_at": "2026-08-17T03:31:25+08:00",
        "generated_at": "2026-08-17T16:00:00+08:00",
        "scope_note": "仅导出当前账号权限及当前查询条件内的费用行",
    }

    output = build_rental_receivable_expense_workbook_file(report)
    try:
        workbook = load_workbook(output, read_only=True, data_only=False)
        assert workbook.sheetnames == ["费用明细", "导出说明"]
        sheet = workbook["费用明细"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        assert row[headers.index("项目编码")] == "00-0303"
        assert row[headers.index("明细余额")] == -42196.24
        assert row[headers.index("应收影响")] == 0
        assert row[headers.index("是否计入应收")] == "否"
        assert row[headers.index("是否销售返款")] == "是"
        assert row[headers.index("备注")] == "'=unsafe"
        notes = [values for values in workbook["导出说明"].iter_rows(values_only=True)]
        assert ("门店筛选", "602") in notes
    finally:
        output.close()
