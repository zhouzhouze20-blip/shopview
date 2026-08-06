from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from python_app.routers.authz import CORE_PERMISSION_DEFINITIONS
from python_app.services.od0002_report import TrustedScopeSql
from python_app.services.store_other_business_income_excel import (
    build_store_other_business_income_workbook,
)
from python_app.services.store_other_business_income_report import (
    TARGET_SUBJECT_CODES,
    build_store_other_business_income_query,
    normalize_store_other_business_income_rows,
)


def source_row(**overrides):
    row = {
        "store_id": "3",
        "store_code": "603",
        "store_name": "新世纪",
        "department_code": "3132",
        "department_name": "新世纪二部",
        "account_year": 2026,
        "account_period": 1,
        "subject_code": "605110",
        "debit_amount": 0,
        "credit_amount": 100_000,
        "net_income": 100_000,
    }
    row.update(overrides)
    return row


def test_query_matches_finance_scope_and_comparison_period():
    sql, params = build_store_other_business_income_query(
        2026,
        7,
        TrustedScopeSql(" AND st.store_id::text = ANY(:allowed_stores)"),
        {"allowed_stores": ["3", "4"]},
        selected_store="603",
    )
    compact = " ".join(sql.lower().split())

    assert "bh_dw_gl_detail_fact2" in compact
    assert "f.account_year = any(:financial_years)" in compact
    assert "f.account_period::int = any(:periods)" in compact
    assert "when f.pk_corp = '1018' and f.valuecode like '22%' then '602'" in compact
    assert "trim(coalesce(f.valuecode, '')) not in ('210109', '220109', '310109', '330109')" in compact
    assert "sum(credit_amount - debit_amount) as net_income" in compact
    assert "st.store_id::text = any(:allowed_stores)" in compact
    assert params["financial_years"] == ["2025", "2026"]
    assert params["periods"] == list(range(1, 8))
    assert params["subject_codes"] == list(TARGET_SUBJECT_CODES)
    assert params["selected_store"] == "603"


def test_normalization_adds_category_and_store_totals_with_yoy():
    result = normalize_store_other_business_income_rows(
        [
            source_row(),
            source_row(account_year=2025, net_income=80_000, credit_amount=80_000),
            source_row(subject_code="605112", net_income=20_000, credit_amount=20_000),
            source_row(
                store_id="4",
                store_code="604",
                store_name="半山",
                department_code="3334",
                department_name="半山非图部",
                net_income=10_000,
                credit_amount=10_000,
            ),
        ],
        financial_year=2026,
        end_period=1,
    )

    store_totals = [row for row in result["store_rows"] if row["row_type"] == "grand_total"]
    assert [row["store_name"] for row in store_totals] == ["半山", "新世纪"]
    assert store_totals[0]["total"]["current"] == 10_000
    assert store_totals[1]["total"]["current"] == 120_000
    assert store_totals[1]["total"]["prior"] == 80_000
    assert store_totals[1]["total"]["difference"] == 40_000
    assert store_totals[1]["total"]["rate"] == pytest.approx(0.5)
    assert result["summary"]["current"] == 130_000
    assert result["quality"] == {
        "source_row_count": 4,
        "store_count": 2,
        "department_count": 2,
    }


def test_workbook_has_reference_store_and_department_layout():
    normalized = normalize_store_other_business_income_rows(
        [source_row(), source_row(account_year=2025, net_income=80_000, credit_amount=80_000)],
        financial_year=2026,
        end_period=1,
    )
    report = {
        "financial_year": 2026,
        "prior_year": 2025,
        "periods": [1],
        **normalized,
    }
    workbook = load_workbook(BytesIO(build_store_other_business_income_workbook(report)))

    assert workbook.sheetnames == ["门店", "部门"]
    store = workbook["门店"]
    department = workbook["部门"]
    assert store["A1"].value.startswith("本期年份：2026年")
    assert store["D1"].value.startswith("同期年份：2025年")
    assert store["A3"].value == "门店"
    assert department["A3"].value == "部门"
    assert store["E3"].value == "01"
    assert store["E4"].value == "本期"
    assert store["F4"].value == "同期"
    assert store.freeze_panes == "E5"
    assert store["A5"].value == "新世纪"
    assert store["B5"].value == "合同收费收入"
    assert store["C5"].value == "综合管理费"
    assert store["E5"].value == 10
    assert store["F5"].value == 8
    assert store["E5"].number_format == "0.00"
    assert store["F5"].number_format == "0.00"
    assert store["I5"].number_format == "0.00"
    assert store["J5"].number_format == "0.00%"


def test_permission_is_registered_and_migration_copies_report_access():
    permissions = {row[0] for row in CORE_PERMISSION_DEFINITIONS}
    assert "sales.store_other_business_income.view" in permissions
    migration = Path(
        "python_app/alembic/versions/"
        "z4b5c6d7e8f9_add_store_other_business_income_permission.py"
    ).read_text(encoding="utf-8")
    assert "sales.store_other_business_income.view" in migration
    assert "source.permission_code = 'sales.od0002.view'" in migration
