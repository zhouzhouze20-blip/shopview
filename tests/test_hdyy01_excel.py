import asyncio
from datetime import date
from io import BytesIO
from tempfile import SpooledTemporaryFile

import pytest
from openpyxl import load_workbook


START = date(2026, 7, 1)
END = date(2026, 7, 10)


def sample_report(*, empty: bool = False):
    rows = [] if empty else [
        {
            "store_name": "新世纪",
            "department_name": "中心三部",
            "group_code": "G01",
            "group_name": "A柜组",
            "area": 10.5,
            "floor_code": "02",
            "level1_code": "10",
            "level1_name": "服装",
            "level2_code": "1001",
            "level2_name": "女装",
            "grade_label": "A",
            "quantity": 2.125,
            "sales_amount": 100.0,
            "tax_cost": 60.0,
            "profit": 40.0,
            "ticket_count": 1,
            "average_ticket": 100.0,
            "member_sales": 80.0,
            "stored_card_sales": 30.0,
        },
        {
            "store_name": "新世纪",
            "department_name": "中心三部",
            "group_code": "G02",
            "group_name": "退货柜组",
            "area": 8.0,
            "floor_code": "02",
            "level1_code": "10",
            "level1_name": "服装",
            "level2_code": "1002",
            "level2_name": "男装",
            "grade_label": "B",
            "quantity": -0.5,
            "sales_amount": -20.0,
            "tax_cost": -12.0,
            "profit": -8.0,
            "ticket_count": 0,
            "average_ticket": None,
            "member_sales": -10.0,
            "stored_card_sales": -5.0,
        },
    ]
    total = {
        "quantity": 0.0 if empty else 1.625,
        "sales_amount": 0.0 if empty else 80.0,
        "tax_cost": 0.0 if empty else 48.0,
        "profit": 0.0 if empty else 32.0,
        "ticket_count": 0 if empty else 1,
        "average_ticket": None if empty else 80.0,
        "member_sales": 0.0 if empty else 70.0,
        "stored_card_sales": 0.0 if empty else 25.0,
    }
    return {
        "dates": {"start_date": START, "end_date": END},
        "selected_store": "603",
        "selected_department": "6030117",
        "rows": rows,
        "total": total,
        "scope_description": "当前用户权限范围：store=4；department=6030117",
    }


HEADERS = [
    "机构", "部门", "柜组编码", "柜组名称", "面积", "楼层", "一级编码", "一级名称",
    "二级编码", "二级名称", "等级", "数量", "销售收入", "含税成本", "毛利",
    "消费次数", "客单价", "会员销售", "储值卡销售",
]


def test_workbook_has_exact_sheets_layout_formats_signed_rows_total_and_notes():
    from python_app.services.hdyy01_excel import build_hdyy01_workbook
    from python_app.services.od0002_report import EXCLUDED_DEPARTMENT_CODES

    workbook = load_workbook(BytesIO(build_hdyy01_workbook(sample_report())))

    assert workbook.sheetnames == ["明细", "报表说明"]
    sheet = workbook["明细"]
    assert sheet["A1"].value == "HDYY01柜组经营分析表"
    assert "2026-07-01" in sheet["A2"].value
    assert "2026-07-10" in sheet["A2"].value
    assert [sheet.cell(4, column).value for column in range(1, 20)] == HEADERS
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref == "A4:S6"
    assert sheet["A4"].fill.fgColor.rgb.endswith("4472C4")
    assert sheet["A4"].font.bold
    assert sheet["A4"].font.color.rgb.endswith("FFFFFF")
    assert sheet["A4"].border.left.style == "thin"
    assert sheet.column_dimensions["A"].width > 0

    assert sheet["E5"].number_format == "0.00"
    assert sheet["L5"].number_format == "0.####"
    assert sheet["M5"].number_format == "0.00"
    assert sheet["P5"].number_format == "0"
    assert sheet["Q6"].value is None
    assert sheet["L6"].value == -0.5
    assert sheet["M6"].value == -20
    assert isinstance(sheet["M6"].value, (int, float))

    assert sheet["A7"].value == "合计"
    assert all(sheet.cell(7, column).value is None for column in range(2, 12))
    assert [sheet.cell(7, column).value for column in range(12, 20)] == [
        1.625, 80, 48, 32, 1, 80, 70, 25,
    ]
    assert sheet["A7"].font.bold
    assert sheet["A7"].fill.fill_type == "solid"

    notes = workbook["报表说明"]["B2"].value
    for phrase in (
        "sglhsrq", "sglxssr", "sgln13+sgln14-sglsupzk", "sgln2", "sglfcard",
        "salehead.hykh", "退货按带符号金额计入", "小票净销售额 > 0", "客单价 = 带符号销售额 / 正向消费次数",
        "报表期间：2026-07-01 至 2026-07-10", "排除租赁业务 sglwmid=5",
        "当前用户权限范围：store=4；department=6030117",
        "mfchr2", "mana_brand_hierarchy", "按编码层级关联", "未匹配",
    ):
        assert phrase in notes
    assert all(code in notes for code in EXCLUDED_DEPARTMENT_CODES)


def test_workbook_escapes_formula_like_text_but_keeps_negative_numbers_numeric():
    from python_app.services.hdyy01_excel import build_hdyy01_workbook

    report = sample_report()
    text_values = [
        "=1+1", "+SUM(A1:A2)", "-G01", "@柜组", "=10", "+02", "-10", "@服装",
        "=1001", "+女装", "-A",
    ]
    for key, value in zip(
        (
            "store_name", "department_name", "group_code", "group_name", "area",
            "floor_code", "level1_code", "level1_name", "level2_code", "level2_name", "grade_label",
        ),
        text_values,
    ):
        report["rows"][0][key] = value
    report["rows"][0]["area"] = -10.25
    report["rows"][0]["quantity"] = -2.5

    sheet = load_workbook(
        BytesIO(build_hdyy01_workbook(report)), data_only=False
    )["明细"]
    for column in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11):
        cell = sheet.cell(5, column)
        assert cell.data_type == "s"
        assert cell.value.startswith("'")
    assert sheet["E5"].value == -10.25
    assert sheet["L5"].value == -2.5
    assert isinstance(sheet["L5"].value, (int, float))


def test_empty_workbook_has_headers_total_notes_and_missing_classification_display():
    from python_app.services.hdyy01_excel import build_hdyy01_workbook

    workbook = load_workbook(BytesIO(build_hdyy01_workbook(sample_report(empty=True))))
    sheet = workbook["明细"]
    assert [sheet.cell(4, column).value for column in range(1, 20)] == HEADERS
    assert sheet["A5"].value == "合计"
    assert sheet.auto_filter.ref == "A4:S4"
    assert "sglxssr" in workbook["报表说明"]["B2"].value

    report = sample_report()
    for key in ("level1_code", "level1_name", "level2_code", "level2_name", "grade_label"):
        report["rows"][0][key] = None
    sheet = load_workbook(BytesIO(build_hdyy01_workbook(report)))["明细"]
    assert [sheet.cell(5, column).value for column in range(7, 12)] == ["未匹配"] * 5


def test_workbook_file_api_is_seeked_and_bytes_helper_closes_its_file(monkeypatch):
    from python_app.services import hdyy01_excel

    export_file = hdyy01_excel.build_hdyy01_workbook_file(sample_report())
    try:
        assert isinstance(export_file, SpooledTemporaryFile)
        assert export_file.tell() == 0
        assert load_workbook(export_file).sheetnames == ["明细", "报表说明"]
    finally:
        export_file.close()
    assert export_file.closed

    tracked_file = SpooledTemporaryFile()
    tracked_file.write(b"xlsx")
    tracked_file.seek(0)
    monkeypatch.setattr(hdyy01_excel, "build_hdyy01_workbook_file", lambda report: tracked_file)
    assert hdyy01_excel.build_hdyy01_workbook({}) == b"xlsx"
    assert tracked_file.closed


def test_workbook_builder_uses_write_only_mode_and_handles_thousands_of_rows():
    from python_app.services import hdyy01_excel

    contract_workbook = hdyy01_excel._build_hdyy01_workbook(
        sample_report(empty=True)
    )
    try:
        assert contract_workbook.write_only is True
    finally:
        contract_workbook.save(BytesIO())

    report = sample_report()
    template = report["rows"][0]
    report["rows"] = [
        {**template, "group_code": f"G{index:05d}"}
        for index in range(5_000)
    ]

    export_file = hdyy01_excel.build_hdyy01_workbook_file(report)
    workbook = load_workbook(export_file, read_only=True)
    try:
        detail = workbook["明细"]
        rows = detail.iter_rows(values_only=True)
        assert next(rows)[0] == "HDYY01柜组经营分析表"
        assert "2026-07-01" in next(rows)[0]
        assert next(rows)[0].startswith("筛选：")
        assert list(next(rows)) == HEADERS
        for _ in range(5_000):
            last = next(rows)
        assert last[2] == "G04999"
        assert next(rows)[0] == "合计"
        with pytest.raises(StopIteration):
            next(rows)
    finally:
        workbook.close()
        export_file.close()


def test_export_route_uses_shared_loader_once_filters_threadpool_and_closes_file(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    calls = {"load": 0}
    report = sample_report()

    def fake_load(start_date, end_date, store_id, db, user, department_id=None):
        calls["load"] += 1
        calls["arguments"] = (start_date, end_date, store_id, db, user, department_id)
        return report, DataScope(allow={"store": {"4"}, "department": {"6030117"}})

    monkeypatch.setattr(sales, "_load_hdyy01_for_request", fake_load)
    export_file = SpooledTemporaryFile()
    export_file.write(b"x" * (64 * 1024 + 3))
    export_file.seek(0)
    captured = {}

    def fake_builder(payload):
        captured["payload"] = payload
        return export_file

    monkeypatch.setattr(sales, "build_hdyy01_workbook_file", fake_builder)
    threadpool_calls = []

    async def fake_threadpool(function, *args, **kwargs):
        threadpool_calls.append((function, args, kwargs))
        return function(*args, **kwargs)

    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)

    db, user = object(), object()
    response = asyncio.run(
        sales.hdyy01_export(START, END, " 603 ", db, user, " 6030117 ")
    )

    assert calls["load"] == 1
    assert calls["arguments"] == (START, END, " 603 ", db, user, " 6030117 ")
    assert threadpool_calls == [(sales.build_hdyy01_workbook_file, (captured["payload"],), {})]
    assert captured["payload"] is not report
    assert "当前用户权限范围" in captured["payload"]["scope_description"]
    assert "store=4" in captured["payload"]["scope_description"]
    assert report["scope_description"] == "当前用户权限范围：store=4；department=6030117"
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert (
        "HDYY01%E6%9F%9C%E7%BB%84%E7%BB%8F%E8%90%A5%E5%88%86%E6%9E%90%E8%A1%A8_"
        "2026-07-01_2026-07-10.xlsx"
    ) in response.headers["content-disposition"]

    async def consume_and_close():
        chunks = [chunk async for chunk in response.body_iterator]
        await response.background()
        return chunks

    chunks = asyncio.run(consume_and_close())
    assert b"".join(chunks) == b"x" * (64 * 1024 + 3)
    assert all(len(chunk) <= 64 * 1024 for chunk in chunks)
    assert export_file.closed


def test_query_and_export_each_use_shared_loader_once(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    calls = []
    monkeypatch.setattr(
        sales,
        "_load_hdyy01_for_request",
        lambda *args: calls.append(args) or (sample_report(), DataScope(all_access=True)),
    )
    export_file = SpooledTemporaryFile()
    monkeypatch.setattr(sales, "build_hdyy01_workbook_file", lambda report: export_file)

    async def fake_threadpool(function, *args, **kwargs):
        return function(*args, **kwargs)

    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)
    asyncio.run(sales.hdyy01_report(START, END, None, object(), object()))
    assert len(calls) == 1
    calls.clear()
    asyncio.run(sales.hdyy01_export(START, END, None, object(), object()))
    assert len(calls) == 1
    export_file.close()


def test_export_builder_failure_propagates_without_response(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    monkeypatch.setattr(
        sales,
        "_load_hdyy01_for_request",
        lambda *args: (sample_report(), DataScope(all_access=True)),
    )

    def fail(_report):
        raise RuntimeError("builder failed")

    monkeypatch.setattr(sales, "build_hdyy01_workbook_file", fail)

    async def fake_threadpool(function, *args, **kwargs):
        return function(*args, **kwargs)

    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)
    with pytest.raises(RuntimeError, match="builder failed"):
        asyncio.run(sales.hdyy01_export(START, END, None, object(), object()))


def test_export_closes_file_when_asgi_send_fails_during_streaming(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    export_file = SpooledTemporaryFile()
    export_file.write(b"xlsx")
    export_file.seek(0)
    monkeypatch.setattr(
        sales,
        "_load_hdyy01_for_request",
        lambda *args: (sample_report(), DataScope(all_access=True)),
    )
    monkeypatch.setattr(
        sales, "build_hdyy01_workbook_file", lambda report: export_file
    )

    async def fake_threadpool(function, *args, **kwargs):
        return function(*args, **kwargs)

    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)
    response = asyncio.run(
        sales.hdyy01_export(START, END, None, object(), object())
    )
    assert response.background is not None

    async def invoke_response():
        never_disconnect = asyncio.Event()

        async def receive():
            await never_disconnect.wait()
            return {"type": "http.disconnect"}

        async def send(message):
            if message["type"] == "http.response.body":
                raise OSError("client disconnected")

        await response({"type": "http"}, receive, send)

    with pytest.raises(Exception) as exc_info:
        asyncio.run(invoke_response())

    def contains_oserror(error):
        if isinstance(error, OSError):
            return True
        return any(
            contains_oserror(child)
            for child in getattr(error, "exceptions", ())
        )

    assert contains_oserror(exc_info.value)
    assert export_file.closed
