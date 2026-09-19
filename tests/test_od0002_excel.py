import asyncio
from datetime import date
from io import BytesIO
from tempfile import SpooledTemporaryFile

from openpyxl import load_workbook


def sample_report(*, empty=False):
    metrics = {
        "ticket_count_current": 100,
        "ticket_count_prior": 80,
        "ticket_count_yoy": 0.25,
        "average_ticket_current": 1200.0,
        "average_ticket_prior": 1250.0,
        "average_ticket_yoy": -0.04,
        "sales_current": 120000.0,
        "sales_prior": 100000.0,
        "sales_yoy": 0.2,
        "profit_current": 24000.0,
        "profit_prior": 15000.0,
        "profit_yoy": 0.6,
        "margin_current": 0.2,
        "margin_prior": 0.15,
        "margin_change": 0.05,
    }
    dimensions = {}
    totals = {}
    for key in ("stores", "departments", "areas", "categories", "groups", "floors"):
        dimensions[key] = [] if empty else [{
            "store_code": "601",
            "store_name": "一店",
            "dimension_code": "D01" if key != "stores" else "601",
            "dimension_name": "女装" if key != "stores" else "一店",
            "metrics": dict(metrics),
        }]
        totals[key] = dict(metrics)
    hierarchy_metrics = dict(metrics)
    hierarchy_total = {
        **metrics,
        "sales_current": 240000.0,
        "sales_prior": 200000.0,
        "profit_current": 48000.0,
        "profit_prior": 30000.0,
    }
    dimensions["department_categories"] = [] if empty else [
        {
            "store_code": "603", "store_name": "商城",
            "department_code": "6030102", "department_name": "新世纪二部",
            "area_code": "A1", "area_name": "女装区",
            "category_code": "C1", "category_name": "中式女装",
            "dimension_code": "C1", "dimension_name": "中式女装",
            "row_type": "category", "metrics": dict(hierarchy_metrics),
        },
        {
            "store_code": "603", "store_name": "商城",
            "department_code": "6030102", "department_name": "新世纪二部",
            "area_code": "A1", "area_name": "女装区",
            "category_code": "C2", "category_name": "中淑女装",
            "dimension_code": "C2", "dimension_name": "中淑女装",
            "row_type": "category", "metrics": dict(hierarchy_metrics),
        },
        {
            "store_code": "603", "store_name": "商城",
            "department_code": "6030102", "department_name": "新世纪二部",
            "area_code": "A1", "area_name": "女装区",
            "category_code": None, "category_name": None,
            "dimension_code": "A1", "dimension_name": "女装区小计",
            "row_type": "area_subtotal", "metrics": dict(hierarchy_total),
        },
        {
            "store_code": "603", "store_name": "商城",
            "department_code": "6030102", "department_name": "新世纪二部",
            "area_code": None, "area_name": None,
            "category_code": None, "category_name": None,
            "dimension_code": "6030102", "dimension_name": "新世纪二部小计",
            "row_type": "department_subtotal", "metrics": dict(hierarchy_total),
        },
    ]
    totals["department_categories"] = dict(hierarchy_total)
    if not empty:
        dimensions["groups"][0].update({
            "department_code": "6010101",
            "department_name": "一店一部(化妆)",
            "area_code": "A01",
            "area_name": "化妆品区域",
            "category_code": "C01",
            "category_name": "国际化妆品",
            "dimension_code": "6010101005",
            "dimension_name": "L'oreal欧莱雅厅",
        })
    dimensions["special_sales"] = [] if empty else [{
        "store_code": "601",
        "store_name": "一店",
        "department_code": "6010101",
        "department_name": "一店一部(化妆)",
        "dimension_code": "6010101999",
        "dimension_name": "一楼特卖厅",
        "brand_code": "00310",
        "brand_name": "Christian dior迪奥",
        "metrics": dict(metrics),
    }]
    totals["special_sales"] = dict(metrics)
    if not empty:
        dimensions["departments"].append({
            "store_code": "602", "store_name": "二店",
            "dimension_code": "D02", "dimension_name": "男装",
            "metrics": dict(metrics),
        })
    return {
        "dates": {
            "start_date": date(2026, 1, 1), "end_date": date(2026, 1, 31),
            "prior_start_date": date(2025, 1, 1), "prior_end_date": date(2025, 1, 31),
        },
        "selected_store": None,
        "dimensions": dimensions,
        "totals": totals,
        "scope_description": "当前用户权限范围：门店 601、602",
    }


def test_workbook_has_required_sheets_headers_formats_totals_and_notes():
    from python_app.services.od0002_excel import build_od0002_workbook
    from python_app.services.od0002_report import EXCLUDED_DEPARTMENT_CODES

    payload = build_od0002_workbook(sample_report())
    assert isinstance(payload, bytes)
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == [
        "分店", "部门", "部门（含品类）", "区域", "品类", "柜组", "特卖", "楼层", "报表说明"
    ]

    sheet = workbook["部门"]
    assert sheet["A1"].value == "OD0002 门店销售毛利汇总表（部门）"
    assert "2026-01-01" in sheet["A2"].value and "2026-01-31" in sheet["A2"].value
    assert "2025-01-01" in sheet["A3"].value and "2025-01-31" in sheet["A3"].value
    assert sheet.freeze_panes == "A8"
    assert sheet["A5"].fill.fgColor.rgb.endswith("4472C4")
    assert sheet["A5"].font.color.type == "rgb" and sheet["A5"].font.color.rgb.endswith("FFFFFF")
    assert [sheet.cell(5, column).value for column in (5, 8, 11, 14, 17)] == [
        "销售收入", "毛利额", "毛利率", "来客数", "客单",
    ]
    assert sheet["E8"].value == 12
    assert sheet["E8"].number_format == "0.00"
    assert sheet["G8"].number_format == "0.00%"
    assert sheet["K8"].value == 0.2
    assert sheet["K8"].number_format == "0.00%"
    assert sheet["N8"].value == 100
    assert sheet["N8"].number_format == "#,##0"
    assert sheet["Q8"].value == 1200
    assert sheet["Q8"].number_format == "0.00"
    assert sheet["A10"].value == "合计"
    assert sheet["E10"].value == 12
    assert sheet["G10"].value == 0.2
    assert sheet["N10"].value == 100
    assert sheet["P10"].value == 0.25
    assert sheet["A10"].border.bottom.style is not None

    notes = workbook["报表说明"]["B2"].value
    for text in (
        "sglhsrq",
        "sglxssr",
        "sgln2",
        "sglbillno",
        "客单=销售收入（元）/来客数",
        "sglwmid=5",
        "楼层00",
        "16部门",
        "当前用户权限范围",
        "manaframe.mflc=16",
        "salegoodslist.sglppcode",
        "codebrand.cbcname",
        "大类暂不提供",
    ):
        assert text in notes
    sorted_codes = sorted(EXCLUDED_DEPARTMENT_CODES)
    assert sorted_codes[0] in notes and sorted_codes[-1] in notes
    assert all(code in notes for code in sorted_codes)
    assert "合计仅包含当前用户权限范围" in notes


def test_department_category_sheet_has_hierarchy_subtotals_and_total():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report())))
    sheet = workbook["部门（含品类）"]

    assert sheet["A1"].value == "OD0002 门店销售毛利汇总表（部门（含品类））"
    assert [sheet.cell(6, column).value for column in range(1, 5)] == ["门店", "部门", "区域", "品类"]
    assert "商城" in sheet["A8"].value and "603" in sheet["A8"].value
    assert "新世纪二部" in sheet["B8"].value and "6030102" in sheet["B8"].value
    assert "女装区" in sheet["C8"].value and "A1" in sheet["C8"].value
    assert "中式女装" in sheet["D8"].value and "C1" in sheet["D8"].value
    assert sheet["C10"].value.startswith("女装区小计")
    assert sheet["B11"].value.startswith("新世纪二部小计")
    assert sheet["A12"].value == "合计"
    assert sheet["E12"].value == 24
    assert sheet["N12"].value == 100
    assert sheet.freeze_panes == "A8"
    assert sheet.max_column == 19
    assert sheet["B11"].font.bold is True
    assert sheet["B11"].fill.fgColor.rgb.endswith("D9EAF7")


def test_empty_department_category_sheet_keeps_headers_and_total():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report(empty=True))))
    sheet = workbook["部门（含品类）"]

    assert sheet["A6"].value == "门店"
    assert sheet["D6"].value == "品类"
    assert sheet["A8"].value == "合计"


def test_group_sheet_has_department_area_category_before_group_and_twenty_five_columns():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report())))
    sheet = workbook["柜组"]

    assert sheet["A1"].value == "OD0002 门店销售毛利汇总表（柜组）"
    assert "A1:Y1" in {str(item) for item in sheet.merged_cells.ranges}
    assert [sheet.cell(6, column).value for column in range(1, 11)] == [
        "门店编码",
        "门店名称",
        "部门编码",
        "部门名称",
        "区域编码",
        "区域名称",
        "品类编码",
        "品类名称",
        "柜组编码",
        "柜组名称",
    ]
    assert [sheet.cell(8, column).value for column in range(1, 11)] == [
        "601",
        "一店",
        "6010101",
        "一店一部(化妆)",
        "A01",
        "化妆品区域",
        "C01",
        "国际化妆品",
        "6010101005",
        "L'oreal欧莱雅厅",
    ]
    assert [sheet.cell(5, column).value for column in (11, 14, 17, 20, 23)] == [
        "销售收入", "毛利额", "毛利率", "来客数", "客单",
    ]
    assert sheet["K8"].value == 12
    assert sheet["M8"].number_format == "0.00%"
    assert sheet["Q8"].value == 0.2
    assert sheet["T8"].value == 100
    assert sheet["W8"].value == 1200
    assert sheet.max_column == 25
    assert sheet.freeze_panes == "A8"


def test_empty_group_sheet_keeps_twenty_five_columns_and_total():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report(empty=True))))
    sheet = workbook["柜组"]

    assert sheet["J6"].value == "柜组名称"
    assert sheet["A8"].value == "合计"
    assert sheet.max_column == 25


def test_special_sale_sheet_has_group_brand_detail_and_twenty_three_columns():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report())))
    sheet = workbook["特卖"]

    assert sheet["A1"].value == "OD0002 门店销售毛利汇总表（特卖）"
    assert "A1:W1" in {str(item) for item in sheet.merged_cells.ranges}
    assert [sheet.cell(6, column).value for column in range(1, 9)] == [
        "门店编码",
        "门店名称",
        "部门编码",
        "部门名称",
        "柜组编码",
        "柜组名称",
        "品牌编码",
        "品牌名称",
    ]
    assert [sheet.cell(8, column).value for column in range(1, 9)] == [
        "601",
        "一店",
        "6010101",
        "一店一部(化妆)",
        "6010101999",
        "一楼特卖厅",
        "00310",
        "Christian dior迪奥",
    ]
    assert [sheet.cell(5, column).value for column in (9, 12, 15, 18, 21)] == [
        "销售收入", "毛利额", "毛利率", "来客数", "客单",
    ]
    assert sheet["I8"].value == 12
    assert sheet["K8"].number_format == "0.00%"
    assert sheet["O8"].value == 0.2
    assert sheet["R8"].value == 100
    assert sheet["U8"].value == 1200
    assert sheet.max_column == 23
    assert sheet.freeze_panes == "A8"


def test_empty_special_sale_sheet_keeps_headers_and_total():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report(empty=True))))
    sheet = workbook["特卖"]

    assert sheet["H6"].value == "品牌名称"
    assert sheet["A8"].value == "合计"
    assert sheet.max_column == 23


def test_special_sale_sheet_escapes_formula_like_brand_and_group_text():
    from python_app.services.od0002_excel import build_od0002_workbook

    report = sample_report()
    report["dimensions"]["special_sales"][0].update({
        "dimension_code": "=1+1",
        "dimension_name": "+SUM(A1:A2)",
        "brand_code": "-2+3",
        "brand_name": "@cmd",
    })
    workbook = load_workbook(BytesIO(build_od0002_workbook(report)), data_only=False)
    sheet = workbook["特卖"]

    for column in range(5, 9):
        cell = sheet.cell(8, column)
        assert cell.data_type == "s"
        assert cell.value.startswith("'")


def test_group_sheet_escapes_formula_like_department_area_category_and_group_text():
    from python_app.services.od0002_excel import build_od0002_workbook

    report = sample_report()
    report["dimensions"]["groups"][0].update({
        "department_code": "=1+1",
        "department_name": "+SUM(A1:A2)",
        "area_code": "-A01",
        "area_name": "@区域",
        "category_code": "=C01",
        "category_name": "+品类",
        "dimension_code": "-2+3",
        "dimension_name": "@cmd",
    })
    workbook = load_workbook(BytesIO(build_od0002_workbook(report)), data_only=False)
    sheet = workbook["柜组"]

    for column in range(3, 11):
        cell = sheet.cell(8, column)
        assert cell.data_type == "s"
        assert cell.value.startswith("'")


def test_workbook_keeps_store_columns_for_multiple_store_rows():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report())))
    sheet = workbook["部门"]
    assert [sheet.cell(8, column).value for column in range(1, 5)] == ["601", "一店", "D01", "女装"]
    assert [sheet.cell(9, column).value for column in range(1, 5)] == ["602", "二店", "D02", "男装"]


def test_workbook_escapes_formula_like_dimension_text_and_has_no_invalid_filter():
    from python_app.services.od0002_excel import build_od0002_workbook

    report = sample_report()
    report["dimensions"]["departments"][0].update({
        "store_code": "=1+1",
        "store_name": "+SUM(A1:A2)",
        "dimension_code": "-2+3",
        "dimension_name": "@cmd",
    })
    workbook = load_workbook(BytesIO(build_od0002_workbook(report)), data_only=False)
    sheet = workbook["部门"]
    for column in range(1, 5):
        cell = sheet.cell(8, column)
        assert cell.data_type == "s"
        assert cell.value.startswith("'")
    assert sheet.auto_filter.ref is None


def test_empty_report_still_has_valid_headers_and_notes():
    from python_app.services.od0002_excel import build_od0002_workbook

    workbook = load_workbook(BytesIO(build_od0002_workbook(sample_report(empty=True))))
    assert workbook["部门"]["A5"].value == "维度"
    assert workbook["部门"]["A8"].value == "合计"
    assert "大类暂不提供" in workbook["报表说明"]["B2"].value


def test_workbook_file_api_returns_seeked_temporary_file():
    from python_app.services.od0002_excel import build_od0002_workbook_file

    export_file = build_od0002_workbook_file(sample_report())
    try:
        assert isinstance(export_file, SpooledTemporaryFile)
        assert export_file.tell() == 0
        assert load_workbook(export_file).sheetnames[0] == "分店"
    finally:
        export_file.close()


def test_export_route_reuses_permission_scope_loader_and_sets_disposition(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    calls = {"load": 0}
    monkeypatch.setattr(sales, "require_permission", lambda db, user, code: calls.setdefault("permission", code))
    monkeypatch.setattr(sales, "load_business_scope", lambda *args, **kwargs: DataScope(allow={"store": {"601"}}))
    monkeypatch.setattr(sales, "_od0002_store_id_for_code", lambda db, code: "601")
    monkeypatch.setattr(sales, "_business_scope_filter_sql", lambda scope, params, **kwargs: " AND 1=1")
    def fake_load(*args, **kwargs):
        calls["load"] += 1
        calls["selected_store"] = kwargs["selected_store"]
        calls["selected_department"] = kwargs["selected_department"]
        calls["prior_start_date"] = kwargs["prior_start_date"]
        calls["prior_end_date"] = kwargs["prior_end_date"]
        return sample_report()
    monkeypatch.setattr(sales, "load_od0002_report", fake_load)
    export_file = SpooledTemporaryFile()
    export_file.write(b"xlsx")
    export_file.seek(0)
    monkeypatch.setattr(sales, "build_od0002_workbook_file", lambda report: export_file)
    threadpool_calls = []
    async def fake_threadpool(function, *args, **kwargs):
        threadpool_calls.append((function, args, kwargs))
        return function(*args, **kwargs)
    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)

    response = asyncio.run(sales.od0002_export(
        date(2026, 1, 1),
        date(2026, 1, 31),
        " 601 ",
        object(),
        object(),
        " 6030117 ",
        date(2024, 12, 29),
        date(2025, 1, 28),
    ))

    assert calls == {
        "load": 1,
        "permission": "sales.od0002.view",
        "prior_end_date": date(2025, 1, 28),
        "prior_start_date": date(2024, 12, 29),
        "selected_store": "601",
        "selected_department": "6030117",
    }
    assert len(threadpool_calls) == 1
    assert threadpool_calls[0][0] is sales.build_od0002_workbook_file
    assert response.background is not None
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    disposition = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert "OD0002_%E9%97%A8%E5%BA%97%E9%94%80%E5%94%AE%E6%AF%9B%E5%88%A9%E6%B1%87%E6%80%BB%E8%A1%A8_2026-01-01_2026-01-31.xlsx" in disposition

    async def consume_and_close():
        chunks = [chunk async for chunk in response.body_iterator]
        await response.background()
        return chunks

    chunks = asyncio.run(consume_and_close())
    assert b"".join(chunks) == b"xlsx"
    assert all(len(chunk) <= 64 * 1024 for chunk in chunks)
    assert export_file.closed


def test_query_and_export_each_load_report_once(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    calls = []
    monkeypatch.setattr(sales, "require_permission", lambda *args: None)
    monkeypatch.setattr(sales, "load_business_scope", lambda *args, **kwargs: DataScope(all_access=True))
    monkeypatch.setattr(sales, "_business_scope_filter_sql", lambda *args, **kwargs: "")
    monkeypatch.setattr(sales, "load_od0002_report", lambda *args, **kwargs: calls.append(kwargs) or sample_report())
    monkeypatch.setattr(sales, "build_od0002_workbook_file", lambda report: SpooledTemporaryFile())
    async def fake_threadpool(function, *args, **kwargs):
        return function(*args, **kwargs)
    monkeypatch.setattr(sales, "run_in_threadpool", fake_threadpool)

    asyncio.run(sales.od0002_report(date(2026, 1, 1), date(2026, 1, 31), None, object(), object()))
    assert len(calls) == 1
    calls.clear()
    asyncio.run(sales.od0002_export(date(2026, 1, 1), date(2026, 1, 31), None, object(), object()))
    assert len(calls) == 1
