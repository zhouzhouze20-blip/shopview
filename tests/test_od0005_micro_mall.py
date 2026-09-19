from datetime import date
from io import BytesIO
from pathlib import Path
import sqlite3

import pytest
from openpyxl import load_workbook
from fastapi import HTTPException
from sqlalchemy.exc import OperationalError

from python_app.routers import od0005_micro_mall
from python_app.routers.authz import CORE_PERMISSION_DEFINITIONS
from python_app.services.od0002_report import TrustedScopeSql
from python_app.services.od0005_micro_mall_excel import build_od0005_workbook
from python_app.services.od0005_micro_mall_report import (
    assemble_od0005_brand_yoy,
    assemble_od0005_daily,
    assemble_od0005_report,
    build_od0005_daily_query,
    build_od0005_query,
    MICRO_MALL_PAYMENT_FILTER_SQL,
)


@pytest.mark.parametrize("store", ["601", "602", "603"])
def test_three_stores_use_same_payment_code(store):
    for builder in (build_od0005_query, build_od0005_daily_query):
        _, params = builder(
            start_date=date(2026, 8, 1), end_date=date(2026, 8, 1),
            selected_store=store, selected_department=None,
            scope_filter_sql=TrustedScopeSql(""), scope_params={},
        )
        assert params["micro_mall_payment_code"] == "0581"
        assert "micro_mall_cashier" not in params
        with pytest.raises(ValueError, match="unsupported OD0005 store"):
            builder(
                start_date=date(2026, 8, 1), end_date=date(2026, 8, 1),
                selected_store="604", selected_department=None,
                scope_filter_sql=TrustedScopeSql(""), scope_params={},
            )


def test_query_keeps_source_columns_and_payment_groups():
    sql, params = build_od0005_query(
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 1),
        selected_store=" 601 ",
        selected_department=" 6010114 ",
        scope_filter_sql=TrustedScopeSql(" AND st.store_id::text = ANY(:allowed_stores)"),
        scope_params={"allowed_stores": ["1"]},
    )
    compact = " ".join(sql.lower().split())

    assert "s.sgldate between :start_date and :end_date" in compact
    assert "s.sglchecker" not in compact
    assert "sum(coalesce(s.sglsl, 0))" in compact
    assert "sum(coalesce(s.sglsjje, 0))" in compact
    assert "coalesce(s.sglxssr, 0) + coalesce(s.sgltotzk, 0)" in compact
    assert "sum(coalesce(s.sgln2, 0))" in compact
    assert "spg.spgpmcode in ('2031', '1014', '3064')" in compact
    assert "spg.spgpmcode = '0511'" in compact
    assert "spg.spgpmtype = '5'" in compact
    assert "join sellpaygoods spg" in compact
    assert params["micro_mall_payment_code"] == "0581"
    assert params["selected_department"] == "6010114"


def test_daily_query_uses_current_and_prior_year_dates_with_same_scope():
    sql, params = build_od0005_daily_query(
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 2),
        selected_store="601",
        selected_department="6010114",
        scope_filter_sql=TrustedScopeSql(" AND st.store_id::text = ANY(:allowed_stores)"),
        scope_params={"allowed_stores": ["1"]},
    )
    compact = " ".join(sql.lower().split())

    assert "('current', cast(:start_date as date), cast(:end_date as date))" in compact
    assert "('prior', cast(:prior_start_date as date), cast(:prior_end_date as date))" in compact
    assert "s.sgldate as sales_date" in compact
    assert "sum(coalesce(s.sglxssr, 0)) as sales_revenue" in compact
    assert "st.store_id::text = any(:allowed_stores)" in compact
    assert params["prior_start_date"] == date(2025, 8, 1)
    assert params["prior_end_date"] == date(2025, 8, 2)
    assert params["selected_department"] == "6010114"

    three_year_sql, three_year_params = build_od0005_daily_query(
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 2),
        selected_store="601",
        selected_department=None,
        scope_filter_sql=TrustedScopeSql(""),
        scope_params={},
        include_two_year_prior=True,
    )
    assert "'two_year_prior'" in three_year_sql
    assert three_year_params["two_year_prior_start_date"] == date(2024, 8, 1)
    assert three_year_params["two_year_prior_end_date"] == date(2024, 8, 2)


def test_micro_mall_queries_use_payment_exists_without_cashier_filter():
    report_sql, _ = build_od0005_query(
        start_date=date(2026, 1, 1),
        end_date=date(2026, 8, 31),
        selected_store="602",
        selected_department=None,
        scope_filter_sql=TrustedScopeSql(""),
        scope_params={},
    )
    daily_sql, _ = build_od0005_daily_query(
        start_date=date(2026, 1, 1),
        end_date=date(2026, 8, 31),
        selected_store="602",
        selected_department=None,
        scope_filter_sql=TrustedScopeSql(""),
        scope_params={},
    )

    for sql in (report_sql, daily_sql):
        compact = " ".join(sql.lower().split())
        assert "sglchecker" not in compact
        assert "micro_mall_cashier" not in compact
        assert MICRO_MALL_PAYMENT_FILTER_SQL in sql
        assert "micro_payment.spgbillno = s.sglbillno" in compact
        assert "micro_payment.spgpmcode) = :micro_mall_payment_code" in compact
        assert "trim(both from s.sglmarket::text) = :selected_store" in compact
        assert "s.sgldate between" in compact
        assert "s.sgldate::date" not in compact


def test_payment_selection_includes_other_cashiers_and_does_not_duplicate_sales():
    # Execute the shared production predicate; only TRIM syntax differs in SQLite.
    predicate = MICRO_MALL_PAYMENT_FILTER_SQL.replace(
        "TRIM(BOTH FROM micro_payment.spgpmcode)", "TRIM(micro_payment.spgpmcode)"
    )
    with sqlite3.connect(":memory:") as db:
        db.executescript("""
            CREATE TABLE salegoodslist (
                sglbillno INTEGER, sglchecker TEXT, sglmarket TEXT,
                sgldate TEXT, sglxssr NUMERIC
            );
            CREATE TABLE sellpaygoods (
                spgbillno INTEGER, spgpmcode TEXT, spgpmtype TEXT
            );
            INSERT INTO salegoodslist VALUES
                (1, 'other-cashier', '601', '2026-08-01', 100),
                (1, 'other-cashier', '601', '2026-08-01', 50),
                (2, '300411', '601', '2026-08-01', 999),
                (3, NULL, '601', '2026-08-01', -20),
                (4, '300411', '601', '2026-08-01', 888),
                (5, 'other-cashier', '602', '2026-08-01', 777),
                (6, 'other-cashier', '601', '2025-08-01', 60);
            INSERT INTO sellpaygoods VALUES
                (1, '0581', '5'), (1, '0581', '5'), (1, '0511', '5'),
                (2, '0511', '5'), (3, ' 0581 ', '1'),
                (5, '0581', '5'), (6, '0581', '5');
        """)
        for period, expected in [("2026-08-01", 130), ("2025-08-01", 60)]:
            amount = db.execute(
                f"""SELECT SUM(s.sglxssr) FROM salegoodslist s
                WHERE s.sglmarket = :store AND s.sgldate = :period {predicate}""",
                {"store": "601", "period": period, "micro_mall_payment_code": "0581"},
            ).fetchone()[0]
            assert amount == expected


def test_statement_timeout_is_reported_as_gateway_timeout(monkeypatch):
    class _Db:
        rolled_back = False

        def rollback(self):
            self.rolled_back = True

    db = _Db()

    def timeout(*_args, **_kwargs):
        raise OperationalError(
            "SELECT 1",
            {},
            RuntimeError("canceling statement due to statement timeout"),
        )

    monkeypatch.setattr(od0005_micro_mall, "load_od0005_report", timeout)

    with pytest.raises(HTTPException) as exc_info:
        od0005_micro_mall._execute_od0005_report(db, report_kwargs={})

    assert exc_info.value.status_code == 504
    assert "查询超时" in exc_info.value.detail
    assert db.rolled_back is True


def test_report_and_workbook_preserve_template_layout():
    report = assemble_od0005_report(
        [
            {
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_quantity": 4,
                "price_amount": 6800,
                "sales_before_discount": 6800,
                "sales_revenue": 5458.8,
                "gross_profit": 1200,
                "gross_margin": 0,
                "yzq_amount": 1661.2,
                "other_payment_amount": 5458.8,
                "nzd_amount": 0,
            }
        ],
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 1),
        selected_store="601",
        selected_department=None,
    )
    assert report["payment_code"] == "0581"
    assert "cashier_code" not in report
    assert report["totals"]["gross_margin"] == pytest.approx(1200 / 5458.8)

    workbook = load_workbook(BytesIO(build_od0005_workbook(report)))
    sheet = workbook["微商城品牌销售统计"]
    assert sheet["A1"].value == "普灵仕集团百货事业部"
    assert sheet["A2"].value == "微商城品牌销售统计"
    assert sheet["A4"].value == "门店"
    assert sheet["D4"].value == "柜组名称"
    assert sheet["E4"].value == "销售数量"
    assert sheet["J4"].value == "毛利率"
    assert sheet["M4"].value == "NZD"
    assert sheet["D5"].value == "SK-II厅"
    assert sheet.freeze_panes == "A5"


def test_workbook_adds_department_sales_summary_sheet():
    report = assemble_od0005_report(
        [
            {
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_quantity": 4,
                "price_amount": 6800,
                "sales_before_discount": 5600,
                "sales_revenue": 5458.8,
                "gross_profit": 1200,
                "yzq_amount": 1661.2,
                "other_payment_amount": 3738.8,
                "nzd_amount": 58.8,
            },
            {
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101006",
                "group_name": "兰蔻厅",
                "sales_quantity": 2,
                "price_amount": 3200,
                "sales_before_discount": 3000,
                "sales_revenue": 2800,
                "gross_profit": 600,
                "yzq_amount": 1000,
                "other_payment_amount": 1800,
                "nzd_amount": 0,
            },
            {
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010119",
                "department_name": "中心儿童游乐园",
                "group_code": "6010119001",
                "group_name": "游乐园",
                "sales_quantity": 1,
                "price_amount": 39.9,
                "sales_before_discount": 39.9,
                "sales_revenue": 39.9,
                "gross_profit": 39.9,
                "yzq_amount": 0,
                "other_payment_amount": 39.9,
                "nzd_amount": 0,
            },
        ],
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 1),
        selected_store="601",
        selected_department=None,
    )

    workbook = load_workbook(BytesIO(build_od0005_workbook(report)))
    assert workbook.sheetnames == ["微商城品牌销售统计", "部门销售统计", "逐日销售", "逐日同期同比", "品牌同比"]

    sheet = workbook["部门销售统计"]
    assert sheet["A2"].value == "微商城部门销售统计"
    assert [sheet.cell(4, column).value for column in range(1, 11)] == [
        "门店",
        "部门",
        "销售数量",
        "应收金额",
        "有赞卡券",
        "礼券",
        "销售收入",
        "销售收入(内转)",
        "毛利额",
        "毛利率",
    ]
    assert sheet["B5"].value == "6010114 中心一部(化妆)"
    assert sheet["C5"].value == pytest.approx(6)
    assert sheet["D5"].value == pytest.approx(10000)
    assert sheet["E5"].value == pytest.approx(2661.2)
    assert sheet["F5"].value == pytest.approx(5538.8)
    assert sheet["G5"].value == pytest.approx(8600)
    assert sheet["H5"].value == pytest.approx(8258.8)
    assert sheet["I5"].value == pytest.approx(1800)
    assert sheet["J5"].value == pytest.approx(1800 / 8258.8)
    assert sheet["A7"].value == "合计"
    assert sheet["D7"].value == pytest.approx(report["totals"]["price_amount"])
    assert sheet["J7"].value == pytest.approx(report["totals"]["gross_margin"])
    assert sheet.freeze_panes == "A5"


def test_workbook_adds_daily_sales_and_daily_yoy_sheets():
    report = assemble_od0005_report(
        [
            {
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_quantity": 2,
                "price_amount": 300,
                "sales_before_discount": 300,
                "sales_revenue": 300,
                "gross_profit": 60,
                "yzq_amount": 0,
                "other_payment_amount": 300,
                "nzd_amount": 0,
            }
        ],
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 2),
        selected_store="601",
        selected_department=None,
    )
    report["daily"] = assemble_od0005_daily(
        [
            {
                "period_key": "current",
                "sales_date": date(2026, 8, 1),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 100,
            },
            {
                "period_key": "current",
                "sales_date": date(2026, 8, 2),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 200,
            },
            {
                "period_key": "prior",
                "sales_date": date(2025, 8, 1),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 80,
            },
        ],
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 2),
    )
    report["brand_yoy"] = assemble_od0005_brand_yoy(
        [
            {
                "period_key": "two_year_prior",
                "sales_date": date(2024, 8, 1),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 60,
            },
            {
                "period_key": "prior",
                "sales_date": date(2025, 8, 1),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 80,
            },
            {
                "period_key": "current",
                "sales_date": date(2026, 8, 1),
                "store_code": "601",
                "store_name": "常州购物中心",
                "department_code": "6010114",
                "department_name": "中心一部(化妆)",
                "group_code": "6010101005",
                "group_name": "SK-II厅",
                "sales_revenue": 100,
            },
        ],
        start_date=date(2026, 8, 1),
        end_date=date(2026, 8, 2),
    )

    assert report["daily"]["rows"][0]["daily"][0]["sales_yoy"] == pytest.approx(0.25)
    assert report["daily"]["rows"][0]["daily"][1]["sales_yoy"] is None

    workbook = load_workbook(BytesIO(build_od0005_workbook(report)))
    assert workbook.sheetnames == ["微商城品牌销售统计", "部门销售统计", "逐日销售", "逐日同期同比", "品牌同比"]

    daily_sheet = workbook["逐日销售"]
    assert daily_sheet["A2"].value == "微商城逐日销售"
    assert daily_sheet["E4"].value == "2026-08-01"
    assert daily_sheet["F4"].value == "2026-08-02"
    assert daily_sheet["E5"].value == pytest.approx(100)
    assert daily_sheet["F5"].value == pytest.approx(200)

    yoy_sheet = workbook["逐日同期同比"]
    assert yoy_sheet["A2"].value == "微商城逐日同期同比"
    assert "同期：2025-08-01 至 2025-08-02" in yoy_sheet["A3"].value
    assert yoy_sheet["E4"].value == "2026-08-01"
    assert [yoy_sheet.cell(5, column).value for column in range(5, 8)] == [
        "本期销售",
        "同期销售",
        "同比",
    ]
    assert yoy_sheet["E6"].value == pytest.approx(100)
    assert yoy_sheet["F6"].value == pytest.approx(80)
    assert yoy_sheet["G6"].value == pytest.approx(0.25)
    assert yoy_sheet.freeze_panes == "A6"

    brand_yoy_sheet = workbook["品牌同比"]
    assert brand_yoy_sheet["A2"].value == "26年8月品牌销售情况及同比"
    assert [brand_yoy_sheet.cell(4, column).value for column in range(1, 10)] == [
        "部门",
        "品牌",
        "24年",
        "25年",
        "26年",
        "差额",
        "同比",
        "原因分析",
        "缺失商品",
    ]
    assert brand_yoy_sheet["B5"].value == "SK-II厅"
    assert brand_yoy_sheet["C5"].value == pytest.approx(60)
    assert brand_yoy_sheet["D5"].value == pytest.approx(80)
    assert brand_yoy_sheet["E5"].value == pytest.approx(100)
    assert brand_yoy_sheet["F5"].value == pytest.approx(20)
    assert brand_yoy_sheet["G5"].value == pytest.approx(0.25)
    assert brand_yoy_sheet["H5"].value is None
    assert brand_yoy_sheet["I5"].value is None


def test_permission_is_registered_and_migration_preserves_report_access():
    registered = {permission[0] for permission in CORE_PERMISSION_DEFINITIONS}
    assert "sales.od0005.view" in registered
    migration = Path(
        "python_app/alembic/versions/u9c0d1e2f3a4_add_od0005_micro_mall_permission.py"
    ).read_text(encoding="utf-8")
    assert "'sales.od0005.view'" in migration
    assert "source.permission_code = 'sales.od0001.view'" in migration
    assert "查看非租赁品牌月度收益表" in migration
