from decimal import Decimal
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python_app"))

import services.revenue_month_close_import as month_close_import
from services.revenue_month_close_import import parse_month_close_workbook


def _build_workbook(path, fee_amount_label="富基全部联营租赁收费"):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "月结结论"
    summary.append(["月结项目", "金额"])
    for label, amount in (
        ("NC6051不含计提税", 110),
        ("NC计提税净额", 0),
        ("NC6051含计提税控制数", 110),
        (fee_amount_label, 100),
        ("NC非富基收费", 20),
        ("月结净调整", -10),
        ("调平后收益", 110),
    ):
        summary.append([label, amount])

    department = workbook.create_sheet("部门科目调平")
    department.append([None] * 14)
    department.append([None] * 14)
    department.append([None] * 14)
    department.append(["NC部门", "部门名称", "6051科目", "科目名称"])
    department.append(["3117", "新世纪九部", "605106", "广告服务费"])

    adjustment = workbook.create_sheet("柜位月结调整")
    adjustment.append([None] * 18)
    adjustment.append([None] * 18)
    adjustment.append([None] * 18)
    adjustment.append([
        "调整类别", "NC部门", "6051科目", "业务", "柜组编码", "柜组名称",
        "供应商编码", "供应商名称", "费用码", "费用名称", "源单数", "源行数",
        "富基原金额", "月结调整", "调平后金额", "分摊依据", "处理原因", "月结月份",
    ])
    adjustment.append([
        "已匹配组合调平", "3117", "605106", "RENTAL", "6030104082", "大鼓米线厅",
        "05236", "供应商", "08", "广告服务费", 1, 1, 100, -10, 90,
        "正向收费金额占比", "NC控制数调平", "2026-07",
    ])

    checks = workbook.create_sheet("数据检查")
    checks.append([None] * 6)
    checks.append([None] * 6)
    checks.append([None] * 6)
    checks.append(["检查项", "实际值", "预期值", "差异", "容差", "状态"])
    checks.append(["最终收益=NC", 110, 110, 0, 0.01, "PASS"])
    workbook.save(path)


def test_parse_month_close_workbook_reconciles_control_amount(tmp_path):
    path = tmp_path / "month-close.xlsx"
    _build_workbook(path)

    parsed = parse_month_close_workbook(path)

    assert parsed.period_month == "2026-07"
    assert parsed.nc_control_amount == Decimal("110.00")
    assert parsed.close_adjustment_amount == Decimal("-10.00")
    assert len(parsed.adjustments) == 1
    assert parsed.adjustments[0].target_component == "FEE"
    assert parsed.adjustments[0].source_department_name == "新世纪九部"
    assert parsed.adjustments[0].source_subject_name == "广告服务费"
    assert parsed.adjustments[0].binding_status == "BOUND"


def test_parse_month_close_workbook_accepts_card_fee_excluded_fuji_label(tmp_path):
    path = tmp_path / "month-close-card-fees-excluded.xlsx"
    _build_workbook(path, fee_amount_label="富基6051可比联营租赁收费")

    parsed = parse_month_close_workbook(path)

    assert parsed.raw_fee_amount == Decimal("100.00")


def test_parse_month_close_workbook_keeps_unmatched_difference_pending(tmp_path):
    path = tmp_path / "month-close-pending.xlsx"
    _build_workbook(path)
    workbook = load_workbook(path)
    row = workbook["柜位月结调整"][5]
    row[0].value = "待人工绑定-富基收费差异"
    row[4].value = None
    row[5].value = "新世纪九部·广告服务费待绑定"
    row[15].value = "不分摊；人工确认柜位后生效"
    workbook.save(path)

    parsed = parse_month_close_workbook(path)

    assert parsed.adjustments[0].binding_status == "PENDING"
    assert parsed.adjustments[0].source_group_code is None
    assert parsed.adjustments[0].adjustment_amount == Decimal("-10.00")


def test_month_close_fee_type_code_schema_accepts_combined_fuji_codes():
    migration = (
        Path(__file__).resolve().parents[1]
        / "python_app"
        / "alembic"
        / "versions"
        / "c7e8f9a0b1c2_expand_month_close_fee_type_code.py"
    ).read_text(encoding="utf-8")

    assert "ALTER COLUMN fee_type_code TYPE VARCHAR(100)" in migration


def test_month_close_manual_binding_schema_is_auditable():
    migration = (
        Path(__file__).resolve().parents[1]
        / "python_app"
        / "alembic"
        / "versions"
        / "d8f9a0b1c2d3_add_month_close_manual_binding.py"
    ).read_text(encoding="utf-8")

    assert "binding_status" in migration
    assert "PENDING" in migration
    assert "BOUND" in migration
    assert "bound_by INTEGER REFERENCES users(user_id)" in migration
    assert "binding_note TEXT" in migration


def test_month_close_import_note_is_store_specific():
    source = Path(month_close_import.__file__).read_text(encoding="utf-8")

    assert "store_name = str(store[\"store_name\"]" in source
    assert 'f"{store_name}{period_year}财务{period_number}月"' in source
    assert "新世纪2026财务7月" not in source
