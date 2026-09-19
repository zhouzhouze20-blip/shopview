from datetime import date, datetime, timezone
from decimal import Decimal as D
from pathlib import Path
import json
import sqlite3
import sys

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python_app"))
from services.fungkids_retail_import import parse_export, import_snapshot, settlement_rate
from sync_fungkids_retail import query_window, INTERVAL_SECONDS


def export_file(tmp_path, rows, total=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "店铺零售报表"
    sheet.append(["会员名称", "单据编号", "渠道编号", "单据日期", "货号", "计收额",
                  "数量", "计收折扣", "销售方式", "原单号", "销售类型", "促销信息", "计收价", "生意额"])
    for row in rows:
        sheet.append(["不得导入的会员信息", *row])
    sheet.append(["合计", None, None, None, None,
                  total if total is not None else sum(row[4] for row in rows)])
    path = tmp_path / "export.xlsx"
    book.save(path)
    return path


def line(amount=1000, discount=1, day="2026-09-05", bill="XS1", product="A1"):
    return [bill, "PBCZ6001", day, product, amount, -1 if amount < 0 else 1,
            discount, "退货" if amount < 0 else "销售", "XS0" if amount < 0 else ""]


def parse(path):
    return parse_export(path, start=date(2026, 9, 4), end=date(2026, 9, 5))


@pytest.mark.parametrize("discount,rate", [("1", ".50"), (".85", ".50"),
    (".8499", ".52"), (".70", ".52"), (".6999", ".55"), ("0", ".55")])
def test_contract_boundaries(discount, rate):
    assert settlement_rate(D(discount)) == D(rate)


def test_user_example(tmp_path):
    result = parse(export_file(tmp_path, [line()])).summary()
    assert result["sales"] == "1000"
    assert result["supplier_cost"] == "500.00"
    assert result["venue_fee"] == "150.00"
    assert result["expected_receipt"] == "850.00"
    assert result["gross_profit"] == "350.00"


def test_returns_reverse_cost_fee_and_profit(tmp_path):
    result = parse(export_file(tmp_path, [line(213.19, .6999), line(-213.19, .6999)])).summary()
    for key in ("sales", "supplier_cost", "venue_fee", "gross_profit"):
        assert D(result[key]) == 0


def test_repeat_import_preserves_identical_lines_without_double_count(tmp_path):
    parsed = parse(export_file(tmp_path, [line(), line(), line(product="A2")]))
    db_path = tmp_path / "stage.sqlite3"
    import_snapshot(db_path, parsed)
    import_snapshot(db_path, parsed)
    with sqlite3.connect(db_path) as db:
        payloads = [json.loads(row[0]) for row in db.execute("SELECT payload FROM fungkids_lines")]
        assert len(payloads) == 3
        assert sum(D(row["sales"]) for row in payloads) == 3000
        assert all("会员" not in str(row) for row in payloads)
        assert db.execute("SELECT count(*) FROM fungkids_batches").fetchone()[0] == 2


def test_reimport_replaces_only_declared_interval(tmp_path):
    db_path = tmp_path / "stage.sqlite3"
    import_snapshot(db_path, parse(export_file(tmp_path, [line(day="2026-09-04"), line()])))
    changed = parse_export(export_file(tmp_path, [line(700)]),
                          start=date(2026, 9, 5), end=date(2026, 9, 5))
    import_snapshot(db_path, changed)
    with sqlite3.connect(db_path) as db:
        rows = list(db.execute("SELECT payload FROM fungkids_lines"))
        assert len(rows) == 2
        assert sum(D(json.loads(row[0])["sales"]) for row in rows) == 1700


@pytest.mark.parametrize("rows,total", [([line()], 999),
    ([line(day="2026-09-06")], None), ([line(discount=85)], None),
    ([line(discount=None)], None)])
def test_invalid_export_is_rejected(tmp_path, rows, total):
    with pytest.raises(ValueError):
        parse(export_file(tmp_path, rows, total))


def test_empty_export_does_not_clear_existing_data(tmp_path):
    db_path = tmp_path / "stage.sqlite3"
    import_snapshot(db_path, parse(export_file(tmp_path, [line()])))
    with pytest.raises(ValueError, match="空报表"):
        import_snapshot(db_path, parse(export_file(tmp_path, [])))
    with sqlite3.connect(db_path) as db:
        assert db.execute("SELECT count(*) FROM fungkids_lines").fetchone()[0] == 1


def test_two_day_window_uses_shanghai_date_across_midnight():
    assert query_window(datetime(2026, 9, 5, 16, 1, tzinfo=timezone.utc)) == (
        date(2026, 9, 5), date(2026, 9, 6))
    assert INTERVAL_SECONDS == 3600


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-0.01", "1.01"])
def test_invalid_discount_is_rejected(value):
    with pytest.raises(ValueError):
        settlement_rate(D(value))


def test_wrong_channel_is_rejected(tmp_path):
    row = line()
    row[1] = "OTHER"
    with pytest.raises(ValueError, match="渠道"):
        parse(export_file(tmp_path, [row]))


@pytest.mark.parametrize("blank", [None, "", "  "])
@pytest.mark.parametrize("promotion", ["买满赠送(按金额)", "CX2023112300000002,会员积分兑换货品,会员600积分兑换礼盒"])
@pytest.mark.parametrize("sale_type,quantity", [("销售", 1), ("退货", -1), ("换货", 1)])
def test_blank_gift_amount_and_discount_preserve_line(tmp_path, blank, sale_type, quantity, promotion):
    gift = line(0, discount=blank, product="GIFT") + ["正常销售", promotion, None, None]
    gift[4], gift[5], gift[7] = blank, quantity, sale_type
    result = parse(export_file(tmp_path, [line(), gift], total=1000))
    row = result.lines[1]
    assert row.product == "GIFT" and row.quantity == quantity and row.sale_type == sale_type
    assert row.discount == 0
    for key in ("sales", "supplier_cost", "venue_fee", "expected_receipt", "gross_profit"):
        assert getattr(row, key) == 0
    assert result.summary()["row_count"] == 2
    assert D(result.summary()["sales"]) == 1000
    assert D(result.summary()["gross_profit"]) == 350


@pytest.mark.parametrize("field,value", [(4, "bad"), (4, "NaN"), (6, "bad"), (6, "Infinity"),
                                         (6, 2), (11, 199), (12, 199)])
@pytest.mark.parametrize("promotion", ["买满赠送(按金额)", "会员积分兑换货品"])
def test_gift_marker_does_not_hide_invalid_or_paid_values(tmp_path, field, value, promotion):
    gift = line(0, discount=None) + ["正常销售", promotion, None, None]
    gift[4] = None
    gift[field] = value
    with pytest.raises(ValueError):
        parse(export_file(tmp_path, [gift], total=0))


def test_unmarked_blank_amount_and_gift_total_mismatch_are_rejected(tmp_path):
    row = line()
    row[4] = None
    with pytest.raises(ValueError, match="计收额缺少有效数值"):
        parse(export_file(tmp_path, [row], total=0))
    row += ["赠品", None, None, None]
    with pytest.raises(ValueError, match="合计不一致"):
        parse(export_file(tmp_path, [row], total=1000))


def test_paid_exchange_retains_actual_amount_and_discount(tmp_path):
    row = line(199, discount=.5) + ["正常销售", "买满赠送(按金额),加199元换购", 199, 199]
    result = parse(export_file(tmp_path, [row]))
    assert result.lines[0].sales == 199
    assert result.lines[0].supplier_cost == D("109.45")


@pytest.mark.parametrize("discount", [.9, .7, .5])
def test_exchange_uses_normal_sales_settlement(tmp_path, discount):
    row = line(328, discount=discount)
    row[7] = "换货"
    result = parse(export_file(tmp_path, [row, line(-100, discount=discount)]))
    normal = parse(export_file(tmp_path, [line(328, discount=discount), line(-100, discount=discount)]))
    assert result.lines[0].sale_type == "换货"
    assert result.summary() == normal.summary()
    assert result.summary()["return_rows"] == 1
    assert result.lines[0].sales == 328


@pytest.mark.parametrize("amount,quantity", [(-328, 1), (328, -1), (-328, -1)])
def test_exchange_rejects_negative_sales_or_quantity(tmp_path, amount, quantity):
    row = line(amount)
    row[5], row[7] = quantity, "换货"
    with pytest.raises(ValueError, match="符号与销售方式不符"):
        parse(export_file(tmp_path, [row]))


def test_paid_points_redemption_retains_actual_amount(tmp_path):
    row = line(199, discount=.5) + ["正常销售", "会员积分兑换货品", 199, 199]
    result = parse(export_file(tmp_path, [row]))
    assert result.lines[0].sales == 199
    assert result.lines[0].supplier_cost == D("109.45")


def unpriced_export(tmp_path, blank=None, sale_type="销售", quantity=1):
    row = line(0, discount=blank, product="PBAXLGBXX") + ["正常销售", None, blank, blank]
    row[4], row[5], row[7] = blank, quantity, sale_type
    path = export_file(tmp_path, [line(), row], total=1000)
    book = load_workbook(path)
    sheet = book.active
    for col, name in enumerate(("吊牌价", "吊牌额", "零售价"), 15):
        sheet.cell(1, col, name)
        sheet.cell(3, col, blank)
    book.save(path)
    book.close()
    return path


@pytest.mark.parametrize("blank", [None, "", "  ", 0, "0.00"])
@pytest.mark.parametrize("sale_type,quantity", [("销售", 1), ("换货", 1), ("退货", -1)])
def test_unmarked_unpriced_line_preserves_quantity_and_zero_amounts(tmp_path, blank, sale_type, quantity):
    result = parse(unpriced_export(tmp_path, blank, sale_type, quantity))
    row = result.lines[1]
    assert row.product == "PBAXLGBXX" and row.quantity == quantity
    assert row.sale_type == sale_type and row.discount == 0
    assert all(getattr(row, field) == 0 for field in (
        "sales", "supplier_cost", "venue_fee", "expected_receipt", "gross_profit"))
    assert result.summary()["row_count"] == 2
    assert D(result.summary()["sales"]) == 1000


@pytest.mark.parametrize("field", ["吊牌价", "吊牌额", "零售价", "生意额", "计收价"])
@pytest.mark.parametrize("value", [199, "bad", "NaN", "Infinity"])
def test_unmarked_missing_amount_with_other_price_is_rejected(tmp_path, field, value):
    path = unpriced_export(tmp_path)
    book = load_workbook(path)
    header = [cell.value for cell in book.active[1]]
    book.active.cell(3, header.index(field) + 1, value)
    book.save(path)
    with pytest.raises(ValueError):
        parse(path)


@pytest.mark.parametrize("field,value", [("吊牌价", "missing_column"),
    ("计收折扣", "bad"), ("计收折扣", 2), ("数量", -1), ("销售方式", "未知")])
def test_unpriced_line_still_requires_schema_and_valid_business_fields(tmp_path, field, value):
    path = unpriced_export(tmp_path)
    book = load_workbook(path)
    header = [cell.value for cell in book.active[1]]
    book.active.cell(1 if value == "missing_column" else 3, header.index(field) + 1, value)
    book.save(path)
    with pytest.raises(ValueError):
        parse(path)


def test_unpriced_line_does_not_bypass_control_total(tmp_path):
    path = unpriced_export(tmp_path)
    book = load_workbook(path)
    book.active.cell(4, 6, 1001)
    book.save(path)
    with pytest.raises(ValueError, match="合计不一致"):
        parse(path)
