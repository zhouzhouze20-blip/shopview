from datetime import date
from decimal import Decimal
from pathlib import Path
import sys

import pytest
from sqlalchemy import create_engine, select, func
from sqlalchemy.pool import StaticPool
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python_app"))
from services.self_operated_sales import metadata, publish_export, read_summary, read_details, lines, imports
from routers.authz import DataScope
from routers import self_operated_sales as api

START, END = date(2026, 9, 1), date(2026, 9, 5)


def sample(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "店铺零售报表"
    sheet.append(["会员名称", "单据编号", "渠道编号", "单据日期", "货号", "品名", "计收额", "数量", "计收折扣", "销售方式", "原单号"])
    sheet.append(["私有会员", "XS1", "PBCZ6001", "2026-09-04", "A1", "童装", 1000, 1, 1, "销售", ""])
    sheet.append(["私有会员", "XS1", "PBCZ6001", "2026-09-04", "A1", "童装", 1000, 1, 1, "销售", ""])
    sheet.append(["私有会员", "XS2", "PBCZ6001", "2026-09-05", "A1", "童装", -1000, -1, 1, "退货", "XS1"])
    sheet.append(["合计", None, None, None, None, None, 1000])
    path = tmp_path / "sample.xlsx"
    book.save(path)
    return path


@pytest.fixture
def engine():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    metadata.create_all(engine)
    yield engine
    engine.dispose()


def test_publish_and_repeat_preserve_multiplicity_and_return(engine, tmp_path):
    path = sample(tmp_path)
    publish_export(engine, path, start=START, end=END)
    publish_export(engine, path, start=START, end=END)
    with engine.connect() as db:
        assert db.execute(select(func.count()).select_from(lines)).scalar() == 3
        assert db.execute(select(func.count()).select_from(imports)).scalar() == 2
        result = read_summary(db, ["CLOTHING", "BAKERY"], START, END)
        assert result["totals"]["sales"] == Decimal("1000")
        assert result["totals"]["gross_profit"] == Decimal("350")
        assert result["sections"][0]["coverage"] == "complete"
        assert result["sections"][1]["coverage"] == "not_imported"
        details = read_details(db, "CLOTHING", START, END, 1, 1)
        assert details["total"] == 3 and len(details["rows"]) == 1
        assert "source_data" not in details["rows"][0]
        assert "会员" not in str(details)


def test_coverage_distinguishes_missing_from_zero_sales(engine, tmp_path):
    publish_export(engine, sample(tmp_path), start=START, end=END)
    with engine.connect() as db:
        assert read_summary(db, ["CLOTHING"], START, START)["sections"][0]["coverage"] == "complete"
        assert read_summary(db, ["CLOTHING"], START, date(2026, 9, 6))["sections"][0]["coverage"] == "partial"
        assert read_summary(db, ["CLOTHING"], date(2026, 9, 6), date(2026, 9, 6))["sections"][0]["coverage"] == "not_imported"


@pytest.mark.parametrize("scope,expected", [
    (DataScope(), []),
    (DataScope(allow={"department": {"CLOTHING"}}), ["CLOTHING"]),
    (DataScope(allow={"department": {"BAKERY"}}), ["BAKERY"]),
    (DataScope(allow={"store": {"SELF_OPERATED"}}), ["CLOTHING", "BAKERY"]),
    (DataScope(all_access=True, deny={"department": {"CLOTHING"}}), ["BAKERY"]),
    (DataScope(all_access=True, deny={"store": {"SELF_OPERATED"}}), []),
    (DataScope(all_access=True, deny={"__all__": {"*"}}), []),
    (DataScope(allow={"store": {"601"}}), []),
])
def test_independent_scope_and_denial(monkeypatch, scope, expected):
    monkeypatch.setattr(api, "require_permission", lambda *args: None)
    def load(db, user, resource, action):
        assert resource == "self_operated_sales" and action == "view"
        return scope
    monkeypatch.setattr(api, "load_data_scope", load)
    assert api.allowed_sections(None, None) == expected


def test_unauthorized_summary_does_not_read_data(monkeypatch):
    monkeypatch.setattr(api, "allowed_sections", lambda *args: [])
    monkeypatch.setattr(api, "require_tables", lambda *args: pytest.fail("must not query tables"))
    assert api.summary(START, END, None, None)["sections"] == []
    with pytest.raises(api.HTTPException) as caught:
        api.details("CLOTHING", START, END, 0, 20, None, None)
    assert caught.value.status_code == 403


def test_invalid_date_range():
    with pytest.raises(api.HTTPException):
        api.validate_dates(END, START)


def test_upload_preview_then_commit_and_history(engine, tmp_path):
    content = sample(tmp_path).read_bytes()
    result = api.process_upload(engine, content, "用户报表.xlsx", START, END, False, "", "tester")
    assert result["status"] == "preview"
    with engine.connect() as db:
        assert db.execute(select(func.count()).select_from(lines)).scalar() == 0
    with pytest.raises(api.HTTPException):
        api.process_upload(engine, content, "用户报表.xlsx", START, END, True, "changed", "tester")
    imported = api.process_upload(engine, content, "用户报表.xlsx", START, END, True, result["sha256"], "tester")
    assert imported["status"] == "imported"
    with engine.connect() as db:
        batch = db.execute(select(imports)).mappings().one()
        assert batch["source_file"] == "用户报表.xlsx"
        assert batch["imported_by"] == "tester"


@pytest.mark.parametrize("promotion", ["买满赠送(按金额)", "会员积分兑换货品,会员600积分兑换礼盒"])
def test_gift_upload_preview_commit_and_readback(engine, tmp_path, promotion):
    path = sample(tmp_path)
    book = load_workbook(path)
    sheet = book.active
    sheet.cell(1, 12, "促销信息")
    sheet.insert_rows(5)
    for column, value in enumerate([None, "GIFT1", "PBCZ6001", "2026-09-05", "G1",
                                   "包挂", None, 1, None, "销售", "", promotion], 1):
        sheet.cell(5, column, value)
    book.save(path)
    book.close()
    content = path.read_bytes()
    preview = api.process_upload(engine, content, "gift.xlsx", START, END, False, "", "tester")
    assert preview["row_count"] == 4 and Decimal(preview["sales"]) == 1000
    api.process_upload(engine, content, "gift.xlsx", START, END, True, preview["sha256"], "tester")
    with engine.connect() as db:
        gift = db.execute(select(lines).where(lines.c.product == "G1")).mappings().one()
        assert gift["quantity"] == 1 and gift["discount"] == 0
        for key in ("sales", "supplier_cost", "venue_fee", "expected_receipt", "gross_profit"):
            assert gift[key] == 0
        assert gift["source_data"]["计收额"] is None
        assert read_summary(db, ["CLOTHING"], START, END)["totals"]["sales"] == 1000


def test_exchange_upload_preview_commit_and_readback(engine, tmp_path):
    path = sample(tmp_path)
    book = load_workbook(path)
    book.active.cell(2, 10, "换货")
    book.save(path)
    book.close()
    content = path.read_bytes()
    preview = api.process_upload(engine, content, "exchange.xlsx", START, END, False, "", "tester")
    assert preview["row_count"] == 3 and preview["return_rows"] == 1
    assert Decimal(preview["sales"]) == 1000
    for _ in range(2):
        api.process_upload(engine, content, "exchange.xlsx", START, END, True, preview["sha256"], "tester")
    with engine.connect() as db:
        exchange = db.execute(select(lines).where(lines.c.sale_type == "换货")).mappings().one()
        assert exchange["sales"] == 1000 and exchange["quantity"] == 1
        assert exchange["supplier_cost"] == 500 and exchange["venue_fee"] == 150
        assert exchange["expected_receipt"] == 850 and exchange["gross_profit"] == 350
        assert exchange["source_data"]["销售方式"] == "换货"
        summary = read_summary(db, ["CLOTHING"], START, END)["totals"]
        assert summary["sales"] == 1000 and summary["gross_profit"] == 350
        assert db.execute(select(func.count()).select_from(lines)).scalar() == 3


def test_import_requires_function_and_clothing_scope(monkeypatch):
    permissions = []
    monkeypatch.setattr(api, "require_permission", lambda db, user, permission: permissions.append(permission))
    monkeypatch.setattr(api, "allowed_sections", lambda *args: ["BAKERY"])
    with pytest.raises(api.HTTPException) as error:
        api.require_import_access(None, None)
    assert error.value.status_code == 403
    assert permissions == ["sales.self_operated.import"]


def test_bad_upload_does_not_write(engine):
    with pytest.raises(api.HTTPException) as error:
        api.process_upload(engine, b"not excel", "bad.xlsx", START, END, True, "", "tester")
    assert error.value.status_code == 422
    with engine.connect() as db:
        assert db.execute(select(func.count()).select_from(imports)).scalar() == 0


def test_http_multipart_upload_summary_and_history(engine, tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import Session
    from types import SimpleNamespace
    app = FastAPI()
    app.include_router(api.router)
    def session():
        with Session(engine) as db:
            yield db
    app.dependency_overrides[api.get_db] = session
    app.dependency_overrides[api.get_current_user] = lambda: SimpleNamespace(username="tester")
    monkeypatch.setattr(api, "require_permission", lambda *args: None)
    monkeypatch.setattr(api, "load_data_scope", lambda *args: DataScope(allow={"department": {"CLOTHING"}}))
    client = TestClient(app)
    content = sample(tmp_path).read_bytes()
    fields = {"start_date": str(START), "end_date": str(END)}
    files = {"file": ("sample.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = client.post("/api/self-operated-sales/upload", data=fields, files=files)
    assert response.status_code == 200
    result = client.post("/api/self-operated-sales/upload", data={**fields, "commit": "true",
        "expected_sha256": response.json()["sha256"]}, files=files)
    assert result.status_code == 200 and result.json()["status"] == "imported"
    summary = client.get("/api/self-operated-sales/summary", params=fields)
    assert summary.status_code == 200
    assert [s["code"] for s in summary.json()["sections"]] == ["CLOTHING"]
    assert summary.json()["totals"]["gross_profit"] == 350
    assert client.get("/api/self-operated-sales/imports").json()[0]["imported_by"] == "tester"
    assert client.get("/api/self-operated-sales/details", params={**fields, "section": "BAKERY"}).status_code == 403


def test_store_summary_and_whole_ticket_paging(engine, tmp_path):
    from services.self_operated_sales import read_tickets
    publish_export(engine, sample(tmp_path), start=START, end=END)
    with engine.begin() as db:
        # A second channel may reuse the same ticket number: never merge them.
        copied = dict(db.execute(select(lines).limit(1)).mappings().one())
        copied.pop("id")
        copied["channel"] = "OTHER"
        db.execute(lines.insert().values(**copied))
        result = read_summary(db, ["CLOTHING"], START, END)
        assert result["sections"][0]["ticket_count"] == 3
        store = result["sections"][0]["stores"][0]
        assert store["channel"] == "PBCZ6001"
        assert store["sales"] == Decimal("1000")
        assert store["gross_profit"] == Decimal("350")
        page = read_tickets(db, "CLOTHING", "PBCZ6001", START, END, 0, 1)
        assert page["total"] == 2
        assert page["tickets"][0]["sales"] == -1000
        assert page["tickets"][0]["gross_profit"] == -350
        second = read_tickets(db, "CLOTHING", "PBCZ6001", START, END, 1, 1)
        assert len(second["tickets"][0]["rows"]) == 2
        assert second["tickets"][0]["sales"] == 2000
        assert "source_data" not in str(second)
        assert "私有会员" not in str(second)
        assert read_tickets(db, "BAKERY", "PBCZ6001", START, END)["total"] == 0
        assert read_tickets(db, "CLOTHING", "unknown", START, END)["total"] == 0
        day = read_summary(db, ["CLOTHING"], END, END)
        assert day["totals"]["sales"] == -1000
        assert day["totals"]["gross_profit"] == -350


def test_ticket_endpoint_denies_unscoped_section(monkeypatch):
    monkeypatch.setattr(api, "allowed_sections", lambda *args: ["CLOTHING"])
    monkeypatch.setattr(api, "require_tables", lambda *args: pytest.fail("no database access"))
    with pytest.raises(api.HTTPException) as error:
        api.tickets("BAKERY", "PBCZ6001", START, END, 0, 20, None, None)
    assert error.value.status_code == 403


def test_ticket_tag_price_comes_from_source_without_exposing_raw_data(engine, tmp_path):
    from services.self_operated_sales import read_tickets
    publish_export(engine, sample(tmp_path), start=START, end=END)
    with engine.begin() as db:
        db.execute(lines.update().values(source_data={"吊牌价": 1999, "会员名称": "private"}))
        result = read_tickets(db, "CLOTHING", "PBCZ6001", START, END)
        row = result["tickets"][0]["rows"][0]
        assert Decimal(str(row["tag_price"])) == 1999
        assert row["sales"] == -1000
        assert "source_data" not in row and "private" not in str(result)
        db.execute(lines.update().values(source_data={}))
        assert read_tickets(db, "CLOTHING", "PBCZ6001", START, END)["tickets"][0]["rows"][0]["tag_price"] is None


def test_discount_settlement_bands_rounding_returns_and_scope(engine, tmp_path):
    from openpyxl import load_workbook
    from services.self_operated_sales import read_discount_settlement
    path = sample(tmp_path)
    book = load_workbook(path)
    sheet = book.active
    sheet.delete_rows(2, sheet.max_row)
    cases = [(100, .85), (200, .8499), (300, .70), (400, .6999), (-100, .85), (.01, 1), (.01, 1)]
    for i, (amount, discount) in enumerate(cases):
        sheet.append([None, f"T{i}", "PBCZ6001", "2026-09-04", "A", "童装", amount,
                      -1 if amount < 0 else 1, discount, "退货" if amount < 0 else "销售", ""])
    sheet.append(["合计", None, None, None, None, None, 900.02])
    book.save(path)
    publish_export(engine, path, start=START, end=END)
    with engine.begin() as db:
        copied = dict(db.execute(select(lines).limit(1)).mappings().one())
        copied.pop("id")
        db.execute(lines.insert().values(**{**copied, "channel": "OTHER"}))
        db.execute(lines.insert().values(**{**copied, "section": "BAKERY"}))
        data = read_discount_settlement(db, START, END, "PBCZ6001")
        assert [r["sales"] for r in data["rows"]] == [Decimal('.02'), Decimal(500), Decimal(400)]
        assert [r["remittance"] for r in data["rows"]] == [Decimal('.02'), Decimal(260), Decimal(220)]
        assert data["totals"]["sales"] == Decimal('900.02')
        assert data["totals"]["remittance"] == Decimal('480.02')
        assert [r["our_settlement"] for r in data["rows"]] == [Decimal(0), Decimal(240), Decimal(180)]
        assert data["totals"]["our_settlement"] == Decimal(420)
        assert sum(r["our_settlement"] for r in data["rows"]) == data["totals"]["our_settlement"]
        assert abs(sum(r['sales_share'] for r in data['rows']) - 1) < Decimal('1e-25')
        assert read_discount_settlement(db, START, END)['totals']['sales'] == Decimal('1000.02')
        empty = read_discount_settlement(db, END, END, 'PBCZ6001')
        assert empty['coverage'] == 'complete'
        assert empty['totals']['sales_share'] is None
        assert empty['totals']['our_settlement'] == 0
        assert all(r['sales_share'] is None for r in empty['rows'])
        assert read_discount_settlement(db, date(2026, 9, 6), date(2026, 9, 6))['coverage'] == 'not_imported'
        assert '私有会员' not in str(data)


def test_discount_settlement_endpoint_enforces_clothing_permission(monkeypatch):
    monkeypatch.setattr(api, 'allowed_sections', lambda *args: ['BAKERY'])
    monkeypatch.setattr(api, 'require_tables', lambda *args: pytest.fail('no database access'))
    with pytest.raises(api.HTTPException) as error:
        api.discount_settlement(START, END, None, None, None)
    assert error.value.status_code == 403


def test_discount_settlement_endpoint_passes_dates_and_channel(engine, tmp_path, monkeypatch):
    publish_export(engine, sample(tmp_path), start=START, end=END)
    monkeypatch.setattr(api, 'allowed_sections', lambda *args: ['CLOTHING'])
    from sqlalchemy.orm import Session
    with Session(engine) as db:
        result = api.discount_settlement(END, END, 'PBCZ6001', db, None)
        assert result['totals']['sales'] == -1000
        assert result['totals']['remittance'] == -500
        assert result['totals']['our_settlement'] == -500
        assert result['totals']['sales_share'] == 1
