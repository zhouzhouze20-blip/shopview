import asyncio
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from python_app.services.od0002_report import TrustedScopeSql
from python_app.services.settled_gross_profit_excel import (
    build_settled_gross_profit_workbook,
)
from python_app.services.settled_gross_profit_report import (
    FLOOR_NAMES,
    OPERATION_MODE_NAMES,
    build_quality,
    build_settled_gross_profit_query,
    build_totals,
    normalize_report_rows,
)


def sample_row(**overrides):
    row = {
        "store_code": "603",
        "store_name": "常州新世纪商城",
        "department_code": "6030101",
        "department_name": "新世纪一部(化妆)",
        "group_code": "6030101005",
        "group_name": "L'oreal欧莱雅厅",
        "supplier_code": "01292",
        "supplier_name": "常州市正大百货有限公司",
        "floor_code": "02",
        "floor_name": "1F",
        "business_area": 52,
        "operation_mode_code": "1",
        "operation_mode_name": "经销",
        "area_code": "01",
        "area_name": "化妆区",
        "contract_code": "01292012",
        "sales_qty": 370,
        "sales_revenue": 68685.4001,
        "tax_excluded_sales": 60783.54,
        "front_profit": -16134.5,
        "tax_excluded_profit_adjustment": 20291.95,
        "floor_adjustment": 0,
        "sales_floor_cost": 0,
        "original_rate_profit": 12764.6,
        "appliance_rebate": 0,
        "contract_end_date": date(2026, 12, 31),
        "contract_profit": None,
    }
    row.update(overrides)
    return row


def test_query_reuses_existing_tables_and_binds_filters_and_scope():
    sql, params = build_settled_gross_profit_query(
        date(2026, 5, 29),
        date(2026, 6, 28),
        TrustedScopeSql(" AND st.store_id::text = ANY(:allowed_stores)"),
        {"allowed_stores": ["3"]},
        selected_store=" 603 ",
        selected_department=" 6030101 ",
    )
    compact = " ".join(sql.lower().split())

    for table in (
        "salegoodslist",
        "manaframe",
        "supplierbase",
        "area_category",
        "contmain",
        "contmanaframe",
        "supsetcharge",
        "salecostday",
        "contbd",
    ):
        assert table in compact
    assert "left join lateral" in compact
    assert "cmf.cmfeffdate" not in compact and "cmf.cmflapdate" not in compact
    assert compact.count("< (:end_date + interval '1 day')") >= 5
    assert ":selected_store" in sql and ":selected_department" in sql
    assert ":allowed_stores" in sql
    assert params == {
        "allowed_stores": ["3"],
        "start_date": date(2026, 5, 29),
        "end_date": date(2026, 6, 28),
        "selected_store": "603",
        "selected_department": "6030101",
    }


def test_floor_and_operation_names_match_the_reference_report():
    assert FLOOR_NAMES["02"] == "1F"
    assert FLOOR_NAMES["16"] == "16"
    assert OPERATION_MODE_NAMES == {
        "1": "经销",
        "2": "成本代销",
        "4": "联营",
        "5": "租赁",
    }


def test_normalization_totals_and_quality_keep_blank_contract_profit_distinct():
    rows = normalize_report_rows([
        sample_row(),
        sample_row(
            group_code="6030101010",
            contract_code=None,
            sales_qty=2,
            sales_revenue=100,
            supplier_name="未匹配",
            area_name="未匹配",
        ),
    ])
    totals = build_totals(rows)
    quality = build_quality(rows)

    assert rows[0]["contract_end_date"] == "2026-12-31"
    assert totals["sales_qty"] == 372
    assert totals["sales_revenue"] == 68785.4001
    assert totals["contract_profit"] is None
    assert quality == {
        "row_count": 2,
        "unresolved_contract_group_count": 1,
        "unresolved_contract_sales": 100.0,
        "missing_supplier_name_count": 1,
        "missing_area_name_count": 1,
    }


def test_workbook_matches_reference_columns_and_includes_quality_and_scope():
    row = normalize_report_rows([sample_row()])[0]
    report = {
        "dates": {"start_date": date(2026, 5, 29), "end_date": date(2026, 6, 28)},
        "rows": [row],
        "totals": build_totals([row]),
        "quality": build_quality([row]),
        "scope_description": "当前用户权限范围：门店 603",
    }
    workbook = load_workbook(BytesIO(build_settled_gross_profit_workbook(report)))
    sheet = workbook["结算后销售毛利排行"]

    assert sheet["A1"].value == "结算后销售毛利排行表"
    assert "2026-05-29" in sheet["A2"].value and "2026-06-28" in sheet["A2"].value
    assert sheet["A3"].value == "当前用户权限范围：门店 603"
    assert sheet.max_column == 20
    assert [sheet.cell(5, column).value for column in range(1, 21)][-4:] == [
        "原扣率毛利", "家电返利", "合同到期日期", "合同毛利"
    ]
    assert sheet["A6"].value.startswith("603 ")
    assert sheet["D6"].value.startswith("[01292]")
    assert sheet["K6"].value == 68685.4001
    assert sheet["A7"].value == "合计"
    assert sheet.freeze_panes == "A6"
    assert "柜位系统 PostgreSQL" in workbook["口径说明"]["B9"].value


def test_endpoint_reuses_od0002_permission_and_business_scope(monkeypatch):
    from python_app.routers import sales
    from python_app.routers.authz import DataScope

    calls = {}
    monkeypatch.setattr(
        sales,
        "require_permission",
        lambda db, user, code: calls.setdefault("permission", code),
    )
    monkeypatch.setattr(
        sales,
        "load_business_scope",
        lambda db, user, **kwargs: DataScope(all_access=True),
    )
    monkeypatch.setattr(
        sales,
        "_business_scope_filter_sql",
        lambda scope, params, **kwargs: (
            calls.setdefault("scope_kwargs", kwargs),
            " AND 1=1",
        )[1],
    )
    monkeypatch.setattr(
        sales,
        "load_settled_gross_profit_report",
        lambda *args, **kwargs: (
            calls.setdefault("load_kwargs", kwargs),
            {"rows": []},
        )[1],
    )

    result = asyncio.run(
        sales.settled_gross_profit_report(
            date(2026, 5, 29),
            date(2026, 6, 28),
            "603",
            object(),
            object(),
            "6030101",
        )
    )

    assert result == {"rows": []}
    assert calls["permission"] == "sales.od0002.view"
    assert calls["scope_kwargs"]["store_expr"] == "st.store_id::text"
    assert calls["scope_kwargs"]["group_expr"] == "mf.mfcode"
    assert calls["load_kwargs"]["selected_store"] == "603"
    assert calls["load_kwargs"]["selected_department"] == "6030101"
