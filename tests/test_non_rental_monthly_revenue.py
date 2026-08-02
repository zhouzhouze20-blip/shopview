from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from python_app.routers.authz import CORE_PERMISSION_DEFINITIONS
from python_app.services.non_rental_monthly_revenue_excel import (
    build_non_rental_monthly_revenue_workbook,
)
from python_app.services.non_rental_monthly_revenue_report import (
    aggregate_metrics,
    build_metrics,
    build_non_rental_monthly_revenue_query,
    financial_month_period,
    financial_year_period,
    normalize_non_rental_monthly_revenue_rows,
)
from python_app.services.od0002_report import TrustedScopeSql


def source_row(**overrides):
    row = {
        "store_code": "602",
        "store_name": "常州百货大楼",
        "department_code": "6020101",
        "department_name": "营运一部",
        "group_code": "6020101001",
        "group_name": "Aupres欧珀莱厅",
        "floor_code": "02",
        "floor_name": "1F",
        "area_code": "01",
        "area_name": "化妆区",
        "category_code": "0103",
        "category_name": "合资品",
        "big_category": "百货类",
        "brand_code": "00703",
        "brand_name": "Aupres欧珀莱",
        "financial_month": 1,
        "tax_included_sales": 9291.2,
        "tax_excluded_sales": 8222.3,
        "gross_profit": 1051.24535,
        "fee": 0,
        "contract_profit": 2040,
        "latest_sales_date": date(2026, 1, 28),
    }
    row.update(overrides)
    return row


def test_financial_months_use_calendar_year_and_extend_december_to_month_end():
    assert financial_month_period(2026, 1) == (
        date(2026, 1, 1),
        date(2026, 1, 28),
    )
    assert financial_month_period(2026, 7) == (
        date(2026, 6, 29),
        date(2026, 7, 28),
    )
    assert financial_month_period(2026, 3) == (
        date(2026, 3, 1),
        date(2026, 3, 28),
    )
    assert financial_month_period(2024, 3) == (
        date(2024, 2, 29),
        date(2024, 3, 28),
    )
    assert financial_month_period(2026, 12) == (
        date(2026, 11, 29),
        date(2026, 12, 31),
    )
    assert financial_year_period(2026) == (
        date(2026, 1, 1),
        date(2026, 12, 31),
    )


def test_metrics_omit_buy_amount_and_recalculate_rates_from_aggregated_amounts():
    january = build_metrics(
        tax_included_sales=100,
        tax_excluded_sales=80,
        gross_profit=20,
        fee=3,
        contract_profit=16,
    )
    february = build_metrics(
        tax_included_sales=200,
        tax_excluded_sales=100,
        gross_profit=10,
        fee=7,
        contract_profit=12,
    )
    total = aggregate_metrics([january, february])

    assert "buy_amount" not in total
    assert total["contribution"] == 40
    assert total["gross_margin"] == pytest.approx(30 / 180)
    assert total["concession_loss"] == 2
    assert total["concession_loss_rate"] == pytest.approx(2 / 180)


def test_normalization_separates_special_sale_and_adds_department_totals():
    result = normalize_non_rental_monthly_revenue_rows(
        [
            source_row(),
            source_row(
                group_code="6020101093",
                group_name="特卖厅",
                floor_code="16",
                floor_name="特卖",
                area_code="14",
                area_name="特卖区",
                category_code="1401",
                category_name="特卖商品",
                brand_code="1026",
                brand_name="resimple",
                financial_month=2,
                tax_included_sales=500,
                tax_excluded_sales=450,
                gross_profit=50,
                fee=0,
                contract_profit=45,
                latest_sales_date=date(2026, 2, 28),
            ),
            source_row(
                group_code="6020101093",
                group_name="特卖厅",
                floor_code="16",
                floor_name="特卖",
                area_code="14",
                area_name="特卖区",
                category_code="1401",
                category_name="特卖商品",
                brand_code="",
                brand_name="收费未分配品牌",
                financial_month=2,
                tax_included_sales=0,
                tax_excluded_sales=0,
                gross_profit=0,
                fee=10,
                contract_profit=0,
                latest_sales_date=None,
            ),
        ]
    )

    assert [row["row_type"] for row in result["regular_rows"]] == [
        "brand",
        "department_total",
        "store_total",
    ]
    assert [row["row_type"] for row in result["special_rows"]] == [
        "brand",
        "brand",
        "special_total",
    ]
    assert result["special_rows"][0]["group_name"] == "特卖厅"
    assert result["special_rows"][0]["brand_name"] == "resimple"
    assert result["special_rows"][1]["brand_name"] == "收费未分配品牌"
    assert result["special_rows"][-1]["annual"]["fee"] == 10
    assert result["quality"]["special_brand_count"] == 1
    assert result["grand_total"]["annual"]["tax_included_sales"] == 9791.2
    assert result["quality"]["latest_sales_date"] == "2026-02-28"


def test_query_uses_non_rental_filter_exact_gross_profit_and_fee_tax_rates():
    sql, params = build_non_rental_monthly_revenue_query(
        2026,
        TrustedScopeSql(" AND st.store_id::text = ANY(:allowed_stores)"),
        {"allowed_stores": ["2"]},
        selected_store=" 602 ",
        selected_department=" 6020101 ",
    )
    compact = " ".join(sql.lower().split())

    assert "trim(coalesce(s.sglwmid, '')) <> '5'" in compact
    assert "extract(month from s.sglhsrq) = 12 then 12" in compact
    assert "extract(day from s.sglhsrq) >= 29" in compact
    assert "left join codebrand cb" in compact
    assert "收费未分配品牌" in sql
    assert "union all" in compact
    assert "coalesce(s.sgln13, 0)" in compact
    assert "coalesce(s.sgln14, 0)" in compact
    assert "coalesce(s.sglsupzk, 0)" in compact
    assert "round(sum(" in compact
    assert "supsetcharge" in compact and "codecharge" in compact
    assert "extract(month from fee.sscfsdate) + 1" in compact
    assert "not in ('38', '61', '94', '95')" in compact
    assert ":selected_department" in sql
    assert params["start_date"] == date(2026, 1, 1)
    assert params["end_date"] == date(2026, 12, 31)
    assert params["selected_store"] == "602"
    assert params["selected_department"] == "6020101"
    assert params["fee_start_date"] == date(2025, 12, 1)
    assert params["fee_end_date"] == date(2026, 12, 1)


def test_workbook_matches_grouped_monthly_layout_without_buy_amount():
    normalized = normalize_non_rental_monthly_revenue_rows(
        [
            source_row(),
            source_row(
                group_code="6020101093",
                group_name="特卖厅",
                floor_code="16",
                floor_name="特卖",
                brand_code="1026",
                brand_name="resimple",
                financial_month=2,
                tax_included_sales=500,
                tax_excluded_sales=450,
                gross_profit=50,
                fee=10,
                contract_profit=45,
            ),
        ]
    )
    report = {
        "financial_year": 2026,
        "dates": {"start_date": "2026-01-01", "end_date": "2026-12-31"},
        "scope_description": "当前用户权限范围",
        **normalized,
    }
    workbook = load_workbook(BytesIO(build_non_rental_monthly_revenue_workbook(report)))

    assert workbook.sheetnames == ["2026年", "2026年特卖", "口径说明"]
    regular = workbook["2026年"]
    special = workbook["2026年特卖"]
    assert regular["A1"].value == "大类"
    assert regular["G1"].value == "2026年小计"
    assert regular["G2"].value == "销售含税"
    assert "买单" not in [
        regular.cell(2, column).value for column in range(1, regular.max_column + 1)
    ]
    assert special["F1"].value == "品牌"
    assert special["G1"].value == "品牌厅"
    assert special["F4"].value == "resimple"
    assert special["G4"].value == "特卖厅"
    assert regular.freeze_panes == "G4"
    assert special.freeze_panes == "H4"


def test_permission_is_registered_and_migration_preserves_od0002_role_access():
    registered = {permission[0] for permission in CORE_PERMISSION_DEFINITIONS}
    assert "sales.non_rental_monthly_revenue.view" in registered
    migration = Path(
        "python_app/alembic/versions/"
        "p4e5f6a7b8c9_add_non_rental_monthly_revenue_permission.py"
    ).read_text(encoding="utf-8")
    assert "'sales.non_rental_monthly_revenue.view'" in migration
    assert "source.permission_code = 'sales.od0002.view'" in migration
